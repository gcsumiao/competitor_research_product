"""Single-pass openpyxl workbook engine for ca_market_reports (house style from ca_common).

Contract (ca_common "Formula policy", findings 20/21/23/24):
  * a workbook is built ONCE in memory and written with ONE `wb.save` (Book.save). Built outputs are never reopened and
    re-saved by this package: openpyxl drops charts on a reload.
  * table layout is computed before anything is written; rows are never shifted after the fact.
  * every numeric cell is a static value computed in pandas; the only formula is the HYPERLINK of a link column, which
    matches ca_common.ALLOWED_FORMULA_RES (destination + display carry the row ASIN and the row's market domain).
  * every text cell is written as a string (data_type "s"): a title starting with "=" never becomes a formula.

Also hosts the builder support shared by both report builders: normalized-frame loading (dev/fixture mode), brand and
listing aggregations, the table registry, the manifest and the overwrite/backup guard.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
import shutil
import subprocess
import weakref
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ca_market_reports import ca_common as C
from ca_market_reports.ca_common import Market

# --------------------------------------------------------------------------------------
# Specs
# --------------------------------------------------------------------------------------
KINDS: tuple[str, ...] = ("text", "int", "money", "money2", "pct", "pct2", "rating", "link", "bool")
NUMERIC_KINDS: tuple[str, ...] = ("int", "money", "money2", "pct", "pct2", "rating")
FMT_PCT = "0.0%"
FMT_PCT2 = "0.00%"            # small market shares (gauge share of the code-reader market)
FMT_RATING = "0.0"
FMT_INT = "#,##0"
TOTAL_POSITIONS = ("bottom", "top")


@dataclass(frozen=True)
class ColumnSpec:
    header: str
    field: str
    kind: str
    width: float = 14.0
    market: str | None = None      # override of the table market for money formats / link domain (cross-market columns)

    def __post_init__(self):
        if self.kind not in KINDS:
            raise ValueError(f"ColumnSpec {self.header!r}: unknown kind {self.kind!r}; expected one of {KINDS}")
        if self.market is not None and self.market not in C.MARKETS:
            raise ValueError(f"ColumnSpec {self.header!r}: unknown market {self.market!r}")
        if "{ccy}" in self.header:
            raise ValueError(f"ColumnSpec header still carries an unformatted {{ccy}}: {self.header!r}")


@dataclass
class TableSpec:
    sheet: str
    role: str
    title: str
    columns: list[ColumnSpec]
    rows: pd.DataFrame
    total_row: dict | None
    residual_row: dict | None
    allowed_markets: tuple[str, ...]
    subtitle: str | None = None          # second banner row (market label, export dates, report month)
    note: str | None = None              # third banner row (method note), italic
    dataset_filter: str = ""             # registry: which dataset rows the table covers
    total_position: str = "bottom"       # "top": the Total row sits directly under the header (Brand x Tier)


@dataclass
class TableRange:
    sheet: str
    role: str
    header_row: int | None
    first_data_row: int
    last_data_row: int                   # == first_data_row - 1 when the table has no data rows
    total_row: int | None
    residual_row: int | None
    first_col: int
    last_col: int
    columns: list[str]
    charts: list = field(default_factory=list)
    title: str = ""
    dataset_filter: str = ""
    allowed_markets: tuple[str, ...] = ()

    @property
    def n_rows(self) -> int:
        return self.last_data_row - self.first_data_row + 1

    def col(self, header_or_index: str | int) -> int:
        if isinstance(header_or_index, int):
            return header_or_index
        if header_or_index not in self.columns:
            raise KeyError(f"{self.sheet}/{self.role}: no column {header_or_index!r}; have {self.columns}")
        return self.first_col + self.columns.index(header_or_index)

    def to_registry(self, workbook: str, brand_sheet_map: dict[str, str]) -> dict:
        return {"workbook": workbook, "sheet": self.sheet, "role": self.role, "title": self.title,
                "header_row": self.header_row, "first_data_row": self.first_data_row, "last_data_row": self.last_data_row,
                "total_row": self.total_row, "residual_row": self.residual_row, "first_col": self.first_col,
                "last_col": self.last_col, "columns": list(self.columns), "dataset_filter": self.dataset_filter,
                "allowed_markets": list(self.allowed_markets), "charts": list(self.charts),
                "brand_sheet_map": dict(brand_sheet_map)}


# --------------------------------------------------------------------------------------
# Styles
# --------------------------------------------------------------------------------------
_THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
FILL_TITLE = PatternFill("solid", fgColor=C.FILL_TITLE)
FILL_HEADER = PatternFill("solid", fgColor=C.FILL_HEADER)
FILL_ALT = PatternFill("solid", fgColor=C.FILL_ALT_ROW)
FILL_TOTAL = PatternFill("solid", fgColor=C.FILL_TOTAL)
FONT_TITLE = Font(bold=True, size=14)
FONT_SUBTITLE = Font(italic=True, size=10, color="404040")
FONT_NOTE = Font(italic=True, size=10, color="404040")
FONT_HEADER = Font(bold=True)
FONT_TOTAL = Font(bold=True, color="FFFFFF")
FONT_RESIDUAL = Font(italic=True)
FONT_INNOVA = Font(color=C.COLOR_INNOVA)
FONT_LINK = Font(color=C.COLOR_LINK, underline="single")
FONT_LINK_TOTAL = Font(color="FFFFFF", underline="single")


def month_label(month: str) -> str:
    """'202609' -> "Sep '26"."""
    if not C.MONTH_RE.match(str(month)):
        raise ValueError(f"invalid month {month!r}")
    d = date(int(month[:4]), int(month[4:]), 1)
    return d.strftime("%b '") + month[2:4]


def subtitle_text(market: Market, export_dates: tuple[date, date], month: str) -> str:
    d0, d1 = export_dates
    return f"{market.label} · export {d0.isoformat()}–{d1.isoformat()} · Report month {month_label(month)}"


def _is_blank(v: Any) -> bool:
    if v is None or v is pd.NA or v is pd.NaT:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    try:
        return bool(pd.isna(v)) if not isinstance(v, (str, bytes, list, dict, tuple)) else False
    except (TypeError, ValueError):
        return False


def set_text(cell, value: Any) -> None:
    """Write `value` as a string cell (never a formula, even when it starts with '=')."""
    if _is_blank(value):
        cell.value = None
        return
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    cell.value = str(value)
    cell.data_type = "s"


def _money_fmt(kind: str, market: Market) -> str:
    return market.money_fmt if kind == "money" else market.money_fmt_2dp


def hyperlink_formula(asin: str, market: Market) -> str:
    asin = str(asin)
    if not C.ASIN_RE.match(asin):
        raise ValueError(f"link column: invalid ASIN {asin!r}")
    tld = market.domain.rsplit(".", 1)[-1]
    formula = f'=HYPERLINK("{market.url(asin)}","{C.HYPERLINK_TEXT.format(tld=tld, asin=asin)}")'
    if not any(rx.match(formula) for rx in C.ALLOWED_FORMULA_RES):
        raise ValueError(f"generated HYPERLINK does not match ALLOWED_FORMULA_RES: {formula}")
    return formula


def _row_market(col: ColumnSpec, rec: dict, market: Market) -> Market:
    if col.market:
        return C.MARKETS[col.market]
    rm = rec.get("market")
    if isinstance(rm, str) and rm:
        if rm not in C.MARKETS:
            raise ValueError(f"row market {rm!r} unknown")
        return C.MARKETS[rm]
    return market


def _write_value(cell, v: Any, col: ColumnSpec, rec: dict, market: Market) -> None:
    if col.kind == "link":
        if _is_blank(v):
            cell.value = None
            return
        cell.value = hyperlink_formula(v, _row_market(col, rec, market))
        cell.font = FONT_LINK
        return
    if _is_blank(v):
        cell.value = None
        return
    if isinstance(v, str):
        if col.kind == "bool" and v.strip() == "":
            cell.value = None
            return
        set_text(cell, v)
        return
    if col.kind == "text":
        set_text(cell, v)
        return
    if col.kind == "bool":
        if not isinstance(v, (bool,)) and type(v).__name__ != "bool_":
            raise TypeError(f"column {col.header!r}: bool expected, got {v!r}")
        set_text(cell, "Y" if bool(v) else "N")
        return
    # numeric kinds
    if isinstance(v, bool) or type(v).__name__ == "bool_":
        raise TypeError(f"column {col.header!r}: numeric expected, got bool {v!r}")
    x = float(v)
    if not math.isfinite(x):
        raise ValueError(f"column {col.header!r}: non-finite value {v!r}")
    if col.kind == "int":
        cell.value = int(x) if x.is_integer() else x
        cell.number_format = FMT_INT
    elif col.kind in ("money", "money2"):
        cell.value = x
        mk = C.MARKETS[col.market] if col.market else market
        cell.number_format = _money_fmt(col.kind, mk)
    elif col.kind == "pct":
        cell.value = x
        cell.number_format = FMT_PCT
    elif col.kind == "pct2":
        cell.value = x
        cell.number_format = FMT_PCT2
    elif col.kind == "rating":
        cell.value = x
        cell.number_format = FMT_RATING


def _banner(ws: Worksheet, row: int, first_col: int, last_col: int, text: str, style: str) -> None:
    cell = ws.cell(row, first_col)
    set_text(cell, text)
    if style == "title":
        cell.font = FONT_TITLE
        for c in range(first_col, last_col + 1):
            ws.cell(row, c).fill = FILL_TITLE
        ws.row_dimensions[row].height = 22
    elif style == "subtitle":
        cell.font = FONT_SUBTITLE
    elif style == "note":
        cell.font = FONT_NOTE
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        width = sum((ws.column_dimensions[get_column_letter(c)].width or 13) for c in range(first_col, last_col + 1))
        lines = max(1, math.ceil(len(text) * 1.1 / max(width, 20)))
        ws.row_dimensions[row].height = 15 * lines
    else:
        raise ValueError(style)
    if last_col > first_col:
        ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)


def _set_width(ws: Worksheet, col_idx: int, width: float) -> None:
    cd = ws.column_dimensions[get_column_letter(col_idx)]
    cur = cd.width if cd.customWidth else 0
    cd.width = max(cur or 0, width)


def write_sheet_header(ws: Worksheet, title: str, subtitle: str | None, ncols: int, *, note: str | None = None,
                       first_col: int = 1) -> int:
    """Sheet-level banner rows (title, optional subtitle, optional note) starting at row 1; returns the next free row."""
    last = first_col + max(ncols, 1) - 1
    r = 1
    _banner(ws, r, first_col, last, title, "title")
    r += 1
    if subtitle:
        _banner(ws, r, first_col, last, subtitle, "subtitle")
        r += 1
    if note:
        _banner(ws, r, first_col, last, note, "note")
        r += 1
    return r


def _validate_spec(spec: TableSpec) -> None:
    if spec.role not in C.TABLE_ROLES:
        raise ValueError(f"table {spec.title!r}: role {spec.role!r} not in TABLE_ROLES")
    if spec.total_position not in TOTAL_POSITIONS:
        raise ValueError(f"total_position must be one of {TOTAL_POSITIONS}")
    if not spec.columns:
        raise ValueError(f"table {spec.title!r}: no columns")
    if not isinstance(spec.rows, pd.DataFrame):
        raise TypeError(f"table {spec.title!r}: rows must be a DataFrame")
    fields = {c.field for c in spec.columns}
    missing = sorted(f for f in fields if f not in spec.rows.columns)
    if missing:
        raise KeyError(f"table {spec.title!r} ({spec.sheet}): rows are missing fields {missing}")
    for name, extra in (("total_row", spec.total_row), ("residual_row", spec.residual_row)):
        if extra is not None:
            unknown = sorted(set(extra) - fields)
            if unknown:
                raise KeyError(f"table {spec.title!r}: {name} has fields not in the columns: {unknown}")


def write_table(ws: Worksheet, spec: TableSpec, top_row: int, market: Market, *, first_col: int = 1,
                freeze: bool = True) -> TableRange:
    """Write one table: title, [subtitle], [note], header, [Total if top], data, [residual], [Total if bottom].

    The whole layout is computed from the spec; freeze panes go under the header of the first frozen table on a sheet.
    """
    _validate_spec(spec)
    if spec.sheet != ws.title:
        raise ValueError(f"spec.sheet {spec.sheet!r} != worksheet {ws.title!r}")
    ncol = len(spec.columns)
    last_col = first_col + ncol - 1
    for j, col in enumerate(spec.columns):
        _set_width(ws, first_col + j, col.width)
    r = top_row
    _banner(ws, r, first_col, last_col, spec.title, "title")
    r += 1
    if spec.subtitle:
        _banner(ws, r, first_col, last_col, spec.subtitle, "subtitle")
        r += 1
    if spec.note:
        _banner(ws, r, first_col, last_col, spec.note, "note")
        r += 1
    header_row = r
    max_lines = 1
    for j, col in enumerate(spec.columns):
        cell = ws.cell(header_row, first_col + j)
        set_text(cell, col.header)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        max_lines = max(max_lines, math.ceil(len(col.header) / max(col.width - 1, 4)))
    ws.row_dimensions[header_row].height = 15 * min(max_lines, 4) + 2
    r += 1

    def emit(rec: dict, style: str) -> None:
        nonlocal r
        innova = style in ("plain", "alt") and rec.get("brand_key") == "innova"
        for j, col in enumerate(spec.columns):
            cell = ws.cell(r, first_col + j)
            v = rec.get(col.field)
            if style in ("total", "residual") and col.kind == "link":
                v = None
            _write_value(cell, v, col, rec, market)
            cell.border = BORDER
            if style == "total":
                cell.fill = FILL_TOTAL
                cell.font = FONT_TOTAL
            else:
                if style == "alt":
                    cell.fill = FILL_ALT
                if col.kind != "link" or cell.value is None:
                    if style == "residual":
                        cell.font = FONT_RESIDUAL
                    elif innova:
                        cell.font = FONT_INNOVA
            if col.kind == "text" and col.width >= 30:
                cell.alignment = Alignment(wrap_text=False, vertical="top")
        r += 1

    total_idx = None
    if spec.total_row is not None and spec.total_position == "top":
        total_idx = r
        emit(spec.total_row, "total")
    first_data = r
    for i, rec in enumerate(spec.rows.to_dict("records")):
        emit(rec, "alt" if i % 2 else "plain")
    last_data = r - 1
    residual_idx = None
    if spec.residual_row is not None:
        residual_idx = r
        emit(spec.residual_row, "residual")
    if spec.total_row is not None and spec.total_position == "bottom":
        total_idx = r
        emit(spec.total_row, "total")
    if freeze and ws.freeze_panes is None:
        ws.freeze_panes = ws.cell(header_row + 1, 1).coordinate
    return TableRange(sheet=ws.title, role=spec.role, header_row=header_row, first_data_row=first_data, last_data_row=last_data,
                      total_row=total_idx, residual_row=residual_idx, first_col=first_col, last_col=last_col,
                      columns=[c.header for c in spec.columns], charts=[], title=spec.title,
                      dataset_filter=spec.dataset_filter, allowed_markets=tuple(spec.allowed_markets))


def table_end(tr: TableRange) -> int:
    """Last occupied row of a table (data, residual or Total)."""
    return max(x for x in (tr.last_data_row, tr.header_row or 0, tr.total_row or 0, tr.residual_row or 0))


def write_kpi_block(ws: Worksheet, top_row: int, items: Sequence[tuple[str, Any, str]], market: Market, *,
                    title: str | None = None, header: bool = True, role: str = "kpi", dataset_filter: str = "",
                    allowed_markets: tuple[str, ...] | None = None) -> TableRange:
    """Label/value rows with a per-row number kind. Brand tabs use rows 3-6 with no title and no header."""
    if role not in C.TABLE_ROLES:
        raise ValueError(role)
    r = top_row
    if title:
        _banner(ws, r, 1, 2, title, "title")
        r += 1
    header_row = None
    if header:
        header_row = r
        for j, h in enumerate(("Metric", "Value")):
            cell = ws.cell(r, 1 + j)
            set_text(cell, h)
            cell.font, cell.fill, cell.border = FONT_HEADER, FILL_HEADER, BORDER
        r += 1
    first = r
    for label, value, kind in items:
        lc = ws.cell(r, 1)
        set_text(lc, label)
        lc.font = Font(bold=True)
        lc.border = BORDER
        vc = ws.cell(r, 2)
        _write_value(vc, value, ColumnSpec(label, "v", kind), {}, market)
        vc.border = BORDER
        r += 1
    _set_width(ws, 1, 34)
    _set_width(ws, 2, 16)
    return TableRange(sheet=ws.title, role=role, header_row=header_row, first_data_row=first, last_data_row=r - 1,
                      total_row=None, residual_row=None, first_col=1, last_col=2, columns=["Metric", "Value"], charts=[],
                      title=title or "", dataset_filter=dataset_filter,
                      allowed_markets=tuple(allowed_markets or (market.code,)))


def write_line(ws: Worksheet, row: int, text: str, market: Market, *, role: str, dataset_filter: str = "",
               bold: bool = True) -> TableRange:
    """A single static text line registered as a one-row table (e.g. the gauge Innova device count)."""
    if role not in C.TABLE_ROLES:
        raise ValueError(role)
    cell = ws.cell(row, 1)
    set_text(cell, text)
    cell.font = Font(bold=bold)
    return TableRange(sheet=ws.title, role=role, header_row=None, first_data_row=row, last_data_row=row, total_row=None,
                      residual_row=None, first_col=1, last_col=1, columns=["Line"], charts=[], title=text,
                      dataset_filter=dataset_filter, allowed_markets=(market.code,))


def hide_zero_revenue_rows(ws: Worksheet, table_range: TableRange) -> int:
    """Hide (never delete) data rows whose revenue is exactly 0. Only for brand summary tables. Returns the hidden count."""
    if table_range.role != "summary_brands":
        raise ValueError(f"hide_zero_revenue_rows is only for summary_brands tables, not {table_range.role!r}")
    rev_cols = [h for h in table_range.columns if h.startswith("Monthly Rev (")]
    if len(rev_cols) != 1:
        raise KeyError(f"{table_range.sheet}: expected exactly one 'Monthly Rev (<ccy>)' column, got {rev_cols}")
    c = table_range.col(rev_cols[0])
    hidden = 0
    for r in range(table_range.first_data_row, table_range.last_data_row + 1):
        v = ws.cell(r, c).value
        if isinstance(v, (int, float)) and v == 0:
            ws.row_dimensions[r].hidden = True
            hidden += 1
    return hidden


def write_metadata_sheet(ws: Worksheet, entries: dict, market: Market | None = None) -> TableRange:
    missing = [k for k in C.METADATA_REQUIRED_KEYS if k not in entries]
    if missing:
        raise KeyError(f"metadata is missing required keys {missing}")
    for k, v in entries.items():
        if _is_blank(v) or str(v).strip() == "":
            raise ValueError(f"metadata key {k!r} has an empty value")
    _banner(ws, 1, 1, 2, "Metadata", "title")
    header_row = 2
    for j, h in enumerate(("Key", "Value")):
        cell = ws.cell(header_row, 1 + j)
        set_text(cell, h)
        cell.font, cell.fill, cell.border = FONT_HEADER, FILL_HEADER, BORDER
    ordered = list(C.METADATA_REQUIRED_KEYS) + [k for k in entries if k not in C.METADATA_REQUIRED_KEYS]
    r = header_row + 1
    for i, k in enumerate(ordered):
        kc, vc = ws.cell(r, 1), ws.cell(r, 2)
        set_text(kc, k)
        set_text(vc, entries[k])
        kc.font = Font(bold=True)
        kc.alignment = Alignment(vertical="top")
        vc.alignment = Alignment(wrap_text=True, vertical="top")
        for c in (kc, vc):
            c.border = BORDER
            if i % 2:
                c.fill = FILL_ALT
        lines = sum(max(1, math.ceil(len(part) / 95)) for part in str(entries[k]).split("\n"))
        if lines > 1:
            ws.row_dimensions[r].height = 15 * lines
        r += 1
    _set_width(ws, 1, 26)
    _set_width(ws, 2, 100)
    ws.freeze_panes = "A3"
    return TableRange(sheet=ws.title, role="metadata", header_row=header_row, first_data_row=header_row + 1, last_data_row=r - 1,
                      total_row=None, residual_row=None, first_col=1, last_col=2, columns=["Key", "Value"], charts=[],
                      title="Metadata", dataset_filter="", allowed_markets=tuple(C.MARKETS) if market is None else (market.code,))


def write_text_sheet(ws: Worksheet, title: str, paragraphs: Iterable[str], *, role: str = "read_me",
                     start_row: int = 1) -> TableRange:
    """Title banner then one paragraph per row in column A (wrapped). A paragraph starting with '# ' is a bold heading."""
    if role not in C.TABLE_ROLES:
        raise ValueError(role)
    _set_width(ws, 1, 120)
    _banner(ws, start_row, 1, 1, title, "title")
    r = start_row + 2
    first = r
    for p in paragraphs:
        cell = ws.cell(r, 1)
        if p.startswith("# "):
            set_text(cell, p[2:])
            cell.font = Font(bold=True, size=12)
        else:
            set_text(cell, p)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            lines = sum(max(1, math.ceil(len(part) / 115)) for part in p.split("\n"))
            if lines > 1:
                ws.row_dimensions[r].height = 15 * lines
        r += 1
    return TableRange(sheet=ws.title, role=role, header_row=None, first_data_row=first, last_data_row=r - 1, total_row=None,
                      residual_row=None, first_col=1, last_col=1, columns=["Text"], charts=[], title=title,
                      dataset_filter="", allowed_markets=())


# --------------------------------------------------------------------------------------
# Charts
# --------------------------------------------------------------------------------------
_CHART_SLOTS: "weakref.WeakKeyDictionary[Worksheet, list[tuple[int, int, int, int]]]" = weakref.WeakKeyDictionary()
_ROW_CM = 0.53
_COL_CM = 1.6


def _axes_guard(chart) -> None:
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.tickLblPos = "nextTo"
    chart.y_axis.tickLblPos = "nextTo"


def _anchor(ws: Worksheet, tr: TableRange, chart, anchor: str | None) -> str:
    rows = math.ceil(chart.height / _ROW_CM) + 1
    cols = math.ceil(chart.width / _COL_CM) + 1
    slots = _CHART_SLOTS.setdefault(ws, [])
    if anchor is None:
        col = max(tr.last_col, ws.max_column) + 2
        row = tr.header_row or tr.first_data_row
        while any(not (col + cols <= c0 or col >= c1 or row + rows <= r0 or row >= r1) for c0, c1, r0, r1 in slots):
            col = max(c1 for c0, c1, r0, r1 in slots if not (row + rows <= r0 or row >= r1)) + 1
        anchor = f"{get_column_letter(col)}{row}"
    m = re.match(r"^([A-Z]+)(\d+)$", anchor)
    if not m:
        raise ValueError(f"bad chart anchor {anchor!r}")
    from openpyxl.utils import column_index_from_string
    c0, r0 = column_index_from_string(m.group(1)), int(m.group(2))
    slots.append((c0, c0 + cols, r0, r0 + rows))
    return anchor


def _data_last(tr: TableRange, include_residual: bool) -> int:
    if tr.n_rows <= 0:
        raise ValueError(f"{tr.sheet}/{tr.role}: cannot chart a table without data rows")
    if include_residual and tr.residual_row:
        if tr.residual_row != tr.last_data_row + 1:
            raise ValueError("residual row must directly follow the data rows to be charted")
        return tr.residual_row
    return tr.last_data_row


def _point_colors(series, n: int) -> None:
    for i in range(n):
        pt = DataPoint(idx=i)
        color = C.BRAND_PALETTE[i % len(C.BRAND_PALETTE)]
        pt.graphicalProperties.solidFill = color
        pt.graphicalProperties.line.solidFill = color
        series.dPt.append(pt)


def _labels(show_val: bool = False, show_pct: bool = False) -> DataLabelList:
    d = DataLabelList()
    d.showVal = show_val
    d.showPercent = show_pct
    d.showCatName = False
    d.showSerName = False
    d.showLegendKey = False
    d.showLeaderLines = False
    return d


def add_bar_chart(ws: Worksheet, table_range: TableRange, cat_col: str | int, val_col: str | int, anchor: str | None,
                  title: str, *, include_residual: bool = False, horizontal: bool = True) -> dict:
    tr = table_range
    last = _data_last(tr, include_residual)
    n = last - tr.first_data_row + 1
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = title
    chart.style = 10
    data = Reference(ws, min_col=tr.col(val_col), min_row=tr.header_row, max_row=last)
    cats = Reference(ws, min_col=tr.col(cat_col), min_row=tr.first_data_row, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    s = chart.series[0]
    _point_colors(s, n)
    s.dLbls = _labels(show_val=True)
    chart.varyColors = True
    chart.legend = None
    _axes_guard(chart)
    if horizontal:
        chart.x_axis.scaling.orientation = "maxMin"     # rank 1 on top
    chart.visible_cells_only = False
    chart.width = 16
    chart.height = max(7.5, 0.55 * n + 2.5)
    a = _anchor(ws, tr, chart, anchor)
    ws.add_chart(chart, a)
    entry = {"type": "bar", "anchor": a, "has_axes": True}
    tr.charts.append(entry)
    return entry


def add_pie_chart(ws: Worksheet, table_range: TableRange, cat_col: str | int, val_col: str | int, anchor: str | None,
                  title: str, *, include_residual: bool = False) -> dict:
    tr = table_range
    last = _data_last(tr, include_residual)
    n = last - tr.first_data_row + 1
    chart = PieChart()
    chart.title = title
    data = Reference(ws, min_col=tr.col(val_col), min_row=tr.header_row, max_row=last)
    cats = Reference(ws, min_col=tr.col(cat_col), min_row=tr.first_data_row, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    s = chart.series[0]
    _point_colors(s, n)
    s.dLbls = _labels(show_pct=True)
    chart.varyColors = True
    chart.visible_cells_only = False
    chart.width = 14
    chart.height = 9
    a = _anchor(ws, tr, chart, anchor)
    ws.add_chart(chart, a)
    entry = {"type": "pie", "anchor": a, "has_axes": False}
    tr.charts.append(entry)
    return entry


def add_scatter_chart(ws: Worksheet, table_range: TableRange, cat_col: str | int, val_col: str | int, anchor: str | None,
                      title: str) -> dict:
    """x = cat_col values, y = val_col values (data rows only)."""
    tr = table_range
    last = _data_last(tr, False)
    chart = ScatterChart()
    chart.title = title
    chart.style = 13
    xs = Reference(ws, min_col=tr.col(cat_col), min_row=tr.first_data_row, max_row=last)
    ys = Reference(ws, min_col=tr.col(val_col), min_row=tr.first_data_row, max_row=last)
    s = Series(ys, xs, title=title)
    s.marker.symbol = "circle"
    s.marker.size = 7
    s.marker.graphicalProperties.solidFill = C.BRAND_PALETTE[0]
    s.graphicalProperties.line.noFill = True
    chart.series.append(s)
    chart.x_axis.title = tr.columns[tr.col(cat_col) - tr.first_col]
    chart.y_axis.title = tr.columns[tr.col(val_col) - tr.first_col]
    chart.legend = None
    _axes_guard(chart)
    chart.visible_cells_only = False
    chart.width = 16
    chart.height = 9
    a = _anchor(ws, tr, chart, anchor)
    ws.add_chart(chart, a)
    entry = {"type": "scatter", "anchor": a, "has_axes": True}
    tr.charts.append(entry)
    return entry


# --------------------------------------------------------------------------------------
# Book: one in-memory workbook + its registered tables
# --------------------------------------------------------------------------------------
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


class Book:
    def __init__(self, name: str, market: Market):
        if not name.endswith(".xlsx"):
            raise ValueError(f"workbook name must end with .xlsx: {name!r}")
        self.name = name
        self.market = market
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.tables: list[TableRange] = []
        self.brand_sheet_map: dict[str, str] = {}
        self._saved = False

    # sheets -------------------------------------------------------------------------
    def sheet(self, title: str) -> Worksheet:
        if len(title) > 31 or not title or _SHEET_BAD.search(title):
            raise ValueError(f"invalid sheet name {title!r}")
        if title.lower() in {s.lower() for s in self.wb.sheetnames}:
            raise ValueError(f"duplicate sheet name (case-insensitive) {title!r}")
        return self.wb.create_sheet(title)

    @property
    def sheetnames(self) -> list[str]:
        return list(self.wb.sheetnames)

    # tables -------------------------------------------------------------------------
    def _check_markets(self, ws_title: str, spec: TableSpec) -> None:
        own = self.market.code
        if own not in spec.allowed_markets:
            raise ValueError(f"{ws_title}/{spec.role}: allowed_markets {spec.allowed_markets} must include {own}")
        foreign_ok = own == "CA" and ws_title in C.CA_SHEETS_ALLOWING_US
        for col in spec.columns:
            eff = col.market or own
            if eff != own:
                if not foreign_ok or eff not in spec.allowed_markets:
                    raise ValueError(f"{self.name}!{ws_title}: column {col.header!r} uses market {eff}; "
                                     f"not allowed on this sheet")
                if col.kind in ("money", "money2") and "(USD)" not in col.header:
                    raise ValueError(f"{ws_title}: foreign-currency column {col.header!r} must say '(USD)'")
                if col.kind == "link" and "US" not in col.header:
                    raise ValueError(f"{ws_title}: amazon.com link column {col.header!r} must say 'US'")
        if "market" in spec.rows.columns:
            foreign = {m for m in spec.rows["market"].dropna().unique() if m != own}
            if foreign and (not foreign_ok or not foreign <= set(spec.allowed_markets)):
                raise ValueError(f"{self.name}!{ws_title}/{spec.role}: rows from market(s) {sorted(foreign)} not allowed here")

    def table(self, ws: Worksheet, spec: TableSpec, top_row: int, *, first_col: int = 1, freeze: bool = True) -> TableRange:
        self._check_markets(ws.title, spec)
        tr = write_table(ws, spec, top_row, self.market, first_col=first_col, freeze=freeze)
        self.tables.append(tr)
        return tr

    def kpi(self, ws: Worksheet, top_row: int, items, **kw) -> TableRange:
        tr = write_kpi_block(ws, top_row, items, self.market, **kw)
        self.tables.append(tr)
        return tr

    def line(self, ws: Worksheet, row: int, text: str, **kw) -> TableRange:
        tr = write_line(ws, row, text, self.market, **kw)
        self.tables.append(tr)
        return tr

    def text(self, ws: Worksheet, title: str, paragraphs, *, role: str) -> TableRange:
        tr = write_text_sheet(ws, title, paragraphs, role=role)
        self.tables.append(tr)
        return tr

    def metadata(self, ws: Worksheet, entries: dict) -> TableRange:
        tr = write_metadata_sheet(ws, entries, self.market)
        self.tables.append(tr)
        return tr

    # charts (no-op when the table has no data rows: nothing to plot, nothing registered)
    def bar(self, ws, tr, cat, val, title, **kw):
        return add_bar_chart(ws, tr, cat, val, None, title, **kw) if tr.n_rows > 0 else None

    def pie(self, ws, tr, cat, val, title, **kw):
        return add_pie_chart(ws, tr, cat, val, None, title, **kw) if tr.n_rows > 0 else None

    def scatter(self, ws, tr, x, y, title):
        return add_scatter_chart(ws, tr, x, y, None, title) if tr.n_rows > 0 else None

    def save(self, path: Path) -> Path:
        if self._saved:
            raise RuntimeError(f"{self.name} was already saved (single-pass writer)")
        path = Path(path)
        if path.name != self.name:
            raise ValueError(f"saving {self.name} under a different name {path.name}")
        self.wb.save(path)
        self._saved = True
        return path


# --------------------------------------------------------------------------------------
# Registry
# --------------------------------------------------------------------------------------
class TableRegistry:
    """Collects every table/chart of the books built in one run -> runs/<m>/table_registry_<market>_<m>.json.

    The file is shared by every builder of a market (CA code-reader + CA gauge): entries of the workbooks written in this
    run replace their previous entries; entries of other workbooks are kept.
    """

    def __init__(self, market: str, month: str):
        if market not in C.MARKETS:
            raise ValueError(market)
        if not C.MONTH_RE.match(month):
            raise ValueError(month)
        self.market = market
        self.month = month
        self.books: list[Book] = []

    def add_book(self, book: Book) -> None:
        if book.market.code != self.market:
            raise ValueError(f"{book.name}: market {book.market.code} != registry market {self.market}")
        self.books.append(book)

    def path(self, runs_dir: Path) -> Path:
        return C.run_file(Path(runs_dir), self.month, "table_registry", self.market, "json")

    def write(self, runs_dir: Path) -> Path:
        path = self.path(runs_dir)
        names = {b.name for b in self.books}
        tables: list[dict] = []
        workbooks: dict[str, dict] = {}
        if path.exists():
            prev = json.loads(path.read_text(encoding="utf-8"))
            if (prev.get("market"), prev.get("month")) != (self.market, self.month):
                raise ValueError(f"{path}: registry is for {prev.get('market')}/{prev.get('month')}")
            tables = [t for t in prev["tables"] if t["workbook"] not in names]
            workbooks = {k: v for k, v in prev["workbooks"].items() if k not in names}
        now = datetime.now().isoformat(timespec="seconds")
        for b in self.books:
            workbooks[b.name] = {"sheets": b.sheetnames, "brand_sheet_map": dict(b.brand_sheet_map), "generated_at": now}
            tables.extend(tr.to_registry(b.name, b.brand_sheet_map) for tr in b.tables)
        out = {"market": self.market, "month": self.month, "generated_at": now, "pipeline_version": pipeline_version(),
               "workbooks": dict(sorted(workbooks.items())), "tables": tables}
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp = path.with_suffix(".json.tmp")
        tmp.write_text(json.dumps(out, indent=1, ensure_ascii=False), encoding="utf-8")
        tmp.replace(path)
        return path


# --------------------------------------------------------------------------------------
# Manifest / provenance / overwrite guard
# --------------------------------------------------------------------------------------
def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def input_hashes(paths: Iterable[Path]) -> dict[str, str]:
    out: dict[str, str] = {}
    for p in paths:
        p = Path(p).resolve()
        if not p.is_file():
            raise FileNotFoundError(f"manifest input not found: {p}")
        out[str(p)] = sha256_file(p)
    return out


_GIT_SHA: str | None = None


def pipeline_version() -> str:
    """git sha of the checkout running the code, '-dirty' when ca_market_reports/ has uncommitted changes."""
    global _GIT_SHA
    if _GIT_SHA is None:
        try:
            sha = subprocess.run(["git", "-C", str(C.REPO_ROOT), "rev-parse", "--short=12", "HEAD"], capture_output=True,
                                 text=True, check=True).stdout.strip()
            dirty = subprocess.run(["git", "-C", str(C.REPO_ROOT), "status", "--porcelain", "--", "ca_market_reports"],
                                   capture_output=True, text=True, check=True).stdout.strip()
        except (OSError, subprocess.CalledProcessError) as exc:
            raise RuntimeError(f"cannot determine the pipeline git sha: {exc}") from exc
        if not sha:
            raise RuntimeError("git rev-parse returned an empty sha")
        _GIT_SHA = sha + ("-dirty" if dirty else "")
    return _GIT_SHA


def read_manifest(out_dir: Path, month: str) -> dict | None:
    p = Path(out_dir) / C.manifest_name(month)
    if not p.exists():
        return None
    m = json.loads(p.read_text(encoding="utf-8"))
    if m.get("month") != month:
        raise ValueError(f"{p}: manifest month {m.get('month')!r} != {month}")
    return m


def write_manifest(out_dir: Path, month: str, outputs: list[Path], inputs: dict[str, str]) -> Path:
    """manifest_<m>.json in out_dir: sha256 of each output (files directly in out_dir; never the manifest itself) +
    input hashes + pipeline git sha + generated_at. Entries of outputs written by other builders are kept."""
    out_dir = Path(out_dir)
    path = out_dir / C.manifest_name(month)
    prev = read_manifest(out_dir, month)
    now = datetime.now().isoformat(timespec="seconds")
    outs: dict[str, dict] = {}
    if prev:
        for name, entry in prev.get("outputs", {}).items():
            if (out_dir / name).exists():
                outs[name] = entry
            else:
                print(f"NOTE: manifest entry dropped (file no longer in {out_dir}): {name}")
    for p in outputs:
        p = Path(p)
        if p.resolve().parent != out_dir.resolve():
            raise ValueError(f"manifest outputs must sit directly in {out_dir} (non-recursive): {p}")
        if p.name == path.name:
            continue
        outs[p.name] = {"sha256": sha256_file(p), "bytes": p.stat().st_size, "generated_at": now}
    ins = dict(prev.get("inputs", {})) if prev else {}
    ins.update(inputs)
    manifest = {"month": month, "generated_at": now, "pipeline_git_sha": pipeline_version(),
                "outputs": dict(sorted(outs.items())), "inputs": dict(sorted(ins.items()))}
    tmp = path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(manifest, indent=1), encoding="utf-8")
    tmp.replace(path)
    return path


def safe_output_path(out_dir: Path, name: str, overwrite: bool, manifest: dict | None) -> Path:
    """Refuse an existing output unless overwrite; when overwriting, move the old file to
    <out_dir>/_backup/<stem>.<YYYYMMDD-HHMMSS>.xlsx first and print HAND-EDITED when it differs from the previous manifest."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / name
    lock = out_dir / f"~${name}"
    if lock.exists():
        raise RuntimeError(f"{path} is open in Excel (lock file {lock.name}); close it first")
    if not path.exists():
        return path
    if not overwrite:
        raise FileExistsError(f"{path} exists; pass --overwrite to replace it (the old file is backed up)")
    old_sha = sha256_file(path)
    prev = (manifest or {}).get("outputs", {}).get(name)
    if prev is None:
        print(f"NOTE: no previous manifest entry for {path}; backing up without a hand-edit check")
    elif prev.get("sha256") != old_sha:
        print(f"HAND-EDITED: {path}")
    backup_dir = out_dir / C.BACKUP_SUBDIR
    backup_dir.mkdir(exist_ok=True)
    stem, suffix = Path(name).stem, Path(name).suffix
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    dest = backup_dir / f"{stem}.{ts}{suffix}"
    n = 2
    while dest.exists():
        dest = backup_dir / f"{stem}.{ts}-{n}{suffix}"
        n += 1
    shutil.move(str(path), str(dest))
    print(f"backup: {path.name} -> {dest}")
    return path


def dated_copy_name(name: str, day: date | None = None) -> str:
    """<stem>_<YYYYMMDD>.xlsx (generation date) for --dated-copy."""
    day = day or date.today()
    p = Path(name)
    return f"{p.stem}_{day.strftime('%Y%m%d')}{p.suffix}"


def refuse_new_product_dir(out_dir: Path) -> None:
    """Dev/fixture runs must never write under NewProductCategory/."""
    o = Path(out_dir).resolve()
    npd = C.NEW_PRODUCT_DIR.resolve()
    if o == npd or npd in o.parents:
        raise ValueError(f"dev (--from-normalized) runs must not write under {C.NEW_PRODUCT_DIR}: {out_dir}")


# --------------------------------------------------------------------------------------
# Normalized frames (dev/fixture mode) and frame coercion
# --------------------------------------------------------------------------------------
NUMERIC_COLUMNS: tuple[str, ...] = ("bsr", "subcategory_bsr", "list_price", "units_month", "revenue_month", "price", "review_count",
                                    "rating", "listing_age_months", "variation_count", "last_year_units", "yoy_units_pct",
                                    "sales_trend_90d_pct", "price_trend_90d_pct", "type_confidence", "gauge_confidence")
STRICT_BOOL_COLUMNS: tuple[str, ...] = ("gauge_in_scope", "gauge_device_scope")   # blank is an error
BLANK_FALSE_BOOL_COLUMNS: tuple[str, ...] = ("frequently_returned", "type_conflict", "borderline", "alarms", "multi_gauge",
                                             "gesture_control", "kmh_mph", "lordco_type_unit")
_TRUE = {"y", "yes", "true", "1"}
_FALSE = {"n", "no", "false", "0"}
INPUT_MODE_NORMALIZED = "normalized frame (dev/fixture: --from-normalized; loader and gauge classifier not run)"
INPUT_MODE_LOADER = "raw Helium 10 exports via ca_load.load_month"


def parse_bool_value(v: Any, *, column: str, blank_is_false: bool) -> bool:
    if isinstance(v, bool) or type(v).__name__ == "bool_":
        return bool(v)
    if _is_blank(v) or (isinstance(v, str) and v.strip() == ""):
        if blank_is_false:
            return False
        raise ValueError(f"column {column!r}: blank boolean")
    s = str(v).strip().lower()
    if s in _TRUE:
        return True
    if s in _FALSE:
        return False
    raise ValueError(f"column {column!r}: cannot parse boolean {v!r}")


def coerce_frame(df: pd.DataFrame) -> pd.DataFrame:
    """Typed copy: numeric columns -> float (blank -> NaN, garbage -> error), boolean columns -> bool. Columns already typed
    by the loader are kept as they are; every column of the input is preserved."""
    out = df.copy()
    for c in NUMERIC_COLUMNS:
        if c not in out.columns:
            continue
        s = out[c]
        if s.dtype == object:
            s = s.map(lambda v: float("nan") if _is_blank(v) or (isinstance(v, str) and v.strip() == "") else v)
            try:
                s = pd.to_numeric(s, errors="raise")
            except (ValueError, TypeError) as exc:
                raise ValueError(f"column {c!r}: non-numeric value ({exc})") from exc
        out[c] = s.astype(float)
        bad = out[c][~out[c].isna() & ~out[c].map(lambda x: math.isfinite(x))]
        if len(bad):
            raise ValueError(f"column {c!r}: non-finite values at rows {list(bad.index)[:5]}")
    # ASSEMBLY (orchestrator): frames that have not been through the gauge classifier carry gauge_class == "" and
    # gauge_in_scope/gauge_device_scope == <NA> (loader contract). Those rows are out of gauge scope by definition, so
    # the scope flags become False there; a CLASSIFIED row with a blank scope flag is still an error.
    if "gauge_class" in out.columns:
        unclassified = out["gauge_class"].map(_is_blank) | (out["gauge_class"].astype(str).str.strip() == "")
        for c in STRICT_BOOL_COLUMNS:
            if c in out.columns:
                out.loc[unclassified, c] = False
    for c in STRICT_BOOL_COLUMNS + BLANK_FALSE_BOOL_COLUMNS:
        if c not in out.columns:
            continue
        blank_ok = c in BLANK_FALSE_BOOL_COLUMNS
        out[c] = [parse_bool_value(v, column=c, blank_is_false=blank_ok) for v in out[c]]
        out[c] = out[c].astype(bool)
    for c in ("asin", "title", "brand_key", "brand_display", "type", "type_source", "price_tier", "gauge_class", "source_set",
              "market", "currency", "source_file", "seller", "fulfillment", "gauge_rule_id"):
        if c in out.columns:
            out[c] = out[c].map(lambda v: "" if _is_blank(v) else str(v))
    return out


def read_normalized_csv(path: Path) -> pd.DataFrame:
    """Read a normalized-row CSV (BASE_ROW_COLUMNS [+ enrichment]) as strings, then coerce types."""
    df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8")
    return coerce_frame(df)


def require_columns(df: pd.DataFrame, cols: Iterable[str], what: str) -> None:
    missing = [c for c in cols if c not in df.columns]
    if missing:
        raise KeyError(f"{what}: missing columns {missing}")


def dataset_from_normalized(df: pd.DataFrame, market: str, month: str) -> C.CaDataset:
    """CaDataset from an already-normalized (and, for the gauge workbook, already-classified) frame.

    code_reader = rows with source_set in {code_reader, both}; gauge_set = rows with source_set in {gauge, both}.
    """
    if market not in C.MARKETS:
        raise ValueError(f"unknown market {market!r}")
    if not C.MONTH_RE.match(str(month)):
        raise ValueError(f"invalid month {month!r}")
    require_columns(df, C.BASE_ROW_COLUMNS, "normalized frame")
    df = coerce_frame(df)
    mk = set(df["market"])
    if mk != {market}:
        raise ValueError(f"normalized frame markets {sorted(mk)} != {market}")
    ccy = set(df["currency"])
    if ccy != {C.MARKETS[market].currency}:
        raise ValueError(f"normalized frame currencies {sorted(ccy)} != {C.MARKETS[market].currency}")
    bad = [a for a in df["asin"] if not C.ASIN_RE.match(a)]
    if bad:
        raise ValueError(f"invalid ASINs {bad[:5]}")
    dup = df["asin"][df["asin"].duplicated()].tolist()
    if dup:
        raise ValueError(f"duplicate ASINs in a normalized frame (rows must be deduped): {dup[:5]}")
    unknown = set(df["source_set"]) - set(C.SOURCE_SETS)
    if unknown:
        raise ValueError(f"unknown source_set values {sorted(unknown)}")
    dates = sorted({date.fromisoformat(d) for d in df["export_date"] if str(d).strip()})
    if not dates:
        raise ValueError("no export_date in the normalized frame (never a fake date)")
    cr = df[df["source_set"].isin(["code_reader", "both"])].reset_index(drop=True)
    gs = df[df["source_set"].isin(["gauge", "both"])].reset_index(drop=True)
    raw_files = [(str(k), int(v)) for k, v in df["source_file"].value_counts(sort=False).items()]
    return C.CaDataset(market=market, month=str(month), code_reader=cr, gauge_set=gs, export_dates=(dates[0], dates[-1]),
                       raw_files=raw_files, audits={"dedupe_audit": pd.DataFrame(columns=list(C.DEDUPE_AUDIT_COLUMNS)),
                                                    "input_mode": INPUT_MODE_NORMALIZED})


def input_mode(ds: C.CaDataset) -> str:
    return ds.audits.get("input_mode", INPUT_MODE_LOADER)


# --------------------------------------------------------------------------------------
# Shared aggregations (static values; Avg Rating = units-weighted mean over rating > 0, blank when none)
# --------------------------------------------------------------------------------------
def weighted_rating(df: pd.DataFrame) -> float:
    r = df["rating"]
    u = df["units_month"].fillna(0.0)
    m = r.notna() & (r > 0)
    w = float(u[m].sum())
    if w <= 0:
        return float("nan")
    return float((r[m] * u[m]).sum() / w)


def safe_div(a: float, b: float) -> float:
    return float(a) / float(b) if b and not math.isnan(b) and b != 0 else float("nan")


def check_sales_columns(df: pd.DataFrame, what: str) -> None:
    require_columns(df, ("asin", "revenue_month", "units_month", "price", "brand_key", "brand_display"), what)
    for c in ("revenue_month", "units_month"):
        bad = df.loc[df[c].isna(), "asin"].tolist()
        if bad:
            raise ValueError(f"{what}: {c} is missing for ASINs {bad[:10]} (must be numeric; '-' -> 0 belongs in the loader)")
        neg = df.loc[df[c] < 0, "asin"].tolist()
        if neg:
            raise ValueError(f"{what}: negative {c} for ASINs {neg[:10]}")
    if (df["brand_key"] == "").any():
        raise ValueError(f"{what}: empty brand_key for ASINs {df.loc[df['brand_key'] == '', 'asin'].tolist()[:10]}")
    disp = df.groupby("brand_key")["brand_display"].nunique()
    if (disp > 1).any():
        raise ValueError(f"{what}: brand_key with several displays: {disp[disp > 1].index.tolist()[:10]}")
    if df["asin"].duplicated().any():
        raise ValueError(f"{what}: duplicate ASINs {df.loc[df['asin'].duplicated(), 'asin'].tolist()[:10]}")


def _agg(sub: pd.DataFrame, total_rev: float) -> dict:
    rev = float(sub["revenue_month"].sum())
    units = float(sub["units_month"].sum())
    return {"listings": int(sub["asin"].nunique()), "rev": rev, "units": units, "share": safe_div(rev, total_rev),
            "ppu": safe_div(rev, units), "reviews": float(sub["review_count"].sum(skipna=True)) if "review_count" in sub else float("nan"),
            "rating": weighted_rating(sub)}


def brand_summary(df: pd.DataFrame, *, top_n: int, total_rev: float | None = None,
                  noun: str = "brands") -> tuple[pd.DataFrame, dict | None, dict]:
    """Brand rows (fields brand, brand_key, listings, rev, units, share, ppu, reviews, rating) sorted by revenue desc, the
    residual 'Other brands (n)' when truncated, and the Total over the FULL frame."""
    total_rev = float(df["revenue_month"].sum()) if total_rev is None else total_rev
    rows = []
    for key, sub in df.groupby("brand_key", sort=False):
        d = _agg(sub, total_rev)
        d.update(brand_key=key, brand=str(sub["brand_display"].iloc[0]))
        rows.append(d)
    cols = ["brand", "brand_key", "listings", "rev", "units", "share", "ppu", "reviews", "rating"]
    full = pd.DataFrame(rows, columns=cols)
    if len(full):
        full = full.sort_values(["rev", "units", "brand"], ascending=[False, False, True], kind="mergesort").reset_index(drop=True)
    shown = full.head(top_n).reset_index(drop=True)
    rest_keys = set(full["brand_key"].iloc[top_n:])
    residual = None
    if rest_keys:
        residual = _agg(df[df["brand_key"].isin(rest_keys)], total_rev)
        residual["brand"] = C.RESIDUAL_ROW_LABEL.format(noun=noun, n=len(rest_keys))
    total = _agg(df, total_rev)
    total["brand"] = C.TOTAL_ROW_LABEL
    return shown, residual, total


def rank_listings(df: pd.DataFrame, by: str) -> pd.DataFrame:
    if by == "revenue":
        keys, asc = ["revenue_month", "units_month", "asin"], [False, False, True]
    elif by == "units":
        keys, asc = ["units_month", "revenue_month", "asin"], [False, False, True]
    elif by == "price":
        keys, asc = ["price", "asin"], [True, True]
    else:
        raise ValueError(by)
    out = df.sort_values(keys, ascending=asc, kind="mergesort").reset_index(drop=True)
    out["_rank"] = range(1, len(out) + 1)
    return out


def listing_totals(sub: pd.DataFrame, label_field: str, label: str) -> dict:
    return {label_field: label, "revenue_month": float(sub["revenue_month"].sum()), "units_month": float(sub["units_month"].sum()),
            "review_count": float(sub["review_count"].sum(skipna=True)), "rating": weighted_rating(sub)}


def top_listings(df: pd.DataFrame, *, by: str, n: int, label_field: str = "_rank") -> tuple[pd.DataFrame, dict | None, dict]:
    ranked = rank_listings(df, by)
    shown = ranked.head(n).reset_index(drop=True)
    rest = ranked.iloc[n:]
    residual = listing_totals(rest, label_field, C.RESIDUAL_ROW_LABEL.format(noun="listings", n=len(rest))) if len(rest) else None
    total = listing_totals(df, label_field, C.TOTAL_ROW_LABEL)
    return shown, residual, total


def add_canonical_url(df: pd.DataFrame, market: Market) -> pd.DataFrame:
    """_url = https://<row market domain>/dp/<asin> (same destination as the HYPERLINK)."""
    out = df.copy()
    if "market" in out.columns:
        out["_url"] = [C.MARKETS[m or market.code].url(a) for a, m in zip(out["asin"], out["market"])]
    else:
        out["_url"] = [market.url(a) for a in out["asin"]]
    return out


def tier_of(price: float, tiers: tuple[tuple[str, float, float], ...]) -> str:
    """Half-open [lo, hi) tier label; NaN price is an error."""
    if price is None or (isinstance(price, float) and math.isnan(price)):
        raise ValueError("price tier of a missing price")
    for label, lo, hi in tiers:
        if lo <= price < hi:
            return label
    raise ValueError(f"price {price!r} outside every tier")


def fmt_money(x: float, market: Market, dp: int = 0) -> str:
    sym = "CA$" if market.code == "CA" else "$"
    return f"{sym}{x:,.{dp}f}"


def fit(d: dict | None, columns: Sequence[ColumnSpec]) -> dict | None:
    """Keep only the keys of a total/residual dict that the table shows (deliberate projection at the call site)."""
    if d is None:
        return None
    fields = {c.field for c in columns}
    return {k: v for k, v in d.items() if k in fields}


# --------------------------------------------------------------------------------------
# Shared column sets (frozen headers from ca_common, {ccy} formatted)
# --------------------------------------------------------------------------------------
def summary_columns(ccy: str) -> list[ColumnSpec]:
    h = [s.format(ccy=ccy) for s in C.SUMMARY_COLUMNS]
    return [ColumnSpec(h[0], "brand", "text", 28), ColumnSpec(h[1], "listings", "int", 11), ColumnSpec(h[2], "rev", "money", 15),
            ColumnSpec(h[3], "units", "int", 12), ColumnSpec(h[4], "share", "pct", 14), ColumnSpec(h[5], "ppu", "money2", 13),
            ColumnSpec(h[6], "reviews", "int", 12), ColumnSpec(h[7], "rating", "rating", 10)]


def top50_columns(ccy: str) -> list[ColumnSpec]:
    h = [s.format(ccy=ccy) for s in C.TOP50_COLUMNS]
    spec = [("_rank", "int", 9), ("asin", "text", 13), ("title", "text", 60), ("brand_display", "text", 18), ("type", "text", 13),
            ("price", "money2", 12), ("revenue_month", "money", 15), ("units_month", "int", 12), ("review_count", "int", 11),
            ("rating", "rating", 9), ("listing_age_months", "int", 11), ("last_year_units", "int", 13),
            ("yoy_units_pct", "pct", 12), ("sales_trend_90d_pct", "pct", 13), ("_url", "text", 34), ("asin", "link", 24)]
    return [ColumnSpec(hh, f, k, w) for hh, (f, k, w) in zip(h, spec, strict=True)]


def brand_tab_columns(ccy: str, *, type_header: str = "Type", type_field: str = "type") -> list[ColumnSpec]:
    h = [s.format(ccy=ccy) for s in C.BRAND_TAB_COLUMNS]
    h[2] = type_header
    spec = [("title", "text", 60), ("asin", "text", 13), (type_field, "text", 16), ("price", "money2", 12),
            ("revenue_month", "money", 15), ("units_month", "int", 12), ("review_count", "int", 11), ("rating", "rating", 9),
            ("listing_age_months", "int", 11), ("yoy_units_pct", "pct", 12), ("_url", "text", 34), ("asin", "link", 24)]
    return [ColumnSpec(hh, f, k, w) for hh, (f, k, w) in zip(h, spec, strict=True)]


def brand_kpi_items(rows: pd.DataFrame, market_rev: float, ccy: str) -> list[tuple[str, Any, str]]:
    """Rows 3-6 of a brand tab: Monthly Rev, Units, # Listings, Rev share."""
    rev = float(rows["revenue_month"].sum())
    return [(f"Monthly Rev ({ccy})", rev, "money"), ("Monthly Units", float(rows["units_month"].sum()), "int"),
            ("# of Listings", int(rows["asin"].nunique()), "int"), ("Monthly Rev Market Share %", safe_div(rev, market_rev), "pct")]


def write_brand_tab(book: Book, sheet_name: str, rows: pd.DataFrame, *, columns: list[ColumnSpec], title: str, subtitle: str,
                    kpi_items: list[tuple[str, Any, str]], dataset_filter: str, rev_role: str = "brand_tab_revenue",
                    units_role: str = "brand_tab_units") -> tuple[Worksheet, TableRange, TableRange]:
    """Brand tab layout (every brand tab): row 1 title, row 2 subtitle, rows 3-6 KPI block, 'Rank by Revenue' title at row 7,
    header at row 8, then 3 blank rows and 'Rank by Units'. Total rows = the brand's full row set (no truncation)."""
    if len(kpi_items) != C.BRAND_TAB_RESERVED_ROWS - 2:
        raise ValueError(f"brand tab KPI block needs exactly {C.BRAND_TAB_RESERVED_ROWS - 2} items")
    ws = book.sheet(sheet_name)
    for j, col in enumerate(columns):
        _set_width(ws, 1 + j, col.width)
    write_sheet_header(ws, title, subtitle, len(columns))
    book.kpi(ws, 3, kpi_items, header=False, dataset_filter=dataset_filter)
    code = book.market.code
    total = fit(listing_totals(rows, columns[0].field, C.TOTAL_ROW_LABEL), columns)
    tr1 = book.table(ws, TableSpec(sheet_name, rev_role, "Rank by Revenue", columns, rank_listings(rows, "revenue"), total, None,
                                   (code,), dataset_filter=dataset_filter), C.BRAND_TAB_RESERVED_ROWS + 1)
    if tr1.header_row != C.BRAND_TAB_RESERVED_ROWS + 2:
        raise AssertionError("brand tab header must be at row 8")
    tr2 = book.table(ws, TableSpec(sheet_name, units_role, "Rank by Units", columns, rank_listings(rows, "units"), total, None,
                                   (code,), dataset_filter=dataset_filter), table_end(tr1) + 4)
    return ws, tr1, tr2


# --------------------------------------------------------------------------------------
# Shared metadata values
# --------------------------------------------------------------------------------------
CURRENCY_TEXT = {"CA": "CAD (Helium 10 Black Box estimates, uncalibrated)",
                 "US": "USD (Helium 10 Black Box estimates, raw; no actuals overlay)"}


def dedupe_summary(audit: pd.DataFrame | None, mode: str) -> str:
    if audit is None:
        raise KeyError("dataset audits carry no 'dedupe_audit' frame")
    require_columns(audit, C.DEDUPE_AUDIT_COLUMNS, "dedupe_audit")
    if audit.empty:
        if mode == INPUT_MODE_NORMALIZED:
            return "no dedupe audit: built from a normalized frame (loader not run)"
        return "dedupe audit is empty (no rows recorded by the loader)"
    n_rows = pd.to_numeric(audit["n_rows"], errors="raise")
    dup = audit[n_rows > 1]
    rules = dup["winning_rule"].value_counts().sort_index()
    rule_txt = ", ".join(f"{k} {v}" for k, v in rules.items()) or "none"
    flags = int((audit["discrepancy_flag"].astype(str) == "Y").sum())
    return (f"{len(audit)} audit rows; {len(dup)} ASIN groups had more than one source row (winning rules: {rule_txt}); "
            f"discrepancy flags: {flags}")


def base_metadata(ds: C.CaDataset, market: Market) -> dict[str, str]:
    d0, d1 = ds.export_dates
    raw = "; ".join(f"{n}: {r}" for n, r in ds.raw_files) or "none recorded"
    return {
        "Marketplace": f"{market.domain} ({market.code})",
        "Currency": CURRENCY_TEXT[market.code],
        "Report month": f"{ds.month} ({month_label(ds.month)})",
        "Export dates": f"{d0.isoformat()} – {d1.isoformat()}",
        "Raw files": raw,
        "Pipeline version": f"git {pipeline_version()}",
        "Generated at": datetime.now().isoformat(timespec="seconds"),
        "Input mode": input_mode(ds),
    }


def type_source_base(s: pd.Series) -> pd.Series:
    base = s.astype(str).str.split(":", n=1).str[0]
    unknown = sorted(set(base) - set(C.TYPE_SOURCES))
    if unknown:
        raise ValueError(f"unknown type_source values {unknown}; expected {C.TYPE_SOURCES} (keyword hits as 'keyword:<rule>')")
    return base


def type_coverage_text(df: pd.DataFrame) -> str:
    base = type_source_base(df["type_source"])
    n, rev = len(df), float(df["revenue_month"].sum())
    parts = []
    for src in C.TYPE_SOURCES:
        m = base == src
        if m.any():
            parts.append(f"{src} {m.sum() / n:.1%} rows / {safe_div(df.loc[m, 'revenue_month'].sum(), rev):.1%} revenue")
    return "; ".join(parts) if parts else "no rows"
