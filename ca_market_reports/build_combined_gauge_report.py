"""Combined CA + US OBD gauge competitor workbook (single month; the two markets side by side, native currencies, no FX).

  CA_US_OBD_Gauge_Competitor_Report_<m>.xlsx   (ca_common.combined_gauge_report_name; lands in the CA-OBD-GAUGE outputs dir)
    Read Me, Summary, Top 50 CA, Top 50 US, Innova, brand tabs,
    Price Ladder (Model A), Feature Matrix (Model A), App-Gauge Proxy (Model B)      [the CA analyses, as in the CA workbook]
    US vs CA Same-ASIN, All Products, Dedupe & Classification Audit, Excluded, Source & Method, Metadata

Same numbers as the two single-market gauge workbooks, by construction:
  * each market goes through build_gauge_report's own path (code_reader_totals, gauge_union_with_flags -> the same loader,
    classifier, maps and frozen decisions of runs/<m>, replayed with unchanged freeze semantics; modelb_universe and
    read_app_feature_matrix for CA);
  * every table this workbook repeats from a single-market workbook (Top 50, Innova hardware, All Products, Excluded, Dedupe
    audit, Read Me taxonomy, Source & Method text, Metadata values, brand-tab qualification) is produced by build_gauge_report's
    own sheet writer, run into a throw-away recorder book, and re-laid here (never re-implemented); the Model A/B and
    Same-ASIN writers run directly into this workbook.
Summary tables are new: ONE table per topic with a CA block and a US block side by side (market header fill + data tint
from ca_common.FILL_MARKET_*), money formats per block (ColumnSpec market override), shares within each market, no
cross-currency ratio. Every number is a static value; the HYPERLINK is the only formula. Run:
  ca_market_reports/run.sh ca_market_reports/build_combined_gauge_report.py --month 202609 [--overwrite]
  dev: ... --from-normalized <CA csv> --us-from-normalized <US csv> --out-dir tmp/ca_scratch/cx --runs-dir tmp/ca_scratch/cx/runs
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Sequence

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from ca_market_reports import build_gauge_report as G
from ca_market_reports import ca_common as C
from ca_market_reports import ca_xlsx_style as X
from ca_market_reports.ca_xlsx_style import ColumnSpec as Col, TableRange, TableSpec

REGISTRY_SCOPE = "CAUS"                     # runs/<m>/table_registry_CAUS_<m>.json
MARKET_CODES: tuple[str, str] = C.COMBINED_MARKETS
if MARKET_CODES != ("CA", "US"):
    raise AssertionError(f"combined workbook expects markets ('CA', 'US'), ca_common has {MARKET_CODES}")
if C.GAUGE_ADJACENT_CLASSES != ("gps_hud",):
    raise AssertionError(f"the Summary GPS-only HUD row assumes GAUGE_ADJACENT_CLASSES == ('gps_hud',), got {C.GAUGE_ADJACENT_CLASSES}")
HEADER_FILL = {m: PatternFill("solid", fgColor=C.FILL_MARKET_HEADER[m]) for m in MARKET_CODES}
DATA_FILL = {m: PatternFill("solid", fgColor=C.FILL_MARKET_DATA[m]) for m in MARKET_CODES}

KEY_FIGURE_LABELS: tuple[str, ...] = ("Core device revenue", "Core device units", "# core device ASINs",
                                      "# core device ASINs with sales > 0", "Incl. borderline revenue", "Accessories revenue",
                                      "Adjacent GPS-only HUD revenue")
KEY_FIGURES_NOTE = ("Core devices = device scope (tuner with gauge display, truck gauge monitor, OBD+GPS HUD, OBD HUD, gauge "
                    "display), not borderline. 'Incl. borderline' = device scope. Accessories and GPS-only HUDs are shown, never "
                    "counted in device totals. Money rows are in each market's own currency (CA in CAD, US in USD).")
_SHARE_B_SENTENCE = "(b) = all core devices ÷ (full code-reader export + core devices found only in the gauge export)."
if _SHARE_B_SENTENCE not in G.SHARE_NOTE or G.SHARE_ROW_LABELS[2] != C.COMBINED_SHARE_ROW_LABEL:
    raise AssertionError("build_gauge_report share definition (b) drifted from the combined share table")
COMBINED_SHARE_NOTE = ("Core devices = device scope, not borderline. " + _SHARE_B_SENTENCE +
                       " Shares of revenue and of units are computed within each market (no cross-currency ratio).")
GPS_ROW_LABEL = C.GAUGE_SUBTYPE_LABELS["gps_hud"]   # shared label; the note row states that the row is excluded from totals
FUEL_SUBTOTAL_LABEL = "{fuel} — subtotal"
NO_LISTINGS_LABEL = "(no listings in this market)"
INNOVA_CA_COUNT_CELL = "B3"      # Innova tab: A3/B3 = CA device-count line + number; A4/B4 = US
INNOVA_US_COUNT_CELL = "B4"
INNOVA_US_NOTE = ("US code-reader Types are not assigned (out of scope), so there is no US Innova hardware list: the "
                  "app-capable hardware below is the CA Model B universe (CA code-reader listings typed Dongle).")
BRAND_KPI_LABELS: tuple[str, ...] = ("Monthly Rev (CAD | USD)", "Monthly Units", "# of Listings", "Rev share within market")
BRAND_NOTE = ("Core devices (device scope, not borderline), each market from its own rows. Brands shown = the top {n} by the "
              "larger of their CA (CAD) and US (USD) revenue — a nominal selection rule, no amount is converted; rows are "
              "ordered by CA revenue, then US revenue. Shares are within each market and sum to 1 with the 'Other brands' "
              "row; Total = the market's full core totals. Avg rating = units-weighted mean over ratings > 0.")
SUBTYPE_NOTE = ("Core devices; Rev share within each market. The GPS-only HUD row is adjacent (not an OBD gauge) and is "
                "excluded from the Total and from the shares.")
TIER_NOTE = ("Core devices. Tiers are half-open [low, high) on the listing price, nominal break points in the market's own "
             "currency ({ccy}). Each tier shows revenue and units; 'All tiers' = the sub-type total.")
COMBINED_LAYOUT_PARAGRAPHS: tuple[str, ...] = (
    "# Combined CA + US layout",
    "Side by side: every Summary table is ONE table with a CA block (purple header, CAD) and a US block (blue header, USD) in "
    "adjacent columns. Each block is computed from its own market's classified rows only; no figure mixes the two markets.",
    "No FX: CAD and USD amounts are never converted, added or divided into one figure (no cross-currency ratio); money "
    "headers carry (CAD) or (USD) and each block uses its own currency format. Shares (revenue shares, the gauge share of "
    "the code-reader market) are computed within each market.",
    "Brand summary: rows are the union of the core-device brands of both markets. The brands shown are the top "
    f"{C.SUMMARY_TOP_BRANDS} by the larger of their CA and US revenue (a nominal selection rule only), ordered by CA "
    "revenue, then US revenue; 'Other brands (n)' holds each market's residual, so each market's shares sum to 1 and each "
    "market's Total equals its full core totals.",
    "Sub-type mix: the GPS-only HUD row is adjacent and excluded from the Total. Price tier × sub-type: one table per "
    "market with revenue and units per tier. Fuel split: each '<fuel> — subtotal' row sums the sub-type rows below it; "
    "the Total equals the sum of the subtotals.",
    "Brand tabs: a brand gets a tab when it meets the single-market rule (core device revenue of at least 1,000 in the "
    f"market currency, or at least {C.GAUGE_BRAND_TAB_MIN_ASINS} device listings) in either market. Each tab holds four "
    "ranking tables (CA — Rank by Revenue, CA — Rank by Units, US — Rank by Revenue, US — Rank by Units) with Total rows; "
    "a market without listings shows '(no listings in this market)' and a zero Total.",
    "Top 50 CA / Top 50 US, All Products, Dedupe & Classification Audit and Excluded repeat the single-market tables, one "
    "per market (the audit sheet keeps the dedupe audit; per-ASIN classification decisions are the matching columns of "
    "All Products). The CA workbook's US Benchmark sheet is replaced by the side-by-side Summary.",
    "Consistency: both markets go through the single-market gauge pipeline unchanged (same loader, classifier, maps and "
    "frozen decisions in runs/<m>), so every figure equals the matching cell of the CA and US OBD gauge workbooks.",
)
SCOPE_NOTE_COMBINED = ("Side-by-side CA + US layout: no FX conversion, shares within each market. "
                       "US figures are raw Helium 10 estimates without the actuals overlay.")


# --------------------------------------------------------------------------------------
# Books
# --------------------------------------------------------------------------------------
class CombinedBook(X.Book):
    """The in-memory CA + US workbook. Book market = CA (the file lands in the CA outputs and hosts the CA analyses); every
    table declares its own market (book.table(..., market=US) for US tables) and side-by-side tables carry per-column
    markets (ColumnSpec.market). Money headers must carry the column market's currency; nothing is converted."""

    def __init__(self, month: str):
        super().__init__(C.combined_gauge_report_name(month), C.MARKETS["CA"])
        self.extras: dict[int, dict] = {}

    def table(self, ws: Worksheet, spec: TableSpec, top_row: int, *, first_col: int = 1, freeze: bool = True,
              market: C.Market | None = None) -> TableRange:
        mk = market or self.market
        cms = self._check_combined(ws.title, spec, mk)
        tr = X.write_table(ws, spec, top_row, mk, first_col=first_col, freeze=freeze)
        self.tables.append(tr)
        self.note(tr, market=mk.code, column_markets=cms)
        return tr

    def add(self, tr: TableRange, **extra) -> TableRange:
        """Register a range written by an engine function directly (KPI matrix, count cells, lines, metadata)."""
        self.tables.append(tr)
        if extra:
            self.note(tr, **extra)
        return tr

    def note(self, tr: TableRange, **extra) -> None:
        self.extras.setdefault(id(tr), {}).update(extra)

    def _check_combined(self, ws_title: str, spec: TableSpec, mk: C.Market) -> list[str | None]:
        allowed = tuple(spec.allowed_markets)
        where = f"{self.name}!{ws_title}/{spec.role} {spec.title!r}"
        if not allowed or not set(allowed) <= set(MARKET_CODES):
            raise ValueError(f"{where}: allowed_markets {allowed} must be a non-empty subset of {MARKET_CODES}")
        if mk.code not in allowed:
            raise ValueError(f"{where}: table market {mk.code} not in allowed_markets {allowed}")
        side = len(allowed) > 1
        cms: list[str | None] = []
        for col in spec.columns:
            eff = col.market or mk.code
            if eff not in allowed:
                raise ValueError(f"{where}: column {col.header!r} uses market {eff}, not allowed {allowed}")
            ccy = C.MARKETS[eff].currency
            if col.kind in ("money", "money2") and f"({ccy})" not in col.header:
                raise ValueError(f"{where}: money column {col.header!r} must say '({ccy})'")
            other = [C.MARKETS[m].currency for m in MARKET_CODES if m != eff]
            if any(f"({o})" in col.header for o in other):
                raise ValueError(f"{where}: column {col.header!r} of market {eff} names another currency")
            if col.kind == "link" and eff != mk.code and eff not in col.header:
                raise ValueError(f"{where}: {C.MARKETS[eff].domain} link column {col.header!r} must say {eff!r}")
            if side:
                implicit = col.kind in ("money", "money2", "link") or col.header.startswith(f"{mk.code} ")
                cms.append(col.market or (mk.code if implicit else None))
            else:
                cms.append(eff)
        if "market" in spec.rows.columns:
            rows_m = {m for m in spec.rows["market"].dropna().unique() if m}
            if not rows_m <= set(allowed) or (not side and rows_m and rows_m != {mk.code}):
                raise ValueError(f"{where}: rows from market(s) {sorted(rows_m)} in a {mk.code} table (allowed {allowed})")
        return cms

    def registry_extras(self, tr: TableRange) -> dict:
        allowed = tuple(tr.allowed_markets)
        market = allowed[0] if len(allowed) == 1 else None
        e = {"market": market, "column_markets": [market] * len(tr.columns), "subtotal_rows": [], "excluded_from_total_rows": [],
             "placeholder_rows": []}
        e.update(self.extras.get(id(tr), {}))
        if len(e["column_markets"]) != len(tr.columns):
            raise AssertionError(f"{tr.sheet}/{tr.title}: column_markets {e['column_markets']} vs columns {tr.columns}")
        return e


class _Recorder(X.Book):
    """Throw-away single-market book: one build_gauge_report sheet writer runs into it and every table spec, text sheet and
    metadata dict it writes is recorded, so the combined workbook repeats single-market tables from the single-market code
    itself. Never saved."""

    def __init__(self, market: C.Market):
        super().__init__("_single_market_recorder.xlsx", market)
        self.specs: list[tuple[TableSpec, int, dict]] = []
        self.texts: list[tuple[str, list[str], str]] = []
        self.meta: dict | None = None

    def table(self, ws, spec, top_row, **kw):
        tr = super().table(ws, spec, top_row, **kw)
        self.specs.append((spec, top_row, dict(kw)))
        return tr

    def text(self, ws, title, paragraphs, *, role):
        paragraphs = list(paragraphs)
        self.texts.append((title, paragraphs, role))
        return super().text(ws, title, paragraphs, role=role)

    def metadata(self, ws, entries):
        self.meta = dict(entries)
        return super().metadata(ws, entries)

    def save(self, path):
        raise RuntimeError("recorder books are never saved")


def _record(writer: Callable[[X.Book, G.GCtx], None], c: G.GCtx) -> _Recorder:
    rec = _Recorder(c.market)
    writer(rec, c)
    return rec


def _specs(rec: _Recorder, sheet: str, role: str | None = None) -> list[tuple[TableSpec, int, dict]]:
    out = [(s, top, kw) for s, top, kw in rec.specs if s.sheet == sheet and (role is None or s.role == role)]
    if not out:
        raise AssertionError(f"build_gauge_report wrote no {role or 'any'} table on {sheet!r}")
    return out


def _fills(ws: Worksheet, book: CombinedBook, tr: TableRange, *, data: bool) -> None:
    """Market highlight: each market column's header gets FILL_MARKET_HEADER[m]; with data=True its data and residual cells
    get FILL_MARKET_DATA[m] (the Total row keeps the house dark fill)."""
    cms = book.extras[id(tr)]["column_markets"]
    rows = list(range(tr.first_data_row, tr.last_data_row + 1)) + ([tr.residual_row] if tr.residual_row else [])
    for j, m in enumerate(cms):
        if not m:
            continue
        col = tr.first_col + j
        if tr.header_row:
            ws.cell(tr.header_row, col).fill = HEADER_FILL[m]
        if data:
            for r in rows:
                ws.cell(r, col).fill = DATA_FILL[m]


# --------------------------------------------------------------------------------------
# Market contexts (build_gauge_report's own path)
# --------------------------------------------------------------------------------------
def market_context(ds: C.CaDataset, *, preclassified: bool, gauge_map: Path, runs_dir: Path, rederive: bool) -> G.GCtx:
    """build_gauge_workbook's context for one market: FULL code-reader totals first, then the union through
    gauge_union_with_flags (candidate pre-filter for US, union, classify with the frozen prior, freeze/replay)."""
    market = C.MARKETS[ds.market]
    cr_totals = G.code_reader_totals(ds.code_reader, ds.market)
    u, notes, fuel_absent = G.gauge_union_with_flags(ds, preclassified=preclassified, gauge_map_path=gauge_map, runs_dir=runs_dir,
                                                     rederive=rederive)
    return G.GCtx(ds=ds, market=market, u=u, notes=notes, month=ds.month, mon=X.month_label(ds.month),
                  sub=X.subtitle_text(market, ds.export_dates, ds.month), ccy=market.currency, cr_totals=cr_totals,
                  fuel_absent=fuel_absent)


def _adjacent(c: G.GCtx) -> pd.DataFrame:
    return c.u[c.u["gauge_class"].isin(C.GAUGE_ADJACENT_CLASSES)]


def _accessories(c: G.GCtx) -> pd.DataFrame:
    return c.u[c.u["gauge_class"].isin(C.GAUGE_ACCESSORY_CLASSES)]


def _fuel(c: G.GCtx) -> pd.Series:
    """fuel_scope of the core rows (build_gauge_report's rule: absent column -> 'unspecified'; unknown values fail)."""
    core = c.core
    fuel = pd.Series("unspecified", index=core.index) if c.fuel_absent else core["fuel_scope"].astype(str)
    bad = sorted(set(fuel) - set(C.FEATURE_FUEL_SCOPE))
    if bad:
        raise ValueError(f"{c.code} core devices with fuel_scope outside FEATURE_FUEL_SCOPE {C.FEATURE_FUEL_SCOPE}: {bad} "
                         f"(e.g. {core.loc[~fuel.isin(C.FEATURE_FUEL_SCOPE), 'asin'].tolist()[:5]})")
    return fuel


def _mstats(sub: pd.DataFrame, total_rev: float) -> dict:
    rev = float(sub["revenue_month"].sum())
    return {"n": int(sub["asin"].nunique()), "rev": rev, "units": float(sub["units_month"].sum()),
            "share": X.safe_div(rev, total_rev), "rating": X.weighted_rating(sub)}


def _put(d: dict, code: str, stats: dict, keys: Sequence[str]) -> dict:
    d.update({f"{code.lower()}_{k}": stats[k] for k in keys})
    return d


# --------------------------------------------------------------------------------------
# Summary column sets (headers frozen here; the validator can import them)
# --------------------------------------------------------------------------------------
def _block(code: str, specs: Sequence[tuple[str, str, str, float]]) -> list[Col]:
    return [Col(f"{code} {h}", f"{code.lower()}_{f}", kind, w, market=code) for h, f, kind, w in specs]


def _ccy(code: str) -> str:
    return C.MARKETS[code].currency


def share_columns() -> list[Col]:
    return [Col("Measure", "label", "text", 40)] + [
        c for m in MARKET_CODES for c in _block(m, (("# ASINs", "n", "int", 10), (f"Monthly Rev ({_ccy(m)})", "rev", "money", 15),
                                                    ("Monthly Units", "units", "int", 12), ("Share of revenue", "s_rev", "pct2", 12),
                                                    ("Share of units", "s_u", "pct2", 12)))]


def brand_columns() -> list[Col]:
    return [Col("Brand", "brand", "text", 28)] + [
        c for m in MARKET_CODES for c in _block(m, (("# of Listings", "n", "int", 11), (f"Monthly Rev ({_ccy(m)})", "rev", "money", 15),
                                                    ("Monthly Units", "units", "int", 12), ("Rev Share", "share", "pct", 10),
                                                    ("Avg Rating", "rating", "rating", 9)))]


def subtype_columns() -> list[Col]:
    return [Col("Sub-type", "label", "text", 28)] + [
        c for m in MARKET_CODES for c in _block(m, (("# ASINs", "n", "int", 10), (f"Monthly Rev ({_ccy(m)})", "rev", "money", 15),
                                                    ("Monthly Units", "units", "int", 12), ("Rev share", "share", "pct", 10)))]


def fuel_columns() -> list[Col]:
    return [Col("Fuel / Sub-type", "label", "text", 28)] + [
        c for m in MARKET_CODES for c in _block(m, (("# ASINs", "n", "int", 10), (f"Monthly Rev ({_ccy(m)})", "rev", "money", 15),
                                                    ("Monthly Units", "units", "int", 12)))]


def tier_matrix_columns(code: str) -> list[Col]:
    ccy = _ccy(code)
    cols = [Col("Sub-type", "label", "text", 28)]
    for i, (label, _, _) in enumerate(C.GAUGE_TIERS):
        cols += [Col(f"{label} Rev ({ccy})", f"t{i}_rev", "money", 14, market=code), Col(f"{label} Units", f"t{i}_units", "int", 10, market=code)]
    return cols + [Col(f"All tiers Rev ({ccy})", "all_rev", "money", 15, market=code), Col("All tiers Units", "all_units", "int", 11, market=code)]


def tier_matrix_headers(ccy: str) -> tuple[str, ...]:
    code = {m.currency: m.code for m in C.MARKETS.values()}[ccy]
    return tuple(c.header for c in tier_matrix_columns(code))


SHARE_HEADERS: tuple[str, ...] = tuple(c.header for c in share_columns())
BRAND_HEADERS: tuple[str, ...] = tuple(c.header for c in brand_columns())
SUBTYPE_HEADERS: tuple[str, ...] = tuple(c.header for c in subtype_columns())
FUEL_HEADERS: tuple[str, ...] = tuple(c.header for c in fuel_columns())


# --------------------------------------------------------------------------------------
# Summary data
# --------------------------------------------------------------------------------------
def key_figure_items(ca: G.GCtx, us: G.GCtx) -> list[tuple[str, dict[str, tuple[Any, str]], str]]:
    def vals(c: G.GCtx) -> list[tuple[Any, str]]:
        core, dev = c.core, c.device
        return [(float(core["revenue_month"].sum()), "money"), (float(core["units_month"].sum()), "int"),
                (int(core["asin"].nunique()), "int"), (int((core["units_month"] > 0).sum()), "int"),
                (float(dev["revenue_month"].sum()), "money"), (float(_accessories(c)["revenue_month"].sum()), "money"),
                (float(_adjacent(c)["revenue_month"].sum()), "money")]
    units = ("CAD / USD", "units", "count", "count", "CAD / USD", "CAD / USD", "CAD / USD")
    per = {c.code: vals(c) for c in (ca, us)}
    return [(label, {m: per[m][i] for m in MARKET_CODES}, units[i]) for i, label in enumerate(KEY_FIGURE_LABELS)]


def brand_side_by_side(ca_core: pd.DataFrame, us_core: pd.DataFrame, *, top_n: int) -> tuple[pd.DataFrame, dict | None, dict]:
    """(rows, residual, total) of the CA vs US brand table over core devices.

    Rows = union of brand keys of both markets; the top_n shown are chosen by max(CA revenue, US revenue) (nominal, a
    selection rule only), displayed by CA revenue desc, then US revenue desc, then name. Per market: listings, revenue,
    units, share of that market's core revenue, units-weighted rating (blank when none). Residual 'Other brands (n)' =
    each market's rest; Total = each market's full core totals."""
    disp = {**dict(zip(us_core["brand_key"], us_core["brand_display"])), **dict(zip(ca_core["brand_key"], ca_core["brand_display"]))}
    revs = {"ca": ca_core.groupby("brand_key")["revenue_month"].sum(), "us": us_core.groupby("brand_key")["revenue_month"].sum()}

    def rev(m: str, k: str) -> float:
        return float(revs[m].get(k, 0.0))

    keys = sorted(disp)
    chosen = sorted(keys, key=lambda k: (-max(rev("ca", k), rev("us", k)), disp[k], k))[:top_n]
    rest = sorted(set(keys) - set(chosen))
    shown = sorted(chosen, key=lambda k: (-rev("ca", k), -rev("us", k), disp[k], k))
    totals = {"CA": float(ca_core["revenue_month"].sum()), "US": float(us_core["revenue_month"].sum())}

    def row(label: str, keyset: set[str], brand_key: str | None) -> dict:
        d: dict = {"brand": label, "brand_key": brand_key}
        for code, core in (("CA", ca_core), ("US", us_core)):
            _put(d, code, _mstats(core[core["brand_key"].isin(keyset)], totals[code]), ("n", "rev", "units", "share", "rating"))
        return d

    fields = ["brand", "brand_key"] + [f"{m}_{k}" for m in ("ca", "us") for k in ("n", "rev", "units", "share", "rating")]
    frame = pd.DataFrame([row(disp[k], {k}, k) for k in shown], columns=fields)
    residual = row(C.RESIDUAL_ROW_LABEL.format(noun="brands", n=len(rest)), set(rest), None) if rest else None
    total = row(C.TOTAL_ROW_LABEL, set(keys), None)
    return frame, residual, total


def _hide_rows_without_revenue(ws: Worksheet, tr: TableRange) -> int:
    """Hide (never delete) brand rows with zero revenue in BOTH markets (the single-market Summary hides zero rows too)."""
    cols = [tr.col(f"{m} Monthly Rev ({_ccy(m)})") for m in MARKET_CODES]
    hidden = 0
    for r in range(tr.first_data_row, tr.last_data_row + 1):
        if all(isinstance(ws.cell(r, c).value, (int, float)) and ws.cell(r, c).value == 0 for c in cols):
            ws.row_dimensions[r].hidden = True
            hidden += 1
    return hidden


def _summary(book: CombinedBook, ca: G.GCtx, us: G.GCtx) -> None:
    ws = book.sheet("Summary")
    T = C.COMBINED_SUMMARY_TITLES
    ctx = {"CA": ca, "US": us}
    # 1. key figures (rows = KPIs, columns CA | US, each row in its market's currency)
    tr = X.write_market_kpi_table(ws, 1, T["key_figures"][0], key_figure_items(ca, us), MARKET_CODES,
                                  subtitle=f"CA vs US · Report month {ca.mon} · each figure in its market's own currency, never converted",
                                  note=f"CA = {ca.sub}. US = {us.sub}. {KEY_FIGURES_NOTE}", header_fills=C.FILL_MARKET_HEADER,
                                  data_fills=C.FILL_MARKET_DATA, role=T["key_figures"][1],
                                  dataset_filter="see metric labels; each market from its own rows")
    book.add(tr, market=None, column_markets=[None, *MARKET_CODES, None])
    # 2. gauge share of the code-reader market: ONE row, (b) all core devices; share cells bold
    row = {"label": C.COMBINED_SHARE_ROW_LABEL}
    for code, c in ctx.items():
        b = G.market_share_rows(c.u, c.cr_totals, code)[2]
        if b["label"] != C.COMBINED_SHARE_ROW_LABEL:
            raise AssertionError(f"{code}: market_share_rows()[2] is {b['label']!r}, not the (b) row")
        _put(row, code, b, ("n", "rev", "units", "s_rev", "s_u"))
    tr = book.table(ws, TableSpec("Summary", T["share"][1], T["share"][0], share_columns(), pd.DataFrame([row]), None, None,
                                  MARKET_CODES, note=COMBINED_SHARE_NOTE,
                                  dataset_filter="core devices vs the full code-reader export, per market"),
                    X.table_end(tr) + 2, freeze=False)
    _fills(ws, book, tr, data=True)
    for code in MARKET_CODES:
        for h in (f"{code} Share of revenue", f"{code} Share of units"):
            ws.cell(tr.first_data_row, tr.col(h)).font = Font(bold=True)
    # 3. brand summary (core devices), two bar charts right of the table
    frame, residual, total = brand_side_by_side(ca.core, us.core, top_n=C.SUMMARY_TOP_BRANDS)
    cols = brand_columns()
    tr_b = book.table(ws, TableSpec("Summary", T["brands"][1], T["brands"][0], cols, frame, X.fit(total, cols), X.fit(residual, cols),
                                    MARKET_CODES, note=BRAND_NOTE.format(n=C.SUMMARY_TOP_BRANDS),
                                    dataset_filter="core devices, both markets: gauge_device_scope & ~borderline"),
                      X.table_end(tr) + 3, freeze=False)
    _fills(ws, book, tr_b, data=True)
    _hide_rows_without_revenue(ws, tr_b)
    for code in MARKET_CODES:
        book.bar(ws, tr_b, "Brand", f"{code} Monthly Rev ({_ccy(code)})", f"{code} core device revenue by brand ({_ccy(code)})")
    # 4. sub-type mix (core) + the adjacent GPS-only HUD row (excluded from the Total)
    rows = []
    for cls in C.GAUGE_DEVICE_CLASSES:
        d = {"label": C.GAUGE_SUBTYPE_LABELS[cls]}
        for code, c in ctx.items():
            _put(d, code, _mstats(c.core[c.core["gauge_class"] == cls], float(c.core["revenue_month"].sum())), ("n", "rev", "units", "share"))
        rows.append(d)
    d = {"label": GPS_ROW_LABEL}
    for code, c in ctx.items():
        _put(d, code, _mstats(_adjacent(c), float("nan")), ("n", "rev", "units", "share"))
    rows.append(d)
    tot = {"label": C.TOTAL_ROW_LABEL}
    for code, c in ctx.items():
        _put(tot, code, _mstats(c.core, float(c.core["revenue_month"].sum())), ("n", "rev", "units", "share"))
    tr = book.table(ws, TableSpec("Summary", T["subtypes"][1], T["subtypes"][0], subtype_columns(), pd.DataFrame(rows), tot, None,
                                  MARKET_CODES, note=SUBTYPE_NOTE,
                                  dataset_filter="core devices, both markets; + gauge_class in GAUGE_ADJACENT_CLASSES (row excluded from Total)"),
                    X.table_end(tr_b) + 3, freeze=False)
    _fills(ws, book, tr, data=True)
    book.note(tr, excluded_from_total_rows=[tr.last_data_row])
    # 5./6. price tier x sub-type, one table per market, revenue and units together
    r = X.table_end(tr) + 3
    for key, c in (("tier_ca", ca), ("tier_us", us)):
        cols = tier_matrix_columns(c.code)
        core = c.core
        mrows = []
        for cls in C.GAUGE_DEVICE_CLASSES + (None,):
            sub = core if cls is None else core[core["gauge_class"] == cls]
            d = {"label": C.TOTAL_ROW_LABEL if cls is None else C.GAUGE_SUBTYPE_LABELS[cls],
                 "all_rev": float(sub["revenue_month"].sum()), "all_units": float(sub["units_month"].sum())}
            for i, (label, _, _) in enumerate(C.GAUGE_TIERS):
                t = sub[sub["_gtier"] == label]
                d[f"t{i}_rev"], d[f"t{i}_units"] = float(t["revenue_month"].sum()), float(t["units_month"].sum())
            mrows.append(d)
        tr = book.table(ws, TableSpec("Summary", C.COMBINED_SUMMARY_TITLES[key][1], C.COMBINED_SUMMARY_TITLES[key][0], cols,
                                      pd.DataFrame(mrows[:-1]), mrows[-1], None, (c.code,), note=TIER_NOTE.format(ccy=c.ccy),
                                      dataset_filter="core devices; tiers = GAUGE_TIERS on price (half-open)"),
                        r, freeze=False, market=c.market)
        _fills(ws, book, tr, data=True)
        r = X.table_end(tr) + 3
    # 7. fuel split: '<fuel> — subtotal' rows, each followed by its sub-type rows (>= 1 ASIN in either market), then Total
    fuels = {code: _fuel(c) for code, c in ctx.items()}
    rows, group_idx = [], []
    for f in C.FEATURE_FUEL_SCOPE:
        d = {"label": FUEL_SUBTOTAL_LABEL.format(fuel=f)}
        for code, c in ctx.items():
            _put(d, code, _mstats(c.core[fuels[code] == f], float("nan")), ("n", "rev", "units"))
        group_idx.append(len(rows))
        rows.append(d)
        for cls in C.GAUGE_DEVICE_CLASSES:
            d = {"label": C.GAUGE_SUBTYPE_LABELS[cls]}
            for code, c in ctx.items():
                m = (fuels[code] == f) & (c.core["gauge_class"] == cls)
                _put(d, code, _mstats(c.core[m], float("nan")), ("n", "rev", "units"))
            if d["ca_n"] + d["us_n"] >= 1:
                rows.append(d)
    tot = {"label": C.TOTAL_ROW_LABEL}
    for code, c in ctx.items():
        _put(tot, code, _mstats(c.core, float("nan")), ("n", "rev", "units"))
    tr = book.table(ws, TableSpec("Summary", T["fuel"][1], T["fuel"][0], fuel_columns(), pd.DataFrame(rows), tot, None, MARKET_CODES,
                                  note=G.FUEL_NOTE, dataset_filter="core devices, both markets; fuel_scope in FEATURE_FUEL_SCOPE"),
                    r, freeze=False)
    _fills(ws, book, tr, data=True)
    sub_rows = [tr.first_data_row + i for i in group_idx]
    for i in range(tr.n_rows):
        cell = ws.cell(tr.first_data_row + i, tr.first_col)
        if i in group_idx:
            for j in range(len(tr.columns)):
                ws.cell(tr.first_data_row + i, tr.first_col + j).font = Font(bold=True)
        else:
            cell.alignment = Alignment(indent=1)
    book.note(tr, subtotal_rows=sub_rows)


# --------------------------------------------------------------------------------------
# Other sheets
# --------------------------------------------------------------------------------------
def _section(paras: list[str], heading: str) -> list[str]:
    """The heading paragraph ('# <heading>') and its body, up to the next heading."""
    h = f"# {heading}"
    if h not in paras:
        raise AssertionError(f"build_gauge_report Read Me has no {h!r} section")
    i = paras.index(h)
    j = next((k for k in range(i + 1, len(paras)) if paras[k].startswith("# ")), len(paras))
    return paras[i:j]


def _files(c: G.GCtx) -> str:
    return "; ".join(f"{n} ({r} rows)" for n, r in c.ds.raw_files) or "none recorded"


def _read_me(book: CombinedBook, ca: G.GCtx, us: G.GCtx) -> None:
    rec = _record(G._read_me, ca)
    ca_paras = rec.texts[0][1]
    ws = book.sheet("Read Me")
    paras = ["# Purpose",
             f"Side-by-side monthly snapshot ({ca.mon}) of OBD gauge, HUD and gauge-tuner listings on amazon.ca (CA, CAD) and "
             f"amazon.com (US, USD), built from Helium 10 Black Box exports. Revenue and unit figures are Helium 10 estimates in "
             f"each market's own currency; CAD and USD are never converted, added or divided into one figure.",
             "# How to read this workbook",
             "Summary: Key figures, the gauge share of the code-reader market, then brand, sub-type, price tier × sub-type and "
             "fuel tables. Each Summary table holds a CA block (purple header) and a US block (blue header) side by side; shares "
             "are computed within each market.",
             "Top 50 CA / Top 50 US: the single-market Top 50 rankings (by revenue, then by units).",
             "Innova: Innova gauge/HUD device counts per market and the CA app-capable Innova hardware (Model B).",
             "Brand tabs: one per brand that qualifies in either market; four tables — CA — Rank by Revenue, CA — Rank by Units, "
             "US — Rank by Revenue, US — Rank by Units.",
             "Price Ladder (Model A), Feature Matrix (Model A), App-Gauge Proxy (Model B): the CA analyses, as in the CA workbook. "
             "US vs CA Same-ASIN: listings sold in both marketplaces.",
             "All Products, Dedupe & Classification Audit, Excluded: one table per market ('— CA', then '— US')."]
    paras += _section(ca_paras, "Business models")
    paras += ["# Source files", f"CA: {_files(ca)}", f"US: {_files(us)}", "# Dedupe rule", G.DEDUPE_RULE]
    paras += _section(ca_paras, "Classification") + _section(ca_paras, "Caveats")
    tr = book.text(ws, f"CA + US OBD Gauge Competitor Report — Read Me ({ca.mon})", paras, role="read_me")
    (spec, _, kw), = _specs(rec, "Read Me", "read_me")
    book.table(ws, replace(spec, allowed_markets=MARKET_CODES), tr.last_data_row + 2, **kw)
    ws.column_dimensions["A"].width = 120


def _top50(book: CombinedBook, c: G.GCtx, sheet: str) -> None:
    """The single-market Top 50 sheet (both rankings, same specs and rows) under the market's own sheet name."""
    rec = _record(G._top50, c)
    ws = book.sheet(sheet)
    for spec, top, kw in _specs(rec, "Top 50"):
        tr = book.table(ws, replace(spec, sheet=sheet), top, market=c.market, **kw)
        _fills(ws, book, tr, data=False)


def _innova(book: CombinedBook, ca: G.GCtx, us: G.GCtx) -> None:
    rec = _record(G._innova, ca)
    (spec, top, kw), = _specs(rec, "Innova", "modelb_top")
    ws = book.sheet("Innova")
    X.write_sheet_header(ws, f"Innova — CA + US OBD gauge view ({ca.mon})", f"CA = {ca.sub} | US = {us.sub}", len(spec.columns))
    flt = "brand_key == 'innova' & gauge_device_scope"
    for row, c, cell in ((3, ca, INNOVA_CA_COUNT_CELL), (4, us, INNOVA_US_COUNT_CELL)):
        n = int(((c.u["brand_key"] == "innova") & c.u["gauge_device_scope"]).sum())
        label = f"Innova gauge/HUD device listings in the {c.code} dataset"
        book.add(X.write_line(ws, row, f"{label}: {n}", c.market, role="innova", dataset_filter=flt))
        tr = book.add(X.write_number(ws, row, 2, n, c.market, label=label, dataset_filter=flt))
        if ws.cell(tr.first_data_row, tr.first_col).coordinate != cell:
            raise AssertionError(f"Innova {c.code} count cell moved")
    X.set_text(ws.cell(5, 1), INNOVA_US_NOTE)
    ws.cell(5, 1).font = X.FONT_NOTE
    book.table(ws, spec, top, market=ca.market, **kw)
    book.brand_sheet_map["innova"] = "Innova"


def qualifying_brands(c: G.GCtx) -> set[str]:
    """Brand keys that get a brand tab in the single-market workbook (build_gauge_report's own rule, via its writer)."""
    return set(_record(G._brand_tabs, c).brand_sheet_map) - {"innova"}


def _brand_tabs(book: CombinedBook, ca: G.GCtx, us: G.GCtx) -> None:
    keys = qualifying_brands(ca) | qualifying_brands(us)
    disp = {**dict(zip(us.u["brand_key"], us.u["brand_display"])), **dict(zip(ca.u["brand_key"], ca.u["brand_display"]))}
    rev = {c.code: c.core.groupby("brand_key")["revenue_month"].sum() for c in (ca, us)}
    order = sorted(keys, key=lambda k: (-float(rev["CA"].get(k, 0.0)), -float(rev["US"].get(k, 0.0)), disp[k], k))
    taken = {s.lower() for s in C.COMBINED_FIXED_SHEETS + C.COMBINED_MODEL_SHEETS + C.COMBINED_TAIL_SHEETS}
    for key in order:
        name = C.sheet_name_for_brand(disp[key], taken)
        _brand_tab(book, name, key, disp[key], ca, us)
        book.brand_sheet_map[key] = name


def _brand_tab(book: CombinedBook, name: str, key: str, display: str, ca: G.GCtx, us: G.GCtx) -> None:
    """Rows 1-6: title, subtitle (A2) + KPI column labels CA | US (B2:C2), KPI block rows 3-6; then the four ranking tables
    (COMBINED_BRAND_TABLE_TITLES) with the single-market brand-tab columns and Total rows."""
    ws = book.sheet(name)
    cols_by = {c.code: X.brand_tab_columns(c.ccy, type_header="Sub-type", type_field="_subtype") for c in (ca, us)}
    ncol = len(cols_by["CA"])
    X.write_sheet_header(ws, f"{display} — CA + US core gauge devices ({ca.mon})", None, ncol)
    sub = ws.cell(2, 1)
    X.set_text(sub, f"CA = amazon.ca, CAD · US = amazon.com, USD (raw Helium 10, no actuals overlay) · Report month {ca.mon}")
    sub.font, sub.alignment = X.FONT_SUBTITLE, Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30
    for j, code in enumerate(MARKET_CODES):
        cell = ws.cell(2, 2 + j)
        X.set_text(cell, code)
        cell.font, cell.fill, cell.border = X.FONT_HEADER, HEADER_FILL[code], X.BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    per = {}
    for c in (ca, us):
        rows = c.core[c.core["brand_key"] == key]
        per[c.code] = [(v, kind) for _, v, kind in X.brand_kpi_items(rows, float(c.core["revenue_month"].sum()), c.ccy)]
    items = [(label, {m: per[m][i] for m in MARKET_CODES}, "") for i, label in enumerate(BRAND_KPI_LABELS)]
    flt = f"core devices & brand_key == {key!r}"
    tr = X.write_market_kpi_table(ws, 3, None, items, MARKET_CODES, header=False, label_header="Metric", unit_header=None,
                                  data_fills=C.FILL_MARKET_DATA, allow_blank=True, dataset_filter=flt)
    book.add(tr, market=None, column_markets=[None, *MARKET_CODES])
    if (tr.first_data_row, tr.last_data_row) != (3, C.BRAND_TAB_RESERVED_ROWS):
        raise AssertionError("combined brand tab KPI block must fill rows 3-6")
    r = C.BRAND_TAB_RESERVED_ROWS + 1
    titles = iter(C.COMBINED_BRAND_TABLE_TITLES)
    for c in (ca, us):
        rows = c.core[c.core["brand_key"] == key]
        cols = cols_by[c.code]
        total = X.fit(X.listing_totals(rows, cols[0].field, C.TOTAL_ROW_LABEL), cols)
        for role, by in (("brand_tab_revenue", "revenue"), ("brand_tab_units", "units")):
            title = next(titles)
            if not title.startswith(f"{c.code} — "):
                raise AssertionError(f"COMBINED_BRAND_TABLE_TITLES order drifted: {title!r} for {c.code}")
            if len(rows):
                body = X.rank_listings(rows, by)
            else:
                body = pd.DataFrame([{col.field: (NO_LISTINGS_LABEL if col.field == cols[0].field else None) for col in cols}])
            t = book.table(ws, TableSpec(name, role, title, cols, body, total, None, (c.code,), dataset_filter=flt), r, market=c.market)
            if r == C.BRAND_TAB_RESERVED_ROWS + 1 and t.header_row != C.BRAND_TAB_RESERVED_ROWS + 2:
                raise AssertionError("combined brand tab: first ranking header must be at row 8")
            _fills(ws, book, t, data=True)
            if not len(rows):
                book.note(t, placeholder_rows=[t.first_data_row])
            r = X.table_end(t) + 4


def _stacked(book: CombinedBook, sheet: str, writer: Callable[[X.Book, G.GCtx], None], role: str, ca: G.GCtx, us: G.GCtx,
             extra_note: Callable[[G.GCtx], str] | None = None) -> None:
    """Two stacked tables '<sheet> — CA' and '<sheet> — US': the first `role` table the single-market writer puts on
    `sheet`, with its columns and rows unchanged (its title moves into the note)."""
    ws = book.sheet(sheet)
    r = 1
    for c in (ca, us):
        spec, _, _ = _specs(_record(writer, c), sheet, role)[0]
        note = ". ".join(x.rstrip(".") for x in (spec.title, spec.note, extra_note(c) if extra_note else None) if x) + "."
        tr = book.table(ws, replace(spec, title=f"{sheet} — {c.code}", subtitle=c.sub, note=note), r, market=c.market, freeze=False)
        _fills(ws, book, tr, data=False)
        r = X.table_end(tr) + 4


def _source_method(book: CombinedBook, ca: G.GCtx, us: G.GCtx) -> None:
    ca_title, ca_paras, _ = _record(G._source_method, ca).texts[0]
    _, us_paras, _ = _record(G._source_method, us).texts[0]
    us_only = [p for p in us_paras if p not in ca_paras]
    paras = list(ca_paras) + list(COMBINED_LAYOUT_PARAGRAPHS) + ["# US inputs"] + us_only + [G.US_BENCHMARK_SOURCE + ".", G.CAVEAT_US_RAW]
    book.text(book.sheet("Source & Method"), ca_title, paras, role="source_method")


def _metadata(book: CombinedBook, ca: G.GCtx, us: G.GCtx) -> None:
    m_ca, m_us = _record(G._metadata, ca).meta, _record(G._metadata, us).meta
    if m_ca is None or m_us is None:
        raise AssertionError("build_gauge_report._metadata wrote no metadata")
    if m_ca["Report month"] != m_us["Report month"]:
        raise ValueError(f"report months differ: {m_ca['Report month']} vs {m_us['Report month']}")
    d = {"Marketplace": "amazon.ca (CA) + amazon.com (US)", "Currency": "CAD and USD, never mixed in one figure",
         "Report month": m_ca["Report month"],
         "Type coverage": f"CA: {m_ca['Type coverage']}\nUS: not typed ({m_us['Type coverage']})",
         "Pipeline version": m_ca["Pipeline version"], "Generated at": datetime.now().isoformat(timespec="seconds"),
         "Scope note": f"CA: {m_ca['Scope note']}\nUS: {m_us['Scope note']}\n{SCOPE_NOTE_COMBINED}"}
    for k in list(m_ca) + [k for k in m_us if k not in m_ca]:
        if k in d:
            continue
        if k in m_ca and k in m_us:
            d[k] = f"CA: {m_ca[k]}\nUS: {m_us[k]}"
        else:
            d[k] = f"{'CA' if k in m_ca else 'US'}: {m_ca.get(k, m_us.get(k))}"
    d["US source"] = G.US_BENCHMARK_SOURCE
    d["Combined layout"] = ("Summary tables hold CA and US blocks side by side (one table each); Top 50 CA / Top 50 US; brand "
                            "tabs with four tables (CA then US); All Products / Audit / Excluded with one table per market")
    book.add(X.write_metadata_sheet(book.sheet("Metadata"), d, None))


def build_combined_book(ca: G.GCtx, us: G.GCtx) -> CombinedBook:
    if (ca.code, us.code) != MARKET_CODES:
        raise ValueError(f"contexts must be CA then US, got {ca.code}, {us.code}")
    if ca.modelb is None or ca.app_matrix is None or us.u is not ca.bu:
        raise ValueError("CA context needs modelb, app_matrix and bu = the US union (Model sheets / Same-ASIN)")
    book = CombinedBook(ca.month)
    _read_me(book, ca, us)
    _summary(book, ca, us)
    _top50(book, ca, "Top 50 CA")
    _top50(book, us, "Top 50 US")
    _innova(book, ca, us)
    _brand_tabs(book, ca, us)
    G._price_ladder(book, ca)
    G._feature_matrix(book, ca)
    G._modelb(book, ca)
    G._same_asin(book, ca)
    _stacked(book, "All Products", G._all_products, "all_rows", ca, us)
    _stacked(book, "Dedupe & Classification Audit", G._audit, "dedupe_audit", ca, us,
             extra_note=lambda c: f"Per-ASIN classification decisions are the matching columns of All Products — {c.code}")
    _stacked(book, "Excluded", G._excluded, "excluded", ca, us)
    _source_method(book, ca, us)
    _metadata(book, ca, us)
    names = book.sheetnames
    nf, nm, nt = len(C.COMBINED_FIXED_SHEETS), len(C.COMBINED_MODEL_SHEETS), len(C.COMBINED_TAIL_SHEETS)
    if (names[:nf] != list(C.COMBINED_FIXED_SHEETS) or names[-nt:] != list(C.COMBINED_TAIL_SHEETS)
            or names[-(nt + nm):-nt] != list(C.COMBINED_MODEL_SHEETS)):
        raise AssertionError(f"combined sheet order drifted: {names}")
    return book


# --------------------------------------------------------------------------------------
# Registry
# --------------------------------------------------------------------------------------
def write_registry(book: CombinedBook, runs_dir: Path, month: str) -> Path:
    """runs/<m>/table_registry_CAUS_<m>.json — the shape of ca_xlsx_style.TableRegistry (market = 'CAUS'); each table entry
    also carries market (table market), column_markets (per column), subtotal_rows, excluded_from_total_rows, placeholder_rows.
    Entries of other workbooks in the file are kept; this workbook's entries are replaced."""
    path = C.run_file(Path(runs_dir), month, "table_registry", REGISTRY_SCOPE, "json")
    tables: list[dict] = []
    workbooks: dict[str, dict] = {}
    if path.exists():
        prev = json.loads(path.read_text(encoding="utf-8"))
        if (prev.get("market"), prev.get("month")) != (REGISTRY_SCOPE, month):
            raise ValueError(f"{path}: registry is for {prev.get('market')}/{prev.get('month')}")
        tables = [t for t in prev["tables"] if t["workbook"] != book.name]
        workbooks = {k: v for k, v in prev["workbooks"].items() if k != book.name}
    now = datetime.now().isoformat(timespec="seconds")
    workbooks[book.name] = {"sheets": book.sheetnames, "brand_sheet_map": dict(book.brand_sheet_map), "generated_at": now}
    for tr in book.tables:
        e = tr.to_registry(book.name, book.brand_sheet_map)
        e.update(book.registry_extras(tr))
        tables.append(e)
    out = {"market": REGISTRY_SCOPE, "month": month, "generated_at": now, "pipeline_version": X.pipeline_version(),
           "workbooks": dict(sorted(workbooks.items())), "tables": tables}
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(out, indent=1, ensure_ascii=False), encoding="utf-8")
    tmp.replace(path)
    return path


# --------------------------------------------------------------------------------------
# Entry points
# --------------------------------------------------------------------------------------
def decision_files(runs_dir: Path, month: str) -> list[Path]:
    """Frozen decision files of both markets that exist (manifest inputs; the type review queue is an output, never listed)."""
    paths = [C.run_file(Path(runs_dir), month, stem, f"{mk}_{src}") for mk in MARKET_CODES for stem in G.DECISION_STEMS
             for src in ("code_reader", "gauge")]
    return [p for p in paths if p.exists()]


def build_combined_gauge_workbook(ca: C.CaDataset, us: C.CaDataset, out_dir: Path, *, overwrite: bool, dated_copy: bool,
                                  runs_dir: Path = C.RUNS_DIR, gauge_map: Path = G.GAUGE_MAP_DEFAULT,
                                  app_feature_matrix: Path = G.APP_FEATURE_MATRIX_DEFAULT, input_paths: Sequence[Path] = (),
                                  app_gauge_brands: Path = G.APP_GAUGE_BRANDS_DEFAULT, rederive: bool = False) -> Path:
    """Build CA_US_OBD_Gauge_Competitor_Report_<m>.xlsx in out_dir; returns its path (a dated copy, when asked, sits next to it).

    Both datasets must come from the same input mode: normalized frames (dev/fixture; already classified) or the loader
    (union/classify/freeze through build_gauge_report.gauge_union_with_flags, replaying runs/<m> unless rederive)."""
    if (ca.market, us.market) != MARKET_CODES:
        raise ValueError(f"datasets must be CA and US, got {ca.market!r} and {us.market!r}")
    if ca.month != us.month:
        raise ValueError(f"months differ: CA {ca.month} vs US {us.month}")
    mode = X.input_mode(ca)
    if X.input_mode(us) != mode:
        raise ValueError(f"input modes differ (never mixed): CA {mode!r} vs US {X.input_mode(us)!r}")
    pre = mode == X.INPUT_MODE_NORMALIZED
    month, out_dir, runs_dir = ca.month, Path(out_dir), Path(runs_dir)
    name = C.combined_gauge_report_name(month)
    if not overwrite and (out_dir / name).exists():
        raise FileExistsError(f"{out_dir / name} exists; pass --overwrite (the old file is backed up)")
    c_ca = market_context(ca, preclassified=pre, gauge_map=gauge_map, runs_dir=runs_dir, rederive=rederive)
    c_us = market_context(us, preclassified=pre, gauge_map=gauge_map, runs_dir=runs_dir, rederive=rederive)
    c_ca.modelb = G.modelb_universe(ca, app_gauge_brands)
    c_ca.app_matrix, c_ca.app_matrix_path = G.read_app_feature_matrix(app_feature_matrix), Path(app_feature_matrix)
    c_ca.bu = c_us.u                       # US vs CA Same-ASIN (the CA workbook's benchmark join)
    book = build_combined_book(c_ca, c_us)
    manifest = X.read_manifest(out_dir, month)
    path = X.safe_output_path(out_dir, book.name, overwrite, manifest)
    book.save(path)
    outputs = [path]
    if dated_copy:
        dp = X.safe_output_path(out_dir, X.dated_copy_name(book.name), overwrite, manifest)
        shutil.copyfile(path, dp)
        outputs.append(dp)
    reg_path = write_registry(book, runs_dir, month)
    G.write_conflict_review(c_ca.u, "CA", month, runs_dir)      # same merge as the CA gauge build (V17); a no-op when already merged
    extra = ([] if pre else [Path(gauge_map), Path(app_gauge_brands)]) + [Path(app_feature_matrix)]
    decisions = [] if pre else decision_files(runs_dir, month)
    inputs = X.input_hashes(list(input_paths) + extra + sorted(C.MAPS_DIR.glob("*.csv")) + decisions)
    man_path = X.write_manifest(out_dir, month, outputs, inputs)
    print(f"registry: {reg_path} ({len(book.tables)} tables)")
    print(f"manifest: {man_path}")
    return path


def parse_args(argv: list[str] | None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Build the combined CA + US OBD gauge competitor workbook")
    p.add_argument("--month", required=True, help="report month YYYYMM")
    p.add_argument("--ca-gauge-raw-dir", type=Path, help="CA gauge export folder (default NewProductCategory/CA-OBD-GAUGE/raw_data/<m>)")
    p.add_argument("--ca-cr-raw-dir", type=Path, help="CA code-reader export folder (default NewProductCategory/CA-CODE-READER/raw_data/<m>)")
    p.add_argument("--us-gauge-raw-dir", type=Path, help="US gauge export folder (default NewProductCategory/US-OBD-GAUGE/raw_data/<m>)")
    p.add_argument("--us-cr-raw-dir", type=Path, help="US code-reader export folder (default the US pipeline raw_data/<m>)")
    p.add_argument("--out-dir", type=Path, help="output folder (default NewProductCategory/CA-OBD-GAUGE/outputs)")
    p.add_argument("--gauge-map", type=Path, default=G.GAUGE_MAP_DEFAULT)
    p.add_argument("--app-feature-matrix", type=Path, default=G.APP_FEATURE_MATRIX_DEFAULT,
                   help="app feature matrix CSV (ca_common.APP_FEATURE_MATRIX_COLUMNS) for App-Gauge Proxy (Model B)")
    p.add_argument("--runs-dir", type=Path, help="frozen run decisions + registry (default ca_market_reports/runs)")
    p.add_argument("--overwrite", action="store_true", help="replace an existing output (old file moved to _backup/)")
    p.add_argument("--rederive", action="store_true", help="re-derive decisions instead of replaying frozen ones (as the single builders)")
    p.add_argument("--dated-copy", action="store_true", help="also write <name>_<YYYYMMDD>.xlsx")
    p.add_argument("--from-normalized", type=Path, help="DEV: CA normalized, already-classified CSV (fixtures)")
    p.add_argument("--us-from-normalized", type=Path, help="DEV: US normalized, already-classified CSV (fixtures)")
    a = p.parse_args(argv)
    if not C.MONTH_RE.match(a.month):
        p.error(f"--month must be a calendar month YYYYMM, got {a.month!r}")
    if bool(a.from_normalized) != bool(a.us_from_normalized):
        p.error("--from-normalized and --us-from-normalized go together (both markets come from the same input mode)")
    raw = [f for f in ("ca_gauge_raw_dir", "ca_cr_raw_dir", "us_gauge_raw_dir", "us_cr_raw_dir") if getattr(a, f) is not None]
    if a.from_normalized and raw:
        p.error(f"dev (--from-normalized) and loader inputs are never mixed: drop {', '.join('--' + f.replace('_', '-') for f in raw)}")
    if a.from_normalized and a.rederive:
        p.error("--rederive needs the loader path (normalized frames are already classified)")
    return a


def main(argv: list[str] | None = None) -> int:
    a = parse_args(argv)
    if a.from_normalized:
        if a.out_dir is None:
            raise SystemExit("--from-normalized needs --out-dir (dev runs never write to the real output folder)")
        X.refuse_new_product_dir(a.out_dir)
        if a.runs_dir is None:
            raise SystemExit("--from-normalized needs --runs-dir (dev runs never write to ca_market_reports/runs)")
        ca = X.dataset_from_normalized(X.read_normalized_csv(a.from_normalized), "CA", a.month)
        us = X.dataset_from_normalized(X.read_normalized_csv(a.us_from_normalized), "US", a.month)
        inputs = [a.from_normalized, a.us_from_normalized]
        out_dir, runs_dir = a.out_dir, a.runs_dir
    else:
        from ca_market_reports.ca_load import load_month  # ASSEMBLY: ca_load.load_month (as build_gauge_report.main)
        runs_dir = a.runs_dir or C.RUNS_DIR
        mca, mus = C.MARKETS["CA"], C.MARKETS["US"]
        ca_cr, ca_g = a.ca_cr_raw_dir or mca.cr_raw_dir(a.month), a.ca_gauge_raw_dir or mca.gauge_raw_dir(a.month)
        us_cr, us_g = a.us_cr_raw_dir or mus.cr_raw_dir(a.month), a.us_gauge_raw_dir or mus.gauge_raw_dir(a.month)
        ca = load_month("CA", a.month, cr_raw_dir=ca_cr, gauge_raw_dir=ca_g, gauge_map=a.gauge_map, runs_dir=runs_dir,
                        assign_types=True, rederive=a.rederive)
        inputs = X.raw_input_files(ca, X.loader_raw_dirs("CA", a.month, ca_cr, ca_g)) + [C.US_TYPE_MAP_DEFAULT]
        us = load_month("US", a.month, cr_raw_dir=us_cr, gauge_raw_dir=us_g, gauge_map=a.gauge_map, runs_dir=runs_dir,
                        assign_types=False, rederive=a.rederive)
        inputs += X.raw_input_files(us, X.loader_raw_dirs("US", a.month, us_cr, us_g))
        out_dir = a.out_dir or mca.gauge_out_dir()
    path = build_combined_gauge_workbook(ca, us, out_dir, overwrite=a.overwrite, dated_copy=a.dated_copy, runs_dir=runs_dir,
                                         gauge_map=a.gauge_map, app_feature_matrix=a.app_feature_matrix, input_paths=inputs,
                                         rederive=a.rederive)
    print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
