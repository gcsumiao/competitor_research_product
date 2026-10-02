"""Combined CA + US OBD gauge workbook tests (track ca-combined-gauge).

Builds, from the normalized fixtures only (dev mode, --from-normalized / --us-from-normalized):
  * the single-market US gauge workbook and the CA gauge workbook (US benchmark) -> the references
  * CA_US_OBD_Gauge_Competitor_Report_<m>.xlsx through the CLI, into the SAME out dir (manifest merge)
into tmp/ca_scratch/test_combined/ and checks the combined layout contract of ca_common ("Combined CA + US gauge workbook").

Run: ca_market_reports/run.sh -m unittest ca_market_reports.tests.test_combined -v
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import json
import math
import shutil
import unittest
from pathlib import Path
from unittest import mock

import openpyxl
import pandas as pd

from ca_market_reports import ca_common as C
from ca_market_reports import ca_xlsx_style as X
from ca_market_reports import build_gauge_report as G
from ca_market_reports import build_combined_gauge_report as CB
from ca_market_reports import render_preview as RP

MONTH = "202609"
OUT = C.SCRATCH_DIR / "test_combined"
RUNS = OUT / "runs"
CA_FIXTURE = C.FIXTURES_DIR / "normalized_rows.csv"
US_FIXTURE = C.FIXTURES_DIR / "normalized_rows_us.csv"
NAME = C.combined_gauge_report_name(MONTH)


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def _core(rows: pd.DataFrame) -> pd.DataFrame:
    """Fixtures carry no borderline rows: core devices = device classes."""
    return rows[rows["gauge_class"].isin(C.GAUGE_DEVICE_CLASSES)]


def _fuel(rows: pd.DataFrame) -> pd.Series:
    core = _core(rows)
    if "fuel_scope" not in core.columns:
        return pd.Series("unspecified", index=core.index)
    return core["fuel_scope"].astype(str)


class _Built:
    done = False

    @classmethod
    def ensure(cls):
        if cls.done:
            return
        if OUT.exists():
            shutil.rmtree(OUT)
        OUT.mkdir(parents=True)
        ca = X.dataset_from_normalized(X.read_normalized_csv(CA_FIXTURE), "CA", MONTH)
        us = X.dataset_from_normalized(X.read_normalized_csv(US_FIXTURE), "US", MONTH)
        with contextlib.redirect_stdout(io.StringIO()):
            cls.us_single = G.build_gauge_workbook(us, OUT, benchmark=None, overwrite=True, dated_copy=False, runs_dir=RUNS,
                                                   preclassified=True, input_paths=[US_FIXTURE])
            cls.ca_single = G.build_gauge_workbook(ca, OUT, benchmark=us, overwrite=True, dated_copy=False, runs_dir=RUNS,
                                                   preclassified=True, input_paths=[CA_FIXTURE, US_FIXTURE])
            cls.rc = CB.main(["--month", MONTH, "--from-normalized", str(CA_FIXTURE), "--us-from-normalized", str(US_FIXTURE),
                              "--out-dir", str(OUT), "--runs-dir", str(RUNS), "--overwrite"])
        cls.path = OUT / NAME
        cls.wb = openpyxl.load_workbook(cls.path)
        cls.ca_wb = openpyxl.load_workbook(cls.ca_single)
        cls.us_wb = openpyxl.load_workbook(cls.us_single)
        cls.reg = json.loads(C.run_file(RUNS, MONTH, "table_registry", CB.REGISTRY_SCOPE, "json").read_text())
        cls.tables = [t for t in cls.reg["tables"] if t["workbook"] == NAME]
        cls.ca_rows = X.read_normalized_csv(CA_FIXTURE)
        cls.us_rows = X.read_normalized_csv(US_FIXTURE)
        cls.done = True


def _t(sheet: str, role: str | None = None, title: str | None = None) -> list[dict]:
    return [t for t in _Built.tables if t["sheet"] == sheet and (role is None or t["role"] == role)
            and (title is None or t["title"] == title)]


def _one(sheet: str, role: str, title: str) -> dict:
    ts = _t(sheet, role, title)
    if len(ts) != 1:
        raise AssertionError(f"expected one {role} table {title!r} on {sheet}, got {len(ts)}")
    return ts[0]


def _cells(ws, t: dict, row: int) -> dict:
    return {h: ws.cell(row, t["first_col"] + i) for i, h in enumerate(t["columns"])}


def _rows(ws, t: dict) -> list[dict]:
    return [_cells(ws, t, r) for r in range(t["first_data_row"], t["last_data_row"] + 1)]


def _val(cell) -> float:
    return 0.0 if cell.value is None else float(cell.value)


class TestCombinedLayout(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _Built.ensure()

    # ---------------------------------------------------------------- sheets
    def test_cli_rc_and_sheet_order(self):
        self.assertEqual(_Built.rc, 0)
        names = _Built.wb.sheetnames
        nf, nm, nt = len(C.COMBINED_FIXED_SHEETS), len(C.COMBINED_MODEL_SHEETS), len(C.COMBINED_TAIL_SHEETS)
        self.assertEqual(names[:nf], list(C.COMBINED_FIXED_SHEETS))
        self.assertEqual(names[-nt:], list(C.COMBINED_TAIL_SHEETS))
        self.assertEqual(names[-(nt + nm):-nt], list(C.COMBINED_MODEL_SHEETS))
        brand_tabs = names[nf:-(nt + nm)]
        bsm = _Built.reg["workbooks"][NAME]["brand_sheet_map"]
        self.assertEqual(sorted(brand_tabs), sorted(v for k, v in bsm.items() if k != "innova"))
        self.assertEqual(bsm["innova"], "Innova")
        self.assertNotIn("Innova", brand_tabs)

    def test_brand_tabs_are_the_union_of_both_single_market_rules(self):
        reg_ca = json.loads(C.run_file(RUNS, MONTH, "table_registry", "CA", "json").read_text())
        reg_us = json.loads(C.run_file(RUNS, MONTH, "table_registry", "US", "json").read_text())
        ca_keys = set(reg_ca["workbooks"][_Built.ca_single.name]["brand_sheet_map"]) - {"innova"}
        us_keys = set(reg_us["workbooks"][_Built.us_single.name]["brand_sheet_map"]) - {"innova"}
        got = set(_Built.reg["workbooks"][NAME]["brand_sheet_map"]) - {"innova"}
        self.assertEqual(got, ca_keys | us_keys)
        self.assertNotEqual(ca_keys, us_keys, "fixture should exercise a brand qualifying in one market only")

    # ---------------------------------------------------------------- Summary: key figures (incl. the folded gauge share)
    def test_summary_table_order_titles_and_roles(self):
        summ = _t("Summary")
        want = [C.COMBINED_SUMMARY_TITLES[k] for k in ("key_figures", "brands", "subtypes", "tier_ca", "tier_us", "fuel")]
        self.assertEqual(list(C.COMBINED_SUMMARY_TITLES), ["key_figures", "brands", "subtypes", "tier_ca", "tier_us", "fuel"])
        self.assertEqual([(t["title"], t["role"]) for t in summ], want)
        tops = [t["header_row"] for t in summ]
        self.assertEqual(tops, sorted(tops))
        kf = summ[0]
        ws = _Built.wb["Summary"]
        self.assertEqual(ws.cell(1, 1).value, "Key figures")
        self.assertEqual(kf["columns"], ["Measure", "CA", "US", "Unit"])
        self.assertEqual(tuple(kf["allowed_markets"]), ("CA", "US"))
        # no separate share table any more: the Brand summary follows the Key figures directly
        texts = [v for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, str)]
        self.assertNotIn("Gauge share of the code-reader market", texts)
        self.assertNotIn(C.COMBINED_SHARE_ROW_LABEL, texts)
        self.assertFalse(any(t["title"] == "Gauge share of the code-reader market" for t in _Built.tables))
        brands = summ[1]
        self.assertLess(X_end(kf), brands["header_row"])
        for r in range(X_end(kf) + 1, brands["header_row"] - 2):
            self.assertTrue(all(ws.cell(r, c).value is None for c in range(1, 12)), r)

    def test_key_figures_values_and_formats(self):
        ws = _Built.wb["Summary"]
        kf = _one("Summary", "kpi", "Key figures")
        got = {ws.cell(r, kf["first_col"]).value: _cells(ws, kf, r) for r in range(kf["first_data_row"], kf["last_data_row"] + 1)}
        self.assertEqual(list(got), list(C.COMBINED_KEY_FIGURE_LABELS))
        self.assertEqual(len(got), 11)
        self.assertEqual(CB.KEY_FIGURE_LABELS, C.COMBINED_KEY_FIGURE_LABELS)
        for code, rows in (("CA", _Built.ca_rows), ("US", _Built.us_rows)):
            core = _core(rows)
            dev = rows[rows["gauge_class"].isin(C.GAUGE_DEVICE_CLASSES)]
            acc = rows[rows["gauge_class"].isin(C.GAUGE_ACCESSORY_CLASSES)]
            adj = rows[rows["gauge_class"].isin(C.GAUGE_ADJACENT_CLASSES)]
            cr = rows[rows["source_set"].isin(["code_reader", "both"])]           # the full code-reader export of the fixture
            want = [core["revenue_month"].sum(), core["units_month"].sum(), core["asin"].nunique(), (core["units_month"] > 0).sum(),
                    dev["revenue_month"].sum(), acc["revenue_month"].sum(), adj["revenue_month"].sum(),
                    cr["revenue_month"].sum(), cr["units_month"].sum()]
            for label, w in zip(C.COMBINED_KEY_FIGURE_LABELS, want):
                self.assertAlmostEqual(got[label][code].value, float(w), places=6, msg=(code, label))
            money_fmt = C.MARKETS[code].money_fmt
            self.assertEqual(got["Core device revenue"][code].number_format, money_fmt)
            self.assertEqual(got["Code-reader market revenue (full export)"][code].number_format, money_fmt)
            self.assertEqual(got["Core device units"][code].number_format, X.FMT_INT)
            self.assertEqual(got["Code-reader market units (full export)"][code].number_format, X.FMT_INT)
            self.assertEqual(got["# core device ASINs"][code].number_format, X.FMT_INT)
            self.assertEqual(got["# core device ASINs"][code].fill.fgColor.rgb[-6:], C.FILL_MARKET_DATA[code])
            self.assertEqual(ws.cell(kf["header_row"], kf["first_col"] + kf["columns"].index(code)).fill.fgColor.rgb[-6:],
                             C.FILL_MARKET_HEADER[code])
            for label in C.COMBINED_KEY_FIGURE_LABELS:
                self.assertEqual(bool(got[label][code].font.b), label in C.COMBINED_KEY_FIGURE_SHARE_LABELS, (code, label))
        self.assertEqual(got["Core device revenue"]["Unit"].value, "CAD / USD")
        self.assertEqual(got["Code-reader market revenue (full export)"]["Unit"].value, "CAD / USD")

    def test_key_figures_share_rows_equal_single_market_b_row(self):
        ws = _Built.wb["Summary"]
        kf = _one("Summary", "kpi", "Key figures")
        notes = [ws.cell(r, 1).value for r in range(1, kf["header_row"])]
        self.assertTrue(any(isinstance(n, str) and CB.COMBINED_SHARE_NOTE in n for n in notes), notes)
        self.assertIn("(b) = all core devices ÷ (full code-reader export + core devices found only in the gauge export)",
                      CB.COMBINED_SHARE_NOTE)
        got = {ws.cell(r, kf["first_col"]).value: _cells(ws, kf, r) for r in range(kf["first_data_row"], kf["last_data_row"] + 1)}
        s_rev_label, s_u_label = C.COMBINED_KEY_FIGURE_SHARE_LABELS
        for code, single in (("CA", _Built.ca_wb), ("US", _Built.us_wb)):
            ref = [r for r in single["Summary"].iter_rows(values_only=True) if r[0] == G.SHARE_ROW_LABELS[2]]
            self.assertEqual(len(ref), 1, code)
            n, rev, units, s_rev, s_u = ref[0][1:6]
            self.assertAlmostEqual(got[s_rev_label][code].value, s_rev, places=12)
            self.assertAlmostEqual(got[s_u_label][code].value, s_u, places=12)
            self.assertEqual(got["# core device ASINs"][code].value, n)
            self.assertAlmostEqual(got["Core device revenue"][code].value, rev, places=9)
            self.assertAlmostEqual(got["Core device units"][code].value, units, places=9)
            cr_ref = [r for r in single["Summary"].iter_rows(values_only=True) if r[0] == G.SHARE_ROW_LABELS[0]][0]
            self.assertAlmostEqual(got["Code-reader market revenue (full export)"][code].value, cr_ref[2], places=9)
            self.assertAlmostEqual(got["Code-reader market units (full export)"][code].value, cr_ref[3], places=9)
            for label in C.COMBINED_KEY_FIGURE_SHARE_LABELS:
                cell = got[label][code]
                self.assertTrue(cell.font.b, (code, label))
                self.assertEqual(cell.number_format, "0.00%")
                self.assertEqual(got[label]["Unit"].value, "share")

    # ---------------------------------------------------------------- Summary: brands
    def test_brand_summary_values_shares_and_charts(self):
        ws = _Built.wb["Summary"]
        t = _one("Summary", "summary_brands", C.COMBINED_SUMMARY_TITLES["brands"][0])
        self.assertEqual(t["columns"], list(CB.BRAND_HEADERS))
        rows = _rows(ws, t)
        ca_core, us_core = _core(_Built.ca_rows), _core(_Built.us_rows)
        keys = set(ca_core["brand_key"]) | set(us_core["brand_key"])
        self.assertEqual(len(rows), len(keys))           # fixtures: fewer than SUMMARY_TOP_BRANDS brands -> no residual
        self.assertIsNone(t["residual_row"])
        disp = {**dict(zip(us_core["brand_key"], us_core["brand_display"])), **dict(zip(ca_core["brand_key"], ca_core["brand_display"]))}
        by_label = {r["Brand"].value: r for r in rows}
        self.assertEqual(set(by_label), {disp[k] for k in keys})
        for code, core in (("CA", ca_core), ("US", us_core)):
            ccy = C.MARKETS[code].currency
            total_rev = float(core["revenue_month"].sum())
            for k in keys:
                sub = core[core["brand_key"] == k]
                r = by_label[disp[k]]
                self.assertEqual(_val(r[f"{code} # of Listings"]), len(sub), (code, k))
                self.assertAlmostEqual(_val(r[f"{code} Monthly Rev ({ccy})"]), float(sub["revenue_month"].sum()), places=6)
                self.assertAlmostEqual(_val(r[f"{code} Monthly Units"]), float(sub["units_month"].sum()), places=6)
                self.assertAlmostEqual(_val(r[f"{code} Rev Share"]), float(sub["revenue_month"].sum()) / total_rev, places=12)
                self.assertEqual(r[f"{code} Monthly Rev ({ccy})"].number_format, C.MARKETS[code].money_fmt)
                self.assertEqual(r[f"{code} Monthly Rev ({ccy})"].fill.fgColor.rgb[-6:], C.FILL_MARKET_DATA[code])
            self.assertAlmostEqual(sum(_val(r[f"{code} Rev Share"]) for r in rows), 1.0, places=9)
            tot = _cells(ws, t, t["total_row"])
            self.assertEqual(tot["Brand"].value, C.TOTAL_ROW_LABEL)
            self.assertAlmostEqual(tot[f"{code} Monthly Rev ({ccy})"].value, total_rev, places=6)
            self.assertEqual(tot[f"{code} # of Listings"].value, core["asin"].nunique())
            self.assertEqual(ws.cell(t["header_row"], t["first_col"] + t["columns"].index(f"{code} Monthly Rev ({ccy})"))
                             .fill.fgColor.rgb[-6:], C.FILL_MARKET_HEADER[code])
        # order: CA revenue desc, then US revenue desc
        ca_rev = [_val(r["CA Monthly Rev (CAD)"]) for r in rows]
        self.assertEqual(ca_rev, sorted(ca_rev, reverse=True))
        # two bar charts side by side, right of the table, at the header row
        self.assertEqual([c["type"] for c in t["charts"]], ["bar", "bar"])
        cols = []
        for c in t["charts"]:
            letters = "".join(ch for ch in c["anchor"] if ch.isalpha())
            self.assertEqual(int(c["anchor"][len(letters):]), t["header_row"])
            cols.append(openpyxl.utils.column_index_from_string(letters))
            self.assertTrue(c["has_axes"])
        self.assertGreaterEqual(min(cols), t["last_col"] + 2)
        self.assertNotEqual(cols[0], cols[1])
        self.assertEqual(len(ws._charts), 2)

    def test_brand_side_by_side_residual_and_shares(self):
        ca_core, us_core = _core(_Built.ca_rows), _core(_Built.us_rows)
        frame, residual, total = CB.brand_side_by_side(ca_core, us_core, top_n=3)
        self.assertEqual(len(frame), 3)
        n_keys = len(set(ca_core["brand_key"]) | set(us_core["brand_key"]))
        self.assertEqual(residual["brand"], C.RESIDUAL_ROW_LABEL.format(noun="brands", n=n_keys - 3))
        for m, core in (("ca", ca_core), ("us", us_core)):
            self.assertAlmostEqual(frame[f"{m}_rev"].sum() + residual[f"{m}_rev"], float(core["revenue_month"].sum()), places=6)
            self.assertAlmostEqual(frame[f"{m}_units"].sum() + residual[f"{m}_units"], float(core["units_month"].sum()), places=6)
            self.assertEqual(int(frame[f"{m}_n"].sum()) + residual[f"{m}_n"], core["asin"].nunique())
            self.assertAlmostEqual(frame[f"{m}_share"].sum() + residual[f"{m}_share"], 1.0, places=12)
            self.assertAlmostEqual(total[f"{m}_rev"], float(core["revenue_month"].sum()), places=6)
            self.assertAlmostEqual(total[f"{m}_share"], 1.0, places=12)
        # selection = top 3 by the larger of the two revenues; display order = CA revenue desc then US revenue desc
        best = {}
        for k in set(ca_core["brand_key"]) | set(us_core["brand_key"]):
            best[k] = max(float(ca_core.loc[ca_core["brand_key"] == k, "revenue_month"].sum()),
                          float(us_core.loc[us_core["brand_key"] == k, "revenue_month"].sum()))
        top3 = sorted(best, key=lambda k: -best[k])[:3]
        self.assertEqual(set(frame["brand_key"]), set(top3))
        order = list(zip(-frame["ca_rev"], -frame["us_rev"]))
        self.assertEqual(order, sorted(order))
        # a brand absent from one market shows 0 and a blank rating there
        one_sided = CB.brand_side_by_side(ca_core, us_core.iloc[0:0], top_n=25)[0]
        self.assertTrue((one_sided["us_n"] == 0).all() and (one_sided["us_rev"] == 0).all())
        self.assertTrue(one_sided["us_rating"].isna().all())

    # ---------------------------------------------------------------- Summary: sub-types, tiers, fuel
    def test_subtype_mix(self):
        ws = _Built.wb["Summary"]
        t = _one("Summary", "subtype_mix", C.COMBINED_SUMMARY_TITLES["subtypes"][0])
        self.assertEqual(t["columns"], list(CB.SUBTYPE_HEADERS))
        rows = _rows(ws, t)
        labels = [r["Sub-type"].value for r in rows]
        self.assertEqual(labels, [C.GAUGE_SUBTYPE_LABELS[c] for c in C.GAUGE_DEVICE_CLASSES] + [CB.GPS_ROW_LABEL])
        gps_row = t["last_data_row"]
        self.assertEqual(t["excluded_from_total_rows"], [gps_row])
        tot = _cells(ws, t, t["total_row"])
        for code, data in (("CA", _Built.ca_rows), ("US", _Built.us_rows)):
            ccy = C.MARKETS[code].currency
            core = _core(data)
            dev_rows = rows[:-1]
            self.assertAlmostEqual(sum(_val(r[f"{code} Monthly Rev ({ccy})"]) for r in dev_rows), float(core["revenue_month"].sum()), places=6)
            self.assertAlmostEqual(tot[f"{code} Monthly Rev ({ccy})"].value, float(core["revenue_month"].sum()), places=6)
            self.assertAlmostEqual(tot[f"{code} Monthly Units"].value, float(core["units_month"].sum()), places=6)
            self.assertEqual(tot[f"{code} # ASINs"].value, core["asin"].nunique())
            self.assertAlmostEqual(sum(_val(r[f"{code} Rev share"]) for r in dev_rows), 1.0, places=9)
            adj = data[data["gauge_class"].isin(C.GAUGE_ADJACENT_CLASSES)]
            self.assertAlmostEqual(_val(rows[-1][f"{code} Monthly Rev ({ccy})"]), float(adj["revenue_month"].sum()), places=6)
            self.assertIsNone(rows[-1][f"{code} Rev share"].value)

    def test_tier_matrices_rev_and_units_in_one_table(self):
        ws = _Built.wb["Summary"]
        for key, code, data in (("tier_ca", "CA", _Built.ca_rows), ("tier_us", "US", _Built.us_rows)):
            title, role = C.COMBINED_SUMMARY_TITLES[key]
            t = _one("Summary", role, title)
            ccy = C.MARKETS[code].currency
            self.assertEqual(t["columns"], list(CB.tier_matrix_headers(ccy)))
            self.assertEqual(tuple(t["allowed_markets"]), (code,))
            self.assertTrue(any(h.endswith(f"Rev ({ccy})") for h in t["columns"]))
            self.assertTrue(any(h.endswith("Units") for h in t["columns"]))
            rows = _rows(ws, t)
            self.assertEqual([r["Sub-type"].value for r in rows], [C.GAUGE_SUBTYPE_LABELS[c] for c in C.GAUGE_DEVICE_CLASSES])
            core = _core(data).copy()
            core["_gtier"] = [X.tier_of(p, C.GAUGE_TIERS) for p in core["price"]]
            tot = _cells(ws, t, t["total_row"])
            for cls, r in zip(C.GAUGE_DEVICE_CLASSES, rows):
                sub = core[core["gauge_class"] == cls]
                for label, _, _ in C.GAUGE_TIERS:
                    st = sub[sub["_gtier"] == label]
                    self.assertAlmostEqual(_val(r[f"{label} Rev ({ccy})"]), float(st["revenue_month"].sum()), places=6)
                    self.assertAlmostEqual(_val(r[f"{label} Units"]), float(st["units_month"].sum()), places=6)
                self.assertAlmostEqual(_val(r[f"All tiers Rev ({ccy})"]), float(sub["revenue_month"].sum()), places=6)
                self.assertEqual(r[f"All tiers Rev ({ccy})"].number_format, C.MARKETS[code].money_fmt)
            self.assertAlmostEqual(tot[f"All tiers Rev ({ccy})"].value, float(core["revenue_month"].sum()), places=6)
            self.assertAlmostEqual(tot["All tiers Units"].value, float(core["units_month"].sum()), places=6)
            self.assertAlmostEqual(sum(tot[f"{label} Rev ({ccy})"].value for label, _, _ in C.GAUGE_TIERS),
                                   float(core["revenue_month"].sum()), places=6)

    def test_fuel_split_groups_and_subtotals(self):
        ws = _Built.wb["Summary"]
        t = _one("Summary", "subtype_mix", C.COMBINED_SUMMARY_TITLES["fuel"][0])
        self.assertEqual(t["columns"], list(CB.FUEL_HEADERS))
        notes = [ws.cell(r, 1).value for r in range(t["header_row"] - 2, t["header_row"])]
        self.assertIn(G.FUEL_NOTE, notes)
        rows = _rows(ws, t)
        sub_rows = t["subtotal_rows"]
        self.assertEqual(len(sub_rows), len(C.FEATURE_FUEL_SCOPE))
        groups: list[tuple[str, dict, list[dict]]] = []
        for r_idx, r in zip(range(t["first_data_row"], t["last_data_row"] + 1), rows):
            if r_idx in sub_rows:
                groups.append((r["Fuel / Sub-type"].value, r, []))
                self.assertTrue(r["Fuel / Sub-type"].font.b)
            else:
                groups[-1][2].append(r)
        self.assertEqual([g[0] for g in groups], [CB.FUEL_SUBTOTAL_LABEL.format(fuel=f) for f in C.FEATURE_FUEL_SCOPE])
        tot = _cells(ws, t, t["total_row"])
        data = {"CA": (_core(_Built.ca_rows), _fuel(_Built.ca_rows)), "US": (_core(_Built.us_rows), _fuel(_Built.us_rows))}
        for f, (label, sub, members) in zip(C.FEATURE_FUEL_SCOPE, groups):
            present = [c for c in C.GAUGE_DEVICE_CLASSES
                       if any(((core["gauge_class"] == c) & (fuel == f)).any() for core, fuel in data.values())]
            self.assertEqual([m["Fuel / Sub-type"].value for m in members], [C.GAUGE_SUBTYPE_LABELS[c] for c in present])
            for code, (core, fuel) in data.items():
                ccy = C.MARKETS[code].currency
                in_f = core[fuel == f]
                self.assertEqual(_val(sub[f"{code} # ASINs"]), len(in_f))
                self.assertAlmostEqual(_val(sub[f"{code} Monthly Rev ({ccy})"]), float(in_f["revenue_month"].sum()), places=6)
                self.assertAlmostEqual(sum(_val(m[f"{code} Monthly Units"]) for m in members), _val(sub[f"{code} Monthly Units"]), places=6)
        for code, (core, _) in data.items():
            ccy = C.MARKETS[code].currency
            self.assertAlmostEqual(sum(_val(g[1][f"{code} Monthly Rev ({ccy})"]) for g in groups), float(core["revenue_month"].sum()), places=6)
            self.assertAlmostEqual(tot[f"{code} Monthly Rev ({ccy})"].value, float(core["revenue_month"].sum()), places=6)
            self.assertEqual(tot[f"{code} # ASINs"].value, len(core))
        # US fixture: one diesel-capable tuner, one gas Lufi gauge
        dsub = groups[list(C.FEATURE_FUEL_SCOPE).index("diesel-capable")][1]
        self.assertEqual(dsub["US Monthly Units"].value, 12)

    # ---------------------------------------------------------------- Top 50 CA / US
    def test_top50_sheets_equal_single_market_content(self):
        for sheet, single in (("Top 50 CA", _Built.ca_wb), ("Top 50 US", _Built.us_wb)):
            a = [list(r) for r in _Built.wb[sheet].iter_rows(values_only=True)]
            b = [list(r) for r in single["Top 50"].iter_rows(values_only=True)]
            self.assertEqual(a, b, sheet)
            fa = [[c.number_format for c in r] for r in _Built.wb[sheet].iter_rows()]
            fb = [[c.number_format for c in r] for r in single["Top 50"].iter_rows()]
            self.assertEqual(fa, fb, sheet)
            self.assertEqual([t["role"] for t in _t(sheet)], ["top_by_revenue", "top_by_units"])
            code = sheet[-2:]
            for t in _t(sheet):
                self.assertEqual(tuple(t["allowed_markets"]), (code,))
                self.assertEqual(t["market"], code)

    # ---------------------------------------------------------------- brand tabs
    def test_every_brand_tab_has_four_titled_tables_with_totals(self):
        bsm = _Built.reg["workbooks"][NAME]["brand_sheet_map"]
        data = {"CA": _core(_Built.ca_rows), "US": _core(_Built.us_rows)}
        for key, sheet in bsm.items():
            if key == "innova":
                continue
            ts = _t(sheet)
            kpi = [t for t in ts if t["role"] == "kpi"]
            self.assertEqual(len(kpi), 1, sheet)
            ranked = [t for t in ts if t["role"] != "kpi"]
            self.assertEqual([t["title"] for t in ranked], list(C.COMBINED_BRAND_TABLE_TITLES), sheet)
            self.assertEqual([t["role"] for t in ranked], ["brand_tab_revenue", "brand_tab_units"] * 2)
            self.assertEqual(ranked[0]["header_row"], C.BRAND_TAB_RESERVED_ROWS + 2)
            ws = _Built.wb[sheet]
            self.assertEqual((ws.cell(2, 2).value, ws.cell(2, 3).value), ("CA", "US"))
            self.assertEqual((kpi[0]["first_data_row"], kpi[0]["last_data_row"]), (3, 6))
            for t in ranked:
                code = t["title"][:2]
                ccy = C.MARKETS[code].currency
                self.assertEqual(tuple(t["allowed_markets"]), (code,))
                self.assertIsNotNone(t["total_row"])
                rows = data[code][data[code]["brand_key"] == key]
                tot = _cells(ws, t, t["total_row"])
                self.assertEqual(tot["Product Name"].value, C.TOTAL_ROW_LABEL)
                self.assertAlmostEqual(_val(tot[f"Monthly Rev ({ccy})"]), float(rows["revenue_month"].sum()), places=6)
                self.assertAlmostEqual(_val(tot["Monthly Units"]), float(rows["units_month"].sum()), places=6)
                body = _rows(ws, t)
                if len(rows):
                    self.assertEqual(sorted(r["ASIN"].value for r in body), sorted(rows["asin"]))
                else:
                    self.assertEqual([r["Product Name"].value for r in body], [CB.NO_LISTINGS_LABEL])
                hdr = ws.cell(t["header_row"], t["first_col"]).fill.fgColor.rgb[-6:]
                self.assertEqual(hdr, C.FILL_MARKET_HEADER[code])
            col = {"CA": 2, "US": 3}
            for code, c in col.items():
                rows = data[code][data[code]["brand_key"] == key]
                self.assertAlmostEqual(_val(ws.cell(3, c)), float(rows["revenue_month"].sum()), places=6)
                self.assertAlmostEqual(_val(ws.cell(4, c)), float(rows["units_month"].sum()), places=6)
                self.assertEqual(_val(ws.cell(5, c)), len(rows))
                self.assertAlmostEqual(_val(ws.cell(6, c)), float(rows["revenue_month"].sum()) / float(data[code]["revenue_month"].sum()),
                                       places=12)
                self.assertEqual(ws.cell(3, c).number_format, C.MARKETS[code].money_fmt)

    def test_brand_missing_in_one_market_shows_dashes(self):
        """Edge Products dropped from the US fixture: CA-only brand. Its US cells show '-' (Brand summary five cells, brand
        tab KPI block, Total revenue/units of the US ranking tables); CA cells, residual and Total rows stay numeric."""
        d = OUT / "one_sided"
        if d.exists():
            shutil.rmtree(d)
        us_rows = X.read_normalized_csv(US_FIXTURE)
        us_rows = us_rows[us_rows["brand_key"] != "edge products"]
        ca = X.dataset_from_normalized(X.read_normalized_csv(CA_FIXTURE), "CA", MONTH)
        us = X.dataset_from_normalized(us_rows, "US", MONTH)
        with contextlib.redirect_stdout(io.StringIO()), mock.patch.object(C, "SUMMARY_TOP_BRANDS", 5):
            p = CB.build_combined_gauge_workbook(ca, us, d, overwrite=True, dated_copy=False, runs_dir=d / "runs",
                                                 input_paths=[CA_FIXTURE, US_FIXTURE])
        reg = json.loads(C.run_file(d / "runs", MONTH, "table_registry", CB.REGISTRY_SCOPE, "json").read_text())
        wb = openpyxl.load_workbook(p)
        # Brand summary: top 5 by max(CA, US) revenue keeps Edge (CA only); 2 brands go to the residual
        ws = wb["Summary"]
        t = [t for t in reg["tables"] if t["sheet"] == "Summary" and t["role"] == "summary_brands"][0]
        rows = {r["Brand"].value: (row, r) for row, r in zip(range(t["first_data_row"], t["last_data_row"] + 1), _rows(ws, t))}
        self.assertIn("Edge Products", rows)
        edge_row, edge = rows["Edge Products"]
        self.assertEqual(t["dash_rows"], [[edge_row, "US"]])
        for h in ("# of Listings", "Monthly Rev (USD)", "Monthly Units", "Rev Share", "Avg Rating"):
            cell = edge[f"US {h}"]
            self.assertEqual(cell.value, CB.DASH, h)
            self.assertEqual(cell.data_type, "s", h)
            self.assertEqual(cell.alignment.horizontal, "right", h)
            self.assertEqual(cell.fill.fgColor.rgb[-6:], C.FILL_MARKET_DATA["US"], h)
        self.assertEqual(edge["CA # of Listings"].value, 2)
        self.assertAlmostEqual(edge["CA Monthly Rev (CAD)"].value, 2644.0, places=6)
        for brand, (_, r) in rows.items():
            if brand != "Edge Products":
                for h in t["columns"][1:]:
                    self.assertNotEqual(r[h].value, CB.DASH, (brand, h))
        for special in (t["residual_row"], t["total_row"]):
            self.assertIsNotNone(special)
            cells = _cells(ws, t, special)
            for h in t["columns"][1:]:
                self.assertNotEqual(cells[h].value, CB.DASH, h)
            self.assertIsInstance(cells["US Monthly Rev (USD)"].value, (int, float))
            self.assertIsInstance(cells["CA # of Listings"].value, (int, float))
        self.assertEqual(ws.cell(t["residual_row"], 1).value, C.RESIDUAL_ROW_LABEL.format(noun="brands", n=2))
        # brand tab: KPI block C3:C6 dashes, US ranking tables keep the text row and show '-' for Total revenue/units
        sheet = reg["workbooks"][p.name]["brand_sheet_map"]["edge products"]
        ws = wb[sheet]
        kpi = [t for t in reg["tables"] if t["sheet"] == sheet and t["role"] == "kpi"][0]
        self.assertEqual(kpi["dash_rows"], [[r, "US"] for r in range(3, 7)])
        for r in range(3, 7):
            self.assertEqual(ws.cell(r, 3).value, CB.DASH, r)
            self.assertEqual(ws.cell(r, 3).alignment.horizontal, "right")
            self.assertIsInstance(ws.cell(r, 2).value, (int, float), r)
        ts = [t for t in reg["tables"] if t["sheet"] == sheet and t["title"].startswith("US — ")]
        self.assertEqual(len(ts), 2)
        for t in ts:
            self.assertEqual(t["first_data_row"], t["last_data_row"])
            self.assertEqual(ws.cell(t["first_data_row"], 1).value, CB.NO_LISTINGS_LABEL)
            self.assertEqual(t["placeholder_rows"], [t["first_data_row"]])
            self.assertEqual(t["dash_rows"], [[t["total_row"], "US"]])
            tot = _cells(ws, t, t["total_row"])
            self.assertEqual(tot["Product Name"].value, C.TOTAL_ROW_LABEL)
            self.assertEqual((tot["Monthly Rev (USD)"].value, tot["Monthly Units"].value), (CB.DASH, CB.DASH))
            for h in t["columns"][1:]:
                if h not in ("Monthly Rev (USD)", "Monthly Units"):
                    self.assertIsNone(tot[h].value, h)
        for t in [t for t in reg["tables"] if t["sheet"] == sheet and t["title"].startswith("CA — ")]:
            self.assertEqual(t["dash_rows"], [])
            self.assertAlmostEqual(_cells(ws, t, t["total_row"])["Monthly Rev (CAD)"].value, 2644.0, places=6)
        # nothing else carries a dash: only brand-summary rows, brand-tab KPI blocks and brand-tab Totals
        for t in reg["tables"]:
            if t["dash_rows"]:
                self.assertTrue(t["role"] in ("summary_brands", "kpi", "brand_tab_revenue", "brand_tab_units"), t["title"])
                self.assertTrue(t["sheet"] == "Summary" or t["sheet"] in reg["workbooks"][p.name]["brand_sheet_map"].values())

    def test_listed_brand_with_zero_revenue_keeps_numbers(self):
        """AEM is listed in both fixture markets with zero revenue: numeric 0, never '-' (main build has no dash rows)."""
        ws = _Built.wb["Summary"]
        t = _one("Summary", "summary_brands", C.COMBINED_SUMMARY_TITLES["brands"][0])
        self.assertEqual(t["dash_rows"], [])
        aem = [r for r in _rows(ws, t) if r["Brand"].value == "AEM"][0]
        for code in ("CA", "US"):
            ccy = C.MARKETS[code].currency
            self.assertEqual(aem[f"{code} # of Listings"].value, 1)
            self.assertEqual(aem[f"{code} Monthly Rev ({ccy})"].value, 0)
            self.assertIsInstance(aem[f"{code} Monthly Rev ({ccy})"].value, (int, float))
        for t in _Built.tables:
            self.assertEqual(t["dash_rows"], [], (t["sheet"], t["title"]))

    # ---------------------------------------------------------------- Innova / model sheets / Same-ASIN
    def test_innova_counts_and_ca_hardware_list(self):
        ws = _Built.wb["Innova"]
        for cell, code, rows in ((CB.INNOVA_CA_COUNT_CELL, "CA", _Built.ca_rows), (CB.INNOVA_US_COUNT_CELL, "US", _Built.us_rows)):
            n = int(((rows["brand_key"] == "innova") & rows["gauge_class"].isin(C.GAUGE_DEVICE_CLASSES)).sum())
            self.assertIsInstance(ws[cell].value, int)
            self.assertEqual(ws[cell].value, n)
            self.assertIn(f"{code} dataset", ws.cell(ws[cell].row, 1).value)
        self.assertEqual((CB.INNOVA_CA_COUNT_CELL, CB.INNOVA_US_COUNT_CELL), ("B3", "B4"))
        hw = _t("Innova", "modelb_top")
        self.assertEqual(len(hw), 1)
        ref = [list(r) for r in _Built.ca_wb["Innova"].iter_rows(min_row=hw[0]["header_row"] - 1, values_only=True)]
        got = [list(r) for r in ws.iter_rows(min_row=hw[0]["header_row"] - 1, values_only=True)]
        self.assertEqual(got, ref)
        texts = [c for row in ws.iter_rows(values_only=True) for c in row if isinstance(c, str)]
        self.assertTrue(any("out of scope" in s and "US" in s for s in texts))
        self.assertEqual([t["role"] for t in _t("Innova") if t["role"] == "innova"], ["innova", "innova"])

    def test_model_and_same_asin_sheets_identical_to_ca_workbook(self):
        for sheet in C.COMBINED_MODEL_SHEETS + ("US vs CA Same-ASIN",):
            a = [list(r) for r in _Built.wb[sheet].iter_rows(values_only=True)]
            b = [list(r) for r in _Built.ca_wb[sheet].iter_rows(values_only=True)]
            self.assertEqual(a, b, sheet)
            self.assertEqual(len(_Built.wb[sheet]._charts), len(_Built.ca_wb[sheet]._charts), sheet)
        self.assertEqual(_Built.wb["US vs CA Same-ASIN"][G.SAME_ASIN_COUNT_CELL].value,
                         _Built.ca_wb["US vs CA Same-ASIN"][G.SAME_ASIN_COUNT_CELL].value)

    # ---------------------------------------------------------------- stacked sheets
    def test_all_products_excluded_audit_two_tables(self):
        for sheet, role in (("All Products", "all_rows"), ("Excluded", "excluded"), ("Dedupe & Classification Audit", "dedupe_audit")):
            ts = _t(sheet)
            self.assertEqual([(t["title"], t["role"]) for t in ts], [(f"{sheet} — CA", role), (f"{sheet} — US", role)], sheet)
            for t, code in zip(ts, ("CA", "US")):
                self.assertEqual(tuple(t["allowed_markets"]), (code,))
                ref = [e for e in json.loads(C.run_file(RUNS, MONTH, "table_registry", code, "json").read_text())["tables"]
                       if e["sheet"] == sheet and e["role"] == role][0]
                self.assertEqual(t["columns"], ref["columns"], (sheet, code))
        for t, rows in zip(_t("All Products"), (_Built.ca_rows, _Built.us_rows)):
            self.assertEqual(t["last_data_row"] - t["first_data_row"] + 1, rows["asin"].nunique())
        for t, rows in zip(_t("Excluded"), (_Built.ca_rows, _Built.us_rows)):
            n = rows["gauge_class"].isin(C.GAUGE_EXCLUDED_CLASSES + ("ambiguous",)).sum()
            self.assertEqual(t["last_data_row"] - t["first_data_row"] + 1, n)
        ws = _Built.wb["All Products"]
        ca_t, us_t = _t("All Products")
        link_c = us_t["first_col"] + us_t["columns"].index("Link")
        self.assertIn("amazon.com", ws.cell(us_t["first_data_row"], link_c).value)
        self.assertIn("amazon.ca", ws.cell(ca_t["first_data_row"], link_c).value)

    # ---------------------------------------------------------------- registry / formats / formulas
    def test_registry_shape_and_markets(self):
        reg = _Built.reg
        self.assertEqual((reg["market"], reg["month"]), (CB.REGISTRY_SCOPE, MONTH))
        self.assertEqual(reg["workbooks"][NAME]["sheets"], _Built.wb.sheetnames)
        for k in ("generated_at", "pipeline_version", "workbooks", "tables"):
            self.assertIn(k, reg)
        bsm = reg["workbooks"][NAME]["brand_sheet_map"]
        for t in _Built.tables:
            for k in ("workbook", "sheet", "role", "title", "header_row", "first_data_row", "last_data_row", "total_row",
                      "residual_row", "first_col", "last_col", "columns", "dataset_filter", "allowed_markets", "charts",
                      "brand_sheet_map", "market", "column_markets", "subtotal_rows", "excluded_from_total_rows", "placeholder_rows",
                      "dash_rows"):
                self.assertIn(k, t)
            self.assertIn(t["role"], C.TABLE_ROLES)
            self.assertEqual(t["brand_sheet_map"], bsm)
            self.assertEqual(len(t["column_markets"]), len(t["columns"]))
            self.assertTrue(set(t["allowed_markets"]) <= {"CA", "US"})
            if t["header_row"]:
                ws = _Built.wb[t["sheet"]]
                got = [ws.cell(t["header_row"], t["first_col"] + i).value for i in range(len(t["columns"]))]
                self.assertEqual(got, t["columns"], (t["sheet"], t["title"]))
        for t in _t("Summary"):
            if t["title"] in (C.COMBINED_SUMMARY_TITLES["tier_ca"][0], C.COMBINED_SUMMARY_TITLES["tier_us"][0]):
                continue
            self.assertEqual(tuple(t["allowed_markets"]), ("CA", "US"), t["title"])
            self.assertEqual({m for m in t["column_markets"] if m}, {"CA", "US"}, t["title"])
            for h, m in zip(t["columns"], t["column_markets"]):
                if h.startswith(("CA ", "US ")) or h in ("CA", "US"):
                    self.assertEqual(m, h[:2], h)

    def test_money_formats_follow_column_market(self):
        n = {"CA": 0, "US": 0}
        for t in _Built.tables:
            if not t["header_row"]:
                continue
            ws = _Built.wb[t["sheet"]]
            rows = list(range(t["first_data_row"], t["last_data_row"] + 1)) + [r for r in (t["residual_row"], t["total_row"]) if r]
            for i, (h, m) in enumerate(zip(t["columns"], t["column_markets"])):
                for r in rows:
                    fmt = ws.cell(r, t["first_col"] + i).number_format or ""
                    if "$" not in fmt:
                        continue
                    self.assertIsNotNone(m, (t["sheet"], t["title"], h))
                    ccy = C.MARKETS[m].currency
                    if h == m:      # Key figures: columns 'CA' / 'US', the row's Unit cell names the currencies
                        self.assertIn(ccy, ws.cell(r, t["first_col"] + t["columns"].index("Unit")).value, (t["sheet"], r))
                    else:
                        self.assertIn(f"({ccy})", h, (t["sheet"], t["title"]))
                    if m == "CA":
                        self.assertIn('"CA$"', fmt, (t["sheet"], h))
                    else:
                        self.assertNotIn("CA$", fmt, (t["sheet"], h))
                    n[m] += 1
        self.assertGreater(n["CA"], 0)
        self.assertGreater(n["US"], 0)
        # every $-format cell on the workbook belongs to a registered table column
        for ws in _Built.wb.worksheets:
            owned = set()
            for t in [t for t in _Built.tables if t["sheet"] == ws.title]:
                last = max(x for x in (t["last_data_row"], t["total_row"] or 0, t["residual_row"] or 0))
                owned |= {(r, c) for r in range(t["first_data_row"], last + 1) for c in range(t["first_col"], t["last_col"] + 1)}
            for row in ws.iter_rows():
                for c in row:
                    if "$" in (c.number_format or ""):
                        self.assertIn((c.row, c.column), owned, f"{ws.title}!{c.coordinate}")

    def test_hyperlink_only_formulas_with_market_domain(self):
        n = 0
        for ws in _Built.wb.worksheets:
            tables = [t for t in _Built.tables if t["sheet"] == ws.title]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    m = next((rx.match(str(cell.value)) for rx in C.ALLOWED_FORMULA_RES if rx.match(str(cell.value))), None)
                    self.assertIsNotNone(m, f"{ws.title}!{cell.coordinate}: {cell.value}")
                    owner = [t for t in tables if t["first_data_row"] <= cell.row <= t["last_data_row"]
                             and t["first_col"] <= cell.column <= t["last_col"]]
                    self.assertEqual(len(owner), 1, f"{ws.title}!{cell.coordinate}")
                    t = owner[0]
                    self.assertEqual(ws.cell(cell.row, t["first_col"] + t["columns"].index("ASIN")).value, m.group("asin"))
                    mk = t["column_markets"][cell.column - t["first_col"]]
                    self.assertEqual(m.group("tld"), "ca" if mk == "CA" else "com", f"{ws.title}!{cell.coordinate}")
                    n += 1
        self.assertGreater(n, 0)

    def test_no_cross_currency_ratio_and_no_error_tokens(self):
        for t in _Built.tables:
            if t["sheet"] == "US vs CA Same-ASIN":
                continue          # units ratio only (identical to the CA workbook)
            for h in t["columns"]:
                self.assertNotIn("ratio", h.lower(), (t["sheet"], h))
        for ws in _Built.wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str):
                        self.assertFalse(any(tok in v for tok in ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!")), (ws.title, v))
                    if isinstance(v, float):
                        self.assertTrue(math.isfinite(v), (ws.title, v))

    def test_metadata_dual_values(self):
        ws = _Built.wb["Metadata"]
        vals = {r[0]: r[1] for r in ws.iter_rows(min_row=3, values_only=True) if r and r[0]}
        for k in C.METADATA_REQUIRED_KEYS + ("Code-reader market totals",):
            self.assertIn(k, vals)
        self.assertEqual(vals["Marketplace"], "amazon.ca (CA) + amazon.com (US)")
        self.assertEqual(vals["Currency"], "CAD and USD, never mixed in one figure")
        self.assertIn("US figures are raw Helium 10 estimates without the actuals overlay", vals["Scope note"])
        for k in ("Raw files", "Dedupe summary", "Gauge overlap", "Code-reader market totals"):
            self.assertIn("CA:", vals[k], k)
            self.assertIn("US:", vals[k], k)
        self.assertIn("US", vals["Type coverage"])
        self.assertIn("not typed", vals["Type coverage"])
        sm = [c for row in _Built.wb["Source & Method"].iter_rows(values_only=True) for c in row if isinstance(c, str)]
        self.assertTrue(any("side by side" in s.lower() for s in sm))
        ca_sm = [c for row in _Built.ca_wb["Source & Method"].iter_rows(min_row=2, values_only=True) for c in row if isinstance(c, str)]
        for para in ca_sm:
            self.assertIn(para, sm)

    def test_manifest_merged(self):
        man = json.loads((OUT / C.manifest_name(MONTH)).read_text())
        self.assertEqual(man["outputs"][NAME]["sha256"], _sha(_Built.path))
        self.assertIn(_Built.ca_single.name, man["outputs"])          # entries of other builders are kept
        self.assertIn(_Built.us_single.name, man["outputs"])
        for p in (CA_FIXTURE, US_FIXTURE, G.APP_FEATURE_MATRIX_DEFAULT):
            self.assertIn(str(p.resolve()), man["inputs"])
        self.assertFalse(any("type_review" in k for k in man["inputs"]))

    def test_source_hygiene(self):
        src = (C.PACKAGE_ROOT / "build_combined_gauge_report.py").read_text()
        for bad in ("insert_rows", "delete_rows", "load_workbook"):
            self.assertNotIn(bad, src)

    def test_cli_guards(self):
        with self.assertRaises(SystemExit), contextlib.redirect_stderr(io.StringIO()):
            CB.main(["--month", "202613", "--from-normalized", str(CA_FIXTURE), "--us-from-normalized", str(US_FIXTURE),
                     "--out-dir", str(OUT)])
        with self.assertRaises(SystemExit), contextlib.redirect_stderr(io.StringIO()):
            CB.main(["--month", MONTH, "--from-normalized", str(CA_FIXTURE), "--out-dir", str(OUT), "--runs-dir", str(RUNS)])
        with self.assertRaises(SystemExit), contextlib.redirect_stderr(io.StringIO()):
            CB.main(["--month", MONTH, "--from-normalized", str(CA_FIXTURE), "--us-from-normalized", str(US_FIXTURE),
                     "--ca-cr-raw-dir", str(OUT), "--out-dir", str(OUT), "--runs-dir", str(RUNS)])
        with self.assertRaises(ValueError):
            CB.main(["--month", MONTH, "--from-normalized", str(CA_FIXTURE), "--us-from-normalized", str(US_FIXTURE),
                     "--out-dir", str(C.NEW_PRODUCT_DIR / "x"), "--runs-dir", str(RUNS)])
        with self.assertRaises(FileExistsError), contextlib.redirect_stdout(io.StringIO()):
            CB.main(["--month", MONTH, "--from-normalized", str(CA_FIXTURE), "--us-from-normalized", str(US_FIXTURE),
                     "--out-dir", str(OUT), "--runs-dir", str(RUNS)])
        ca = X.dataset_from_normalized(X.read_normalized_csv(CA_FIXTURE), "CA", MONTH)
        with self.assertRaises(ValueError):
            CB.build_combined_gauge_workbook(ca, ca, OUT / "never", overwrite=True, dated_copy=False, runs_dir=OUT / "never_runs")

    def test_preview_and_excerpt(self):
        html, md = OUT / "preview.html", OUT / "excerpt.md"
        with contextlib.redirect_stdout(io.StringIO()):
            rc = RP.main(["--workbook", str(_Built.path), "--out", str(html), "--excerpt", str(md)])
        self.assertEqual(rc, 0)
        text = html.read_text()
        for s in ("Summary", "Top 50 CA", "Top 50 US", "Metadata", C.FILL_MARKET_HEADER["US"]):
            self.assertIn(s, text)
        ex = md.read_text()
        self.assertIn("Core device revenue", ex)
        self.assertIn("CA$", ex)


class TestMarketKpiTable(unittest.TestCase):
    def test_per_row_kind_and_per_column_market(self):
        book = X.Book("kpi_matrix.xlsx", C.MARKETS["CA"])
        ws = book.sheet("Summary")
        items = [("Revenue", {"CA": (1234.5, "money"), "US": (99.0, "money")}, "CAD / USD"),
                 ("# ASINs", {"CA": (3, "int"), "US": (4, "int")}, "count")]
        tr = X.write_market_kpi_table(ws, 1, "Key figures", items, ("CA", "US"), header_fills=C.FILL_MARKET_HEADER,
                                      data_fills=C.FILL_MARKET_DATA)
        self.assertEqual(tr.columns, ["Measure", "CA", "US", "Unit"])
        self.assertEqual(tr.allowed_markets, ("CA", "US"))
        r = tr.first_data_row
        self.assertEqual(ws.cell(r, 2).number_format, C.MARKETS["CA"].money_fmt)
        self.assertEqual(ws.cell(r, 3).number_format, C.MARKETS["US"].money_fmt)
        self.assertEqual(ws.cell(r + 1, 2).number_format, X.FMT_INT)
        self.assertEqual(ws.cell(r + 1, 4).value, "count")
        self.assertEqual(ws.cell(tr.header_row, 3).fill.fgColor.rgb[-6:], C.FILL_MARKET_HEADER["US"])
        with self.assertRaises(ValueError):
            X.write_market_kpi_table(ws, 10, "bad", [("x", {"CA": (1, "money")}, "CAD")], ("CA", "US"))
        with self.assertRaises(ValueError):
            X.write_market_kpi_table(ws, 20, "bad", [("x", {"CA": (1, "text"), "US": (1, "int")}, "")], ("CA", "US"))


def X_end(t: dict) -> int:
    return max(x for x in (t["last_data_row"], t["header_row"] or 0, t["total_row"] or 0, t["residual_row"] or 0))


if __name__ == "__main__":
    unittest.main()
