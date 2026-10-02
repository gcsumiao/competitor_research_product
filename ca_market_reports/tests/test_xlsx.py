"""Workbook engine + builder tests (track T4: ca-workbooks).

Builds, from the normalized fixtures only (the loader/classifier are not needed):
  * both CA code-reader workbooks (report + analysis)
  * a US gauge workbook
  * a CA gauge workbook with the US fixture as benchmark
into tmp/ca_scratch/test_out/ and checks the formula/style/registry/manifest contracts of ca_common.

Run: ca_market_reports/run.sh -m unittest ca_market_reports.tests.test_xlsx -v
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import json
import math
import re
import shutil
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd

from ca_market_reports import ca_common as C
from ca_market_reports import ca_xlsx_style as X
from ca_market_reports import build_ca_code_reader_report as CR
from ca_market_reports import build_gauge_report as G
from ca_market_reports import render_preview as RP

MONTH = "202609"
OUT = C.SCRATCH_DIR / "test_out"
RUNS = OUT / "runs"
CA_FIXTURE = C.FIXTURES_DIR / "normalized_rows.csv"
US_FIXTURE = C.FIXTURES_DIR / "normalized_rows_us.csv"
OWN_MODULES = ("ca_xlsx_style.py", "build_ca_code_reader_report.py", "build_gauge_report.py", "render_preview.py")
BUILD_MODULES = ("ca_xlsx_style.py", "build_ca_code_reader_report.py", "build_gauge_report.py")
CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
AXIS_TAGS = ("catAx", "valAx", "dateAx", "serAx")


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


class _Built:
    """Builds every workbook once for the whole module."""
    done = False

    @classmethod
    def ensure(cls):
        if cls.done:
            return
        if OUT.exists():
            shutil.rmtree(OUT)
        OUT.mkdir(parents=True)
        ca = X.dataset_from_normalized(X.read_normalized_csv(CA_FIXTURE), "CA", MONTH)
        us = X.dataset_from_normalized(X.read_normalized_csv(US_FIXTURE), "US", MONTH)
        with contextlib.redirect_stdout(io.StringIO()):
            cls.cr_paths = CR.build_code_reader_workbooks(ca, OUT, overwrite=True, dated_copy=False, runs_dir=RUNS,
                                                          input_paths=[CA_FIXTURE])
            cls.us_gauge = G.build_gauge_workbook(us, OUT, benchmark=None, overwrite=True, dated_copy=False, runs_dir=RUNS,
                                                  preclassified=True, input_paths=[US_FIXTURE])
            cls.ca_gauge = G.build_gauge_workbook(ca, OUT, benchmark=us, overwrite=True, dated_copy=False, runs_dir=RUNS,
                                                  preclassified=True, input_paths=[CA_FIXTURE, US_FIXTURE])
        cls.ca_rows = X.read_normalized_csv(CA_FIXTURE)
        cls.us_rows = X.read_normalized_csv(US_FIXTURE)
        cls.reg = {m: json.loads(C.run_file(RUNS, MONTH, "table_registry", m, "json").read_text()) for m in ("CA", "US")}
        cls.done = True


def _all_paths():
    return list(_Built.cr_paths) + [_Built.ca_gauge, _Built.us_gauge]


def _tables(market: str, workbook: str):
    return [t for t in _Built.reg[market]["tables"] if t["workbook"] == workbook]


def _market_of(path: Path) -> str:
    return path.name.split("_", 1)[0]


class TestBuildOutputs(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _Built.ensure()

    # ---------------------------------------------------------------- names / sheets
    def test_output_names_and_sheet_order(self):
        names = [p.name for p in _Built.cr_paths]
        self.assertEqual(names, [C.cr_report_name("CA", MONTH), C.cr_analysis_name("CA", MONTH)])
        self.assertEqual(_Built.ca_gauge.name, C.gauge_report_name("CA", MONTH))
        self.assertEqual(_Built.us_gauge.name, C.gauge_report_name("US", MONTH))
        rep = openpyxl.load_workbook(_Built.cr_paths[0]).sheetnames
        self.assertEqual(rep[:3], list(C.CR_REPORT_FIXED_SHEETS))
        self.assertEqual(rep[-2:], list(C.CR_REPORT_TAIL_SHEETS))
        self.assertEqual(len(rep), 3 + min(C.CR_REPORT_BRAND_TABS, 10) + 2)   # fixture has exactly 10 non-Innova brands
        ana = openpyxl.load_workbook(_Built.cr_paths[1]).sheetnames
        self.assertEqual(ana, list(C.CR_ANALYSIS_SHEETS))
        cag = openpyxl.load_workbook(_Built.ca_gauge).sheetnames
        self.assertEqual(cag[:4], list(C.GAUGE_FIXED_SHEETS))
        self.assertEqual(cag[-5:], list(C.GAUGE_TAIL_SHEETS))
        tail = cag[-(5 + len(C.GAUGE_MODEL_SHEETS) + len(C.GAUGE_BENCHMARK_SHEETS)):-5]
        self.assertEqual(tail, list(C.GAUGE_MODEL_SHEETS + C.GAUGE_BENCHMARK_SHEETS))
        usg = openpyxl.load_workbook(_Built.us_gauge).sheetnames
        self.assertEqual(usg[:4], list(C.GAUGE_FIXED_SHEETS))
        self.assertEqual(usg[-5:], list(C.GAUGE_TAIL_SHEETS))
        for s in C.GAUGE_MODEL_SHEETS + C.GAUGE_BENCHMARK_SHEETS:
            self.assertNotIn(s, usg)

    def test_sheet_names_valid_and_unique(self):
        for p in _all_paths():
            names = openpyxl.load_workbook(p).sheetnames
            self.assertTrue(all(len(n) <= 31 for n in names), names)
            self.assertEqual(len({n.lower() for n in names}), len(names), names)
            self.assertFalse(any(re.search(r"[\[\]:*?/\\]", n) for n in names), names)

    # ---------------------------------------------------------------- charts
    def test_charts_registered_and_axes_kept(self):
        for p in _all_paths():
            tables = _tables(_market_of(p), p.name)
            charts = [c for t in tables for c in t["charts"]]
            with zipfile.ZipFile(p) as z:
                xmls = sorted(n for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))
                bodies = [z.read(n).decode("utf-8") for n in xmls]
            self.assertEqual(len(xmls), len(charts), f"{p.name}: registered {charts}")
            if p != _Built.cr_paths[1]:   # the Analysis workbook specifies no charts
                self.assertGreater(len(charts), 0, p.name)
            if p in (_Built.cr_paths[0], _Built.ca_gauge, _Built.us_gauge):
                summary = [c["type"] for t in tables if t["sheet"] == "Summary" for c in t["charts"]]
                self.assertEqual(sorted(summary), ["bar", "pie"], p.name)
            if p == _Built.ca_gauge:
                ladder = [c["type"] for t in tables if t["sheet"] == "Price Ladder (Model A)" for c in t["charts"]]
                self.assertEqual(ladder, ["scatter"])
            # namespace-aware: openpyxl writes the chart namespace as the default namespace (<delete val="0"/> == <c:delete val="0"/>)
            trees = [ET.fromstring(b) for b in bodies]
            with_axes = [t for t in trees if any(t.find(f".//{{{CHART_NS}}}{ax}") is not None for ax in AXIS_TAGS)]
            self.assertEqual(len(with_axes), sum(1 for c in charts if c["has_axes"]), p.name)
            for t in with_axes:
                for ax_tag in AXIS_TAGS:
                    for ax in t.iter(f"{{{CHART_NS}}}{ax_tag}"):
                        d = ax.find(f"{{{CHART_NS}}}delete")
                        self.assertIsNotNone(d, f"{p.name}: {ax_tag} without <c:delete>")
                        self.assertEqual(d.get("val"), "0", p.name)
                        self.assertEqual(ax.find(f"{{{CHART_NS}}}tickLblPos").get("val"), "nextTo", p.name)
            for t in trees:
                self.assertEqual(t.find(f"{{{CHART_NS}}}chart/{{{CHART_NS}}}plotVisOnly").get("val"), "0", p.name)
            for b, t in zip(bodies, trees):
                kinds = {el.tag.split("}")[1] for el in t.iter() if el.tag.endswith(("barChart", "pieChart"))}
                if kinds:   # bar/pie: one colour per point from BRAND_PALETTE
                    fills = [el.get("val") for el in t.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")]
                    self.assertTrue(fills and set(fills) <= set(C.BRAND_PALETTE), p.name)
            for c in charts:
                self.assertIn(c["type"], ("bar", "pie", "scatter"))
                self.assertRegex(c["anchor"], r"^[A-Z]{1,3}\d+$")

    def test_chart_anchor_beside_table_header(self):
        for p in _all_paths():
            for t in _tables(_market_of(p), p.name):
                for c in t["charts"]:
                    m = re.match(r"^([A-Z]+)(\d+)$", c["anchor"])
                    col = openpyxl.utils.column_index_from_string(m.group(1))
                    self.assertEqual(int(m.group(2)), t["header_row"], (p.name, t["sheet"], c))
                    self.assertGreaterEqual(col, t["last_col"] + 2, (p.name, t["sheet"], c))

    # ---------------------------------------------------------------- formulas / text
    def test_only_hyperlink_formulas_with_row_asin(self):
        for p in _all_paths():
            market = _market_of(p)
            tables = _tables(market, p.name)
            wb = openpyxl.load_workbook(p)
            n_links = 0
            for ws in wb.worksheets:
                sheet_tables = [t for t in tables if t["sheet"] == ws.title]
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.data_type != "f":
                            continue
                        m = None
                        for rx in C.ALLOWED_FORMULA_RES:
                            m = rx.match(str(cell.value))
                            if m:
                                break
                        self.assertIsNotNone(m, f"{p.name}!{ws.title}!{cell.coordinate}: {cell.value}")
                        owner = [t for t in sheet_tables if t["first_data_row"] and t["first_data_row"] <= cell.row <= t["last_data_row"]]
                        self.assertEqual(len(owner), 1, f"{p.name}!{ws.title}!{cell.coordinate}")
                        t = owner[0]
                        asin_col = t["first_col"] + t["columns"].index("ASIN")
                        self.assertEqual(ws.cell(cell.row, asin_col).value, m.group("asin"), f"{ws.title}!{cell.coordinate}")
                        header = ws.cell(t["header_row"], cell.column).value
                        if m.group("tld") == "com":
                            self.assertTrue(market == "US" or (ws.title in C.CA_SHEETS_ALLOWING_US and "US" in header),
                                            f"{p.name}!{ws.title}!{cell.coordinate} header={header}")
                        else:
                            self.assertEqual(market, "CA", f"{p.name}!{ws.title}!{cell.coordinate}")
                        self.assertEqual(cell.font.color.rgb[-6:], C.COLOR_LINK)
                        n_links += 1
            self.assertGreater(n_links, 0, p.name)

    def test_formula_like_titles_stay_strings(self):
        # the US fixture carries a product title "=1+1 trap ..."
        wb = openpyxl.load_workbook(_Built.us_gauge)
        hits = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=1+1 trap")]
        self.assertGreater(len(hits), 0)
        for c in hits:
            self.assertEqual(c.data_type, "s", c.coordinate)
        # engine level: a table title and a text cell starting with "=" are strings
        book = X.Book("engine_trap.xlsx", C.MARKETS["CA"])
        ws = book.sheet("Summary")
        rows = pd.DataFrame([{"asin": "B0TESTTRP1", "title": "=1+1 trap", "rev": 1.0, "market": "CA"}])
        spec = X.TableSpec(sheet="Summary", role="summary_brands", title="=SUM(A1:A2) title trap",
                           columns=[X.ColumnSpec("ASIN", "asin", "text", 12), X.ColumnSpec("Product Name", "title", "text", 30),
                                    X.ColumnSpec("Monthly Rev (CAD)", "rev", "money", 12), X.ColumnSpec("Link", "asin", "link", 20)],
                           rows=rows, total_row=None, residual_row=None, allowed_markets=("CA",))
        tr = book.table(ws, spec, 1)
        path = OUT / "engine_trap.xlsx"
        book.save(path)
        ws2 = openpyxl.load_workbook(path)["Summary"]
        self.assertEqual(ws2.cell(tr.header_row - 1, 1).data_type, "s")
        self.assertEqual(ws2.cell(tr.header_row - 1, 1).value, "=SUM(A1:A2) title trap")
        self.assertEqual(ws2.cell(tr.first_data_row, 2).data_type, "s")
        self.assertEqual(ws2.cell(tr.first_data_row, 2).value, "=1+1 trap")
        self.assertEqual(ws2.cell(tr.first_data_row, 4).value,
                         '=HYPERLINK("https://amazon.ca/dp/B0TESTTRP1","Amazon.ca - B0TESTTRP1")')

    def test_no_error_tokens(self):
        for p in _all_paths():
            wb = openpyxl.load_workbook(p)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for v in row:
                        if isinstance(v, str):
                            self.assertFalse(any(tok in v for tok in ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!")), (p.name, ws.title, v))
                        if isinstance(v, float):
                            self.assertTrue(math.isfinite(v), (p.name, ws.title, v))

    # ---------------------------------------------------------------- currency
    def test_money_formats_carry_market_currency(self):
        for p in _all_paths():
            market = _market_of(p)
            wb = openpyxl.load_workbook(p)
            ca_fmt = 0
            for ws in wb.worksheets:
                hdr_by_col = {}
                for t in _tables(market, p.name):
                    if t["sheet"] == ws.title and t["header_row"]:
                        for i, h in enumerate(t["columns"]):
                            hdr_by_col.setdefault(t["first_col"] + i, []).append(h)
                for row in ws.iter_rows():
                    for c in row:
                        fmt = c.number_format or ""
                        if "$" not in fmt:
                            continue
                        if market == "US":
                            self.assertNotIn("CA$", fmt, f"{p.name}!{ws.title}!{c.coordinate}")
                        elif "CA$" in fmt:
                            ca_fmt += 1
                        else:
                            self.assertIn(ws.title, C.CA_SHEETS_ALLOWING_US, f"{p.name}!{ws.title}!{c.coordinate} {fmt}")
                            self.assertTrue(any("(USD)" in h for h in hdr_by_col.get(c.column, [])),
                                            f"{p.name}!{ws.title}!{c.coordinate} {hdr_by_col.get(c.column)}")
            if market == "CA":
                self.assertGreater(ca_fmt, 0, p.name)

    def test_headers_carry_currency(self):
        for p in _all_paths():
            market = _market_of(p)
            ccy = C.MARKETS[market].currency
            for t in _tables(market, p.name):
                for h in t["columns"]:
                    if "(CAD)" in h or "(USD)" in h:
                        if market == "US":
                            self.assertNotIn("(CAD)", h, (p.name, t["sheet"], h))
                        elif "(USD)" in h:
                            self.assertIn(t["sheet"], C.CA_SHEETS_ALLOWING_US, (p.name, t["sheet"], h))
                self.assertFalse(any("{ccy}" in h for h in t["columns"]), (p.name, t["sheet"]))
            self.assertTrue(any(f"({ccy})" in h for t in _tables(market, p.name) for h in t["columns"]))

    # ---------------------------------------------------------------- totals / shares / residuals
    def _summary_check(self, path: Path, dataset: pd.DataFrame):
        market = _market_of(path)
        t = [t for t in _tables(market, path.name) if t["sheet"] == "Summary" and t["role"] == "summary_brands"][0]
        ws = openpyxl.load_workbook(path)["Summary"]
        col = {h: t["first_col"] + i for i, h in enumerate(t["columns"])}
        ccy = C.MARKETS[market].currency
        rev_c, units_c = col[f"Monthly Rev ({ccy})"], col["Monthly Units"]
        share_c, list_c = col["Monthly Rev Market Share %"], col["# of Listings"]
        tr = t["total_row"]
        self.assertAlmostEqual(ws.cell(tr, rev_c).value, float(dataset["revenue_month"].sum()), places=6)
        self.assertAlmostEqual(ws.cell(tr, units_c).value, float(dataset["units_month"].sum()), places=6)
        self.assertEqual(ws.cell(tr, list_c).value, dataset["asin"].nunique())
        last = t["residual_row"] or t["last_data_row"]
        shares = [ws.cell(r, share_c).value for r in range(t["first_data_row"], last + 1)]
        self.assertAlmostEqual(sum(shares), 1.0, delta=1e-9)
        revs = [ws.cell(r, rev_c).value for r in range(t["first_data_row"], last + 1)]
        self.assertAlmostEqual(sum(revs), ws.cell(tr, rev_c).value, places=6)
        self.assertEqual(ws.cell(tr, 1).value, C.TOTAL_ROW_LABEL)
        self.assertEqual(ws.cell(tr, 1).fill.fgColor.rgb[-6:], C.FILL_TOTAL)
        # hidden rows are exactly the zero-revenue data rows (never deleted)
        for r in range(t["first_data_row"], t["last_data_row"] + 1):
            self.assertEqual(bool(ws.row_dimensions[r].hidden), ws.cell(r, rev_c).value == 0, (path.name, r))
        return t, ws

    def test_cr_summary_totals_full_dataset(self):
        rows = _Built.ca_rows[_Built.ca_rows["source_set"].isin(["code_reader", "both"])]
        t, ws = self._summary_check(_Built.cr_paths[0], rows)
        # Innova row red
        brand_cells = [ws.cell(r, 1) for r in range(t["first_data_row"], t["last_data_row"] + 1)]
        innova = [c for c in brand_cells if c.value == "Innova"]
        self.assertEqual(len(innova), 1)
        self.assertEqual(innova[0].font.color.rgb[-6:], C.COLOR_INNOVA)

    def test_gauge_summary_totals_core_devices(self):
        for path, rows in ((_Built.ca_gauge, _Built.ca_rows), (_Built.us_gauge, _Built.us_rows)):
            core = rows[rows["gauge_class"].isin(C.GAUGE_DEVICE_CLASSES)]
            self._summary_check(path, core)

    def test_residual_rows_when_truncated(self):
        cr = _Built.ca_rows[_Built.ca_rows["source_set"].isin(["code_reader", "both"])]
        shown, residual, total = X.brand_summary(cr, top_n=3)
        self.assertEqual(len(shown), 3)
        n_brands = cr["brand_key"].nunique()
        self.assertEqual(residual["brand"], C.RESIDUAL_ROW_LABEL.format(noun="brands", n=n_brands - 3))
        self.assertAlmostEqual(shown["share"].sum() + residual["share"], 1.0, delta=1e-9)
        self.assertAlmostEqual(shown["rev"].sum() + residual["rev"], total["rev"], places=6)
        self.assertEqual(total["listings"], cr["asin"].nunique())
        # every truncated registered table has a residual row labelled "Other <noun> (n)"
        seen = 0
        for p in _all_paths():
            wb = openpyxl.load_workbook(p)
            for t in _tables(_market_of(p), p.name):
                if t["residual_row"]:
                    ws = wb[t["sheet"]]
                    cells = [ws.cell(t["residual_row"], c) for c in range(t["first_col"], t["last_col"] + 1)]
                    labels = [c for c in cells if isinstance(c.value, str) and re.match(r"^Other (brands|listings) \(\d+\)$", c.value)]
                    self.assertEqual(len(labels), 1, (p.name, t["sheet"], [c.value for c in cells]))
                    self.assertTrue(labels[0].font.i)
                    self.assertTrue(t["total_row"], (p.name, t["sheet"]))   # a truncated table always carries the full-dataset Total
                    seen += 1
        self.assertGreater(seen, 0)
        # Category tab: Total tier has 11 brands > CATEGORY_TOP_BRANDS -> residual present
        cat = [t for t in _tables("CA", _Built.cr_paths[1].name) if t["role"] == "category" and t["title"].startswith("Total —")]
        self.assertEqual(len(cat), 1)
        self.assertIsNotNone(cat[0]["residual_row"])

    def test_top50_rank1_and_rankings(self):
        rows = _Built.ca_rows[_Built.ca_rows["source_set"].isin(["code_reader", "both"])]
        p = _Built.cr_paths[0]
        t = [t for t in _tables("CA", p.name) if t["role"] == "top_by_revenue"][0]
        ws = openpyxl.load_workbook(p)["Top 50"]
        rev_c = t["first_col"] + t["columns"].index("Est. Monthly Retail Rev (CAD)")
        self.assertAlmostEqual(ws.cell(t["first_data_row"], rev_c).value, float(rows["revenue_month"].max()))
        u = [t for t in _tables("CA", p.name) if t["role"] == "top_by_units"][0]
        self.assertEqual(u["sheet"], "Top 50")
        prev_end = t["total_row"] or t["residual_row"] or t["last_data_row"]
        self.assertEqual(u["header_row"] - 1 - prev_end, 4)   # 3 blank rows then the table title row

    def test_brand_tab_layout(self):
        for p in (_Built.cr_paths[0], _Built.ca_gauge):
            for t in _tables(_market_of(p), p.name):
                if t["role"] in ("brand_tab_revenue",) or (t["role"] == "innova" and p == _Built.cr_paths[0]):
                    self.assertEqual(t["header_row"], C.BRAND_TAB_RESERVED_ROWS + 2, (p.name, t["sheet"]))
                if t["role"] == "kpi" and t["sheet"] not in ("Summary",):
                    self.assertEqual((t["first_data_row"], t["last_data_row"]), (3, 6), (p.name, t["sheet"]))

    def test_innova_tab(self):
        rows = _Built.ca_rows[_Built.ca_rows["source_set"].isin(["code_reader", "both"])]
        inn = rows[rows["brand_key"] == "innova"]
        p = _Built.cr_paths[0]
        t = [t for t in _tables("CA", p.name) if t["role"] == "innova" and t["sheet"] == "Innova"][0]
        self.assertEqual(t["last_data_row"] - t["first_data_row"] + 1, len(inn))
        ws = openpyxl.load_workbook(p)["Innova"]
        item_c = t["first_col"] + t["columns"].index("Item #")
        sold_c = t["first_col"] + t["columns"].index("Sold by Amazon")
        items = {ws.cell(r, item_c).value for r in range(t["first_data_row"], t["last_data_row"] + 1)}
        self.assertEqual(items, {"5610", "1000 V2"})
        sold = {ws.cell(r, sold_c).value for r in range(t["first_data_row"], t["last_data_row"] + 1)}
        self.assertEqual(sold, {"Yes", "No"})
        self.assertEqual(CR.item_number("Innova 3020RS OBD2 Code Reader"), "3020RS")
        self.assertEqual(CR.sold_by_amazon("Amazon US"), "Amazon US (cross-border)")
        self.assertEqual(CR.sold_by_amazon("Amazon.ca"), "Yes")
        # gauge Innova tab: static device-count line
        for g in (_Built.ca_gauge, _Built.us_gauge):
            ws = openpyxl.load_workbook(g)["Innova"]
            vals = [c for row in ws.iter_rows(values_only=True) for c in row if isinstance(c, str)]
            self.assertIn("Innova gauge/HUD device listings in this dataset: 0", vals)

    def test_same_asin_and_benchmark(self):
        p = _Built.ca_gauge
        t = [t for t in _tables("CA", p.name) if t["role"] == "same_asin"][0]
        self.assertGreaterEqual(t["last_data_row"] - t["first_data_row"] + 1, 1)
        self.assertEqual(t["sheet"], "US vs CA Same-ASIN")
        self.assertEqual(tuple(t["allowed_markets"]), ("CA", "US"))
        bench = [t for t in _tables("CA", p.name) if t["role"].startswith("benchmark_")]
        self.assertEqual({t["role"] for t in bench}, {"benchmark_brands", "benchmark_subtypes", "benchmark_tiers"})
        for t in bench:
            self.assertTrue(any("(CAD)" in h for h in t["columns"]) and any("(USD)" in h for h in t["columns"]), t["columns"])
            self.assertFalse(any("ratio" in h.lower() and "rev" in h.lower() for h in t["columns"]), t["columns"])

    def test_all_products_and_excluded_reconcile(self):
        p = _Built.ca_gauge
        rows = _Built.ca_rows
        allp = [t for t in _tables("CA", p.name) if t["role"] == "all_rows"][0]
        self.assertEqual(allp["last_data_row"] - allp["first_data_row"] + 1, rows["asin"].nunique())
        exc = [t for t in _tables("CA", p.name) if t["role"] == "excluded"][0]
        n_exc = rows["gauge_class"].isin(C.GAUGE_EXCLUDED_CLASSES + ("ambiguous",)).sum()
        self.assertEqual(exc["last_data_row"] - exc["first_data_row"] + 1, n_exc)

    def test_no_duplicate_asin_within_a_table(self):
        for p in _all_paths():
            wb = openpyxl.load_workbook(p)
            for t in _tables(_market_of(p), p.name):
                if "ASIN" not in t["columns"] or not t["first_data_row"]:
                    continue
                ws = wb[t["sheet"]]
                c = t["first_col"] + t["columns"].index("ASIN")
                asins = [ws.cell(r, c).value for r in range(t["first_data_row"], t["last_data_row"] + 1)]
                asins = [a for a in asins if a]
                self.assertEqual(len(asins), len(set(asins)), (p.name, t["sheet"], t["role"]))

    # ---------------------------------------------------------------- registry / metadata / manifest
    def test_registry_every_table_and_brand_sheet_map(self):
        need = {
            C.cr_report_name("CA", MONTH): {"summary_brands", "top_by_revenue", "top_by_units", "innova", "kpi", "brand_tab_revenue",
                                            "brand_tab_units", "all_rows", "metadata"},
            C.cr_analysis_name("CA", MONTH): {"tier_pivot", "category", "tier_tab", "trend_proxy", "type_coverage", "metadata"},
            C.gauge_report_name("CA", MONTH): {"read_me", "summary_brands", "subtype_mix", "tier_matrix", "kpi", "top_by_revenue",
                                               "top_by_units", "innova", "modelb_top", "brand_tab_revenue", "brand_tab_units",
                                               "price_ladder", "feature_matrix", "modelb_tiers", "modelb_brands", "modelb_app_matrix",
                                               "benchmark_brands", "benchmark_subtypes", "benchmark_tiers", "same_asin", "all_rows",
                                               "dedupe_audit", "excluded", "source_method", "metadata"},
            C.gauge_report_name("US", MONTH): {"read_me", "summary_brands", "subtype_mix", "tier_matrix", "kpi", "top_by_revenue",
                                               "top_by_units", "innova", "all_rows", "dedupe_audit", "excluded", "source_method",
                                               "metadata"},
        }
        for p in _all_paths():
            market = _market_of(p)
            reg = _Built.reg[market]
            self.assertEqual((reg["market"], reg["month"]), (market, MONTH))
            tables = _tables(market, p.name)
            roles = {t["role"] for t in tables}
            self.assertTrue(roles <= set(C.TABLE_ROLES), roles - set(C.TABLE_ROLES))
            self.assertTrue(need[p.name] <= roles, need[p.name] - roles)
            self.assertIn(p.name, reg["workbooks"])
            bsm = reg["workbooks"][p.name]["brand_sheet_map"]
            self.assertIsInstance(bsm, dict)
            wb = openpyxl.load_workbook(p)
            self.assertEqual(reg["workbooks"][p.name]["sheets"], wb.sheetnames)
            for key, sheet in bsm.items():
                self.assertIn(sheet, wb.sheetnames, key)
            for t in tables:
                self.assertEqual(t["brand_sheet_map"], bsm)
                for k in ("workbook", "sheet", "role", "header_row", "first_data_row", "last_data_row", "total_row", "residual_row",
                          "columns", "dataset_filter", "allowed_markets", "charts", "brand_sheet_map"):
                    self.assertIn(k, t)
                if t["header_row"]:
                    ws = wb[t["sheet"]]
                    got = [ws.cell(t["header_row"], t["first_col"] + i).value for i in range(len(t["columns"]))]
                    self.assertEqual(got, t["columns"], (p.name, t["sheet"], t["role"]))
            if p.name.startswith("CA_Code_Reader_Competitor"):
                self.assertEqual(bsm["innova"], "Innova")
                self.assertEqual(len(bsm), 1 + 10)

    def test_metadata_required_keys(self):
        for p in _all_paths():
            ws = openpyxl.load_workbook(p)["Metadata"]
            keys = {r[0] for r in ws.iter_rows(values_only=True) if r and r[0]}
            for k in C.METADATA_REQUIRED_KEYS:
                self.assertIn(k, keys, p.name)
            vals = {r[0]: r[1] for r in ws.iter_rows(values_only=True) if r and r[0]}
            if _market_of(p) == "CA":
                self.assertTrue(vals["Currency"].startswith("CAD"))
            self.assertIn("normalized", vals["Input mode"].lower())
        g = {r[0]: r[1] for r in openpyxl.load_workbook(_Built.ca_gauge)["Metadata"].iter_rows(values_only=True) if r and r[0]}
        self.assertIn("US benchmark source", g)

    def test_manifest_hashes_and_excludes_itself(self):
        mpath = OUT / C.manifest_name(MONTH)
        man = json.loads(mpath.read_text())
        self.assertNotIn(mpath.name, man["outputs"])
        for p in _all_paths():
            self.assertEqual(man["outputs"][p.name]["sha256"], _sha(p), p.name)
        self.assertTrue(man["inputs"])
        self.assertIn(str(CA_FIXTURE), man["inputs"])
        self.assertEqual(man["inputs"][str(CA_FIXTURE)], _sha(CA_FIXTURE))
        self.assertRegex(man["pipeline_git_sha"], r"^[0-9a-f]{7,40}(-dirty)?$")
        self.assertIn("generated_at", man)


class TestSafeOutputPath(unittest.TestCase):
    def test_refuse_backup_and_hand_edited(self):
        d = C.SCRATCH_DIR / "test_out_safe"
        if d.exists():
            shutil.rmtree(d)
        d.mkdir(parents=True)
        name = "CA_Code_Reader_Competitor_Report_202609.xlsx"
        f = d / name
        f.write_bytes(b"original")
        with self.assertRaises(FileExistsError):
            X.safe_output_path(d, name, overwrite=False, manifest=None)
        manifest = {"outputs": {name: {"sha256": hashlib.sha256(b"something else").hexdigest()}}}
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            out = X.safe_output_path(d, name, overwrite=True, manifest=manifest)
        self.assertEqual(out, f)
        self.assertFalse(f.exists())
        self.assertIn(f"HAND-EDITED: {f}", buf.getvalue())
        backups = list((d / C.BACKUP_SUBDIR).glob("CA_Code_Reader_Competitor_Report_202609.*.xlsx"))
        self.assertEqual(len(backups), 1)
        self.assertRegex(backups[0].name, r"\.\d{8}-\d{6}\.xlsx$")
        self.assertEqual(backups[0].read_bytes(), b"original")
        # unchanged file (sha matches the manifest): backed up, no HAND-EDITED line
        f.write_bytes(b"v2")
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            X.safe_output_path(d, name, overwrite=True, manifest={"outputs": {name: {"sha256": hashlib.sha256(b"v2").hexdigest()}}})
        self.assertNotIn("HAND-EDITED", buf.getvalue())
        # Excel lock file -> refuse
        f.write_bytes(b"v3")
        (d / f"~${name}").write_bytes(b"")
        with self.assertRaises(RuntimeError):
            X.safe_output_path(d, name, overwrite=True, manifest=None)


class TestSourceHygiene(unittest.TestCase):
    def test_no_row_insertion_or_deletion(self):
        for m in OWN_MODULES:
            src = (C.PACKAGE_ROOT / m).read_text()
            self.assertNotIn("insert_rows", src, m)
            self.assertNotIn("delete_rows", src, m)

    def test_build_code_never_reopens_outputs(self):
        for m in BUILD_MODULES:
            src = (C.PACKAGE_ROOT / m).read_text()
            self.assertNotIn("load_workbook", src, m)

    def test_cli_dev_mode_end_to_end(self):
        out = C.SCRATCH_DIR / "test_out_cli"
        if out.exists():
            shutil.rmtree(out)
        with contextlib.redirect_stdout(io.StringIO()):
            rc = CR.main(["--month", MONTH, "--from-normalized", str(CA_FIXTURE), "--out-dir", str(out),
                          "--runs-dir", str(out / "runs"), "--overwrite"])
        self.assertEqual(rc, 0)
        self.assertTrue((out / C.cr_report_name("CA", MONTH)).exists())
        with self.assertRaises(SystemExit):
            with contextlib.redirect_stderr(io.StringIO()):
                CR.main(["--month", "202613", "--from-normalized", str(CA_FIXTURE), "--out-dir", str(out)])
        with self.assertRaises(ValueError):   # dev runs never write under NewProductCategory
            CR.main(["--month", MONTH, "--from-normalized", str(CA_FIXTURE), "--out-dir", str(C.NEW_PRODUCT_DIR / "x")])

    def test_missing_column_fails_loudly(self):
        df = X.read_normalized_csv(CA_FIXTURE).drop(columns=["revenue_month"])
        with self.assertRaises(KeyError):
            X.dataset_from_normalized(df, "CA", MONTH)

    def test_market_mismatch_fails(self):
        ca = X.dataset_from_normalized(X.read_normalized_csv(CA_FIXTURE), "CA", MONTH)
        with self.assertRaises(ValueError):
            X.dataset_from_normalized(X.read_normalized_csv(CA_FIXTURE), "US", MONTH)
        ca.market = "US"
        with self.assertRaises(ValueError):
            CR.build_code_reader_workbooks(ca, C.SCRATCH_DIR / "never", overwrite=True, dated_copy=False, runs_dir=RUNS)


class TestPreview(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _Built.ensure()

    def test_preview_and_excerpt(self):
        html = OUT / "preview.html"
        md = OUT / "excerpt.md"
        with contextlib.redirect_stdout(io.StringIO()):
            rc = RP.main(["--workbook", str(_Built.cr_paths[0]), "--out", str(html), "--excerpt", str(md)])
        self.assertEqual(rc, 0)
        text = html.read_text()
        for s in ("Summary", "Top 50", "Metadata", C.FILL_TITLE, C.FILL_TOTAL):
            self.assertIn(s, text)
        lines = [l for l in md.read_text().splitlines() if l.startswith("|")]
        self.assertLessEqual(len(lines), 20 + 10 + 4)
        self.assertIn("CA$", md.read_text())


if __name__ == "__main__":
    unittest.main()
