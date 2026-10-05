"""validate_outputs tests: fixture workbooks built through the builders' normalized (dev) path into
tmp/ca_scratch/validate_test/, then validated with a tiny memo fixture. Tampered copies (openpyxl edit of a copied
workbook, edited manifest) must make the matching checks FAIL. Writes only under tmp/ca_scratch/validate_test/.
"""
from __future__ import annotations

import contextlib
import io
import json
import re
import shutil
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

from ca_market_reports import build_ca_code_reader_report as CR
from ca_market_reports import build_gauge_report as G
from ca_market_reports import ca_common as C
from ca_market_reports import ca_xlsx_style as X
from ca_market_reports import validate_outputs as V

MONTH = "202609"
ROOT = C.SCRATCH_DIR / "validate_test"
BUILD = ROOT / "build"
RUNS = BUILD / "runs"
CA_FIXTURE = C.FIXTURES_DIR / "normalized_rows.csv"
US_FIXTURE = C.FIXTURES_DIR / "normalized_rows_us.csv"
ABSENT = ROOT / "absent_raw"          # never created: the "raw dir absent" case
CR_REPORT = C.cr_report_name("CA", MONTH)
CHECK_LINE_RE = re.compile(r"^(PASS|FAIL|SKIP) (V\d{2}) ([^:]+): (.*)$")
SRC_URL = "https://example.org/gauge-source"


def build_set(build: Path, runs: Path, ca_csv: Path, us_csv: Path) -> None:
    """CR workbooks, US gauge and CA gauge (with the US benchmark) built through the builders' normalized (dev) path."""
    for d in ("cr", "gauge_ca", "gauge_us"):
        (build / d).mkdir(parents=True)
    ca = X.dataset_from_normalized(X.read_normalized_csv(ca_csv), "CA", MONTH)
    us = X.dataset_from_normalized(X.read_normalized_csv(us_csv), "US", MONTH)
    with contextlib.redirect_stdout(io.StringIO()):
        CR.build_code_reader_workbooks(ca, build / "cr", overwrite=True, dated_copy=False, runs_dir=runs, input_paths=[ca_csv])
        G.build_gauge_workbook(us, build / "gauge_us", benchmark=None, overwrite=True, dated_copy=False, runs_dir=runs,
                               preclassified=True, input_paths=[us_csv])
        G.build_gauge_workbook(ca, build / "gauge_ca", benchmark=us, overwrite=True, dated_copy=False, runs_dir=runs,
                               preclassified=True, input_paths=[ca_csv, us_csv])


class _Built:
    done = False

    @classmethod
    def ensure(cls) -> None:
        if cls.done:
            return
        if ROOT.exists():
            shutil.rmtree(ROOT)
        build_set(BUILD, RUNS, CA_FIXTURE, US_FIXTURE)
        reg = json.loads(C.run_file(RUNS, MONTH, "table_registry", "CA", "json").read_text(encoding="utf-8"))
        summ = next(t for t in reg["tables"] if t["workbook"] == CR_REPORT and t["role"] == "summary_brands")
        cls.summary = summ
        rev_col = summ["first_col"] + summ["columns"].index("Monthly Rev (CAD)")
        lst_col = summ["first_col"] + summ["columns"].index("# of Listings")
        wb = load_workbook(BUILD / "cr" / CR_REPORT)
        ws = wb["Summary"]
        cls.rev_ref = f"{ws.cell(summ['total_row'], rev_col).coordinate}"
        cls.lst_ref = f"{ws.cell(summ['total_row'], lst_col).coordinate}"
        cls.rev_value = float(ws.cell(summ["total_row"], rev_col).value)
        cls.lst_value = int(ws.cell(summ["total_row"], lst_col).value)
        all_asins = next(t for t in reg["tables"] if t["workbook"] == CR_REPORT and t["role"] == "all_rows")
        cls.all_rev_cell = (all_asins["first_data_row"], all_asins["first_col"] + all_asins["columns"].index("Monthly Rev (CAD)"))
        sources = ROOT / "sources.csv"
        sources.write_text(",".join(C.SOURCES_COLUMNS) + "\n" + f"{SRC_URL},2026-10-02,Example,claim,,1\n", encoding="utf-8")
        cls.sources = sources
        cls.done = True


def memo_text(*, fixed: bool) -> str:
    b = _Built
    good = f"Observed revenue CA${round(b.rev_value):,} [WB: {CR_REPORT}!Summary!{b.rev_ref}] in the exports."
    listings = b.lst_value if fixed else b.lst_value + 7
    wrong = f"Listing count {listings} [WB: {CR_REPORT}!Summary!{b.lst_ref}]."
    lines = ["# Memo fixture", "", good, wrong, f"Lordco carries the unit [SRC: {SRC_URL}, accessed 2026-10-02]."]
    if not fixed:
        lines.append("Gauge share placeholder [WB: TBD].")
    return "\n".join(lines) + "\n"


def write_memo(name: str, *, fixed: bool) -> Path:
    p = ROOT / name
    p.write_text(memo_text(fixed=fixed), encoding="utf-8")
    return p


def run_validator(*, cr: Path = BUILD / "cr", gauge_ca: Path = BUILD / "gauge_ca", gauge_us: Path = BUILD / "gauge_us",
                  memo: Path, extra: list[str] | None = None, json_path: Path | None = None,
                  runs: Path = RUNS, us_csv: Path = US_FIXTURE) -> tuple[int, list[str]]:
    argv = ["--month", MONTH, "--cr-out-dir", str(cr), "--gauge-out-dir", str(gauge_ca), "--us-gauge-out-dir", str(gauge_us),
            "--runs-dir", str(runs), "--raw-dir", str(ABSENT / "ca_cr"), "--gauge-raw-dir", str(ABSENT / "ca_gauge"),
            "--us-gauge-raw-dir", str(ABSENT / "us_gauge"), "--us-cr-raw-dir", str(ABSENT / "us_cr"),
            "--from-normalized", str(CA_FIXTURE), "--us-from-normalized", str(us_csv),
            "--memo", str(memo), "--sources", str(_Built.sources)] + (extra or [])
    if json_path is not None:
        argv += ["--json", str(json_path)]
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = V.main(argv)
    return rc, buf.getvalue().splitlines()


def _core(df):
    """Core devices of a normalized fixture: device class and not borderline (columns derived when absent)."""
    dev = df["gauge_class"].isin(C.GAUGE_DEVICE_CLASSES)
    border = df["borderline"].astype(bool) if "borderline" in df.columns else pd.Series(False, index=df.index)
    return df[dev & ~border]


def _registry(mk: str) -> dict:
    return json.loads(C.run_file(RUNS, MONTH, "table_registry", mk, "json").read_text(encoding="utf-8"))


def _special_cells() -> dict:
    """Values of the tables covered by the new V09 rules, read from the fixture workbooks via the registry."""
    reg = _registry("CA")
    gname, aname = C.gauge_report_name("CA", MONTH), C.cr_analysis_name("CA", MONTH)
    wb = load_workbook(BUILD / "gauge_ca" / gname)
    out: dict = {}

    def rows(t, ws):
        return {ws.cell(r, t["first_col"]).value: {h: ws.cell(r, t["first_col"] + j).value for j, h in enumerate(t["columns"])}
                for r in range(t["first_data_row"], t["last_data_row"] + 1)}

    for t in reg["tables"]:
        if t["workbook"] != gname:
            continue
        ws = wb[t["sheet"]]
        if t["title"] == V.SHARE_TABLE_TITLE:
            out["share"] = rows(t, ws)
        elif t["title"] == V.BENCH_SHARE_TABLE_TITLE:
            out["bench"] = rows(t, ws)
        elif t["title"] == V.FUEL_TABLE_TITLE:
            out["fuel_total"] = {h: ws.cell(t["total_row"], t["first_col"] + j).value for j, h in enumerate(t["columns"])}
        elif t["role"] == "kpi" and t["sheet"] == "Innova":
            out["innova_b3"] = ws.cell(t["first_data_row"], t["first_col"]).value
        elif t["role"] == "kpi" and t["sheet"] == "US vs CA Same-ASIN":
            out["same_asin_b3"] = ws.cell(t["first_data_row"], t["first_col"] + 1).value
        elif t["role"] == "modelb_app_matrix":
            out["apps"] = [ws.cell(r, t["first_col"]).value for r in range(t["first_data_row"], t["last_data_row"] + 1)]
    trend = next(t for t in reg["tables"] if t["workbook"] == aname and t["role"] == "trend_proxy")
    ws = load_workbook(BUILD / "cr" / aname)["Trend Proxy"]
    out["trend_total"] = {h: ws.cell(trend["total_row"], trend["first_col"] + j).value for j, h in enumerate(trend["columns"])}
    return out


def statuses(lines: list[str]) -> dict[str, tuple[str, str]]:
    out = {}
    for ln in lines[:-1]:
        m = CHECK_LINE_RE.match(ln)
        if m:
            out[m.group(2)] = (m.group(1), m.group(4))
    return out


class ValidatorFixtureTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _Built.ensure()
        cls.rc_bad, cls.lines_bad = run_validator(memo=write_memo("memo_bad.md", fixed=False),
                                                  json_path=ROOT / "result_bad.json")
        cls.rc_ok, cls.lines_ok = run_validator(memo=write_memo("memo_ok.md", fixed=True))

    def test_exactly_23_check_lines_in_order_then_final_line(self):
        for lines in (self.lines_bad, self.lines_ok):
            self.assertEqual(len(lines), 24, "\n".join(lines))
            ids = []
            for ln in lines[:23]:
                m = CHECK_LINE_RE.match(ln)
                self.assertIsNotNone(m, ln)
                ids.append(m.group(2))
            self.assertEqual(ids, list(C.VALIDATION_CHECKS))
            self.assertRegex(lines[-1], C.VALIDATION_FINAL_RE)

    def test_memo_defects_fail_v20_and_the_fixed_memo_passes(self):
        st, ev = statuses(self.lines_bad)["V20"]
        self.assertEqual(st, "FAIL", ev)
        self.assertIn(f"{CR_REPORT}!Summary!{_Built.lst_ref}", ev)     # the wrong number
        self.assertIn("1 [WB: TBD]", ev)                                  # the placeholder count
        self.assertNotIn(f"Summary!{_Built.rev_ref}", ev)                 # the correct tag is not reported
        self.assertEqual(self.rc_bad, 1)
        self.assertRegex(self.lines_bad[-1], r"^VALIDATION: FAIL \(\d+/23; failed: .*V20.*\)$")
        st, ev = statuses(self.lines_ok)["V20"]
        self.assertEqual(st, "PASS", ev)

    def test_clean_fixture_passes_with_real_data_only_checks_skipped(self):
        st = statuses(self.lines_ok)
        failed = {k: v for k, v in st.items() if v[0] == "FAIL"}
        self.assertEqual(failed, {}, "\n".join(self.lines_ok))
        self.assertEqual(self.rc_ok, 0)
        self.assertEqual(st["V17"][0], "SKIP")
        self.assertIn("raw dir absent", st["V17"][1])
        self.assertIn(str(ABSENT / "ca_cr"), st["V17"][1])
        self.assertIn("raw-input hashes skipped", st["V19"][1])
        self.assertIn("Bully Dog check skipped", st["V12"][1])
        self.assertRegex(self.lines_ok[-1], r"^VALIDATION: PASS \(22/23\)$")

    def test_json_results(self):
        data = json.loads((ROOT / "result_bad.json").read_text(encoding="utf-8"))
        self.assertEqual([r["id"] for r in data["checks"]], list(C.VALIDATION_CHECKS))
        self.assertEqual(data["final"], self.lines_bad[-1])
        self.assertEqual(data["exit_code"], 1)

    def test_tampered_summary_total_fails_v01_and_v09(self):
        d = ROOT / "tamper_total"
        shutil.copytree(BUILD / "cr", d)
        wb = load_workbook(d / CR_REPORT)
        ws = wb["Summary"]
        ws[_Built.rev_ref].value = _Built.rev_value + 250.0
        wb.save(d / CR_REPORT)
        rc, lines = run_validator(cr=d, memo=write_memo("memo_ok2.md", fixed=True))
        st = statuses(lines)
        self.assertEqual(rc, 1)
        self.assertEqual(st["V01"][0], "FAIL", st["V01"])
        self.assertEqual(st["V09"][0], "FAIL", st["V09"])
        self.assertEqual(st["V19"][0], "FAIL", st["V19"])          # the edited file no longer matches its manifest hash

    def test_tampered_all_asins_row_fails_v18_and_v09(self):
        d = ROOT / "tamper_row"
        shutil.copytree(BUILD / "cr", d)
        wb = load_workbook(d / CR_REPORT)
        ws = wb["All ASINs"]
        r, c = _Built.all_rev_cell
        ws.cell(r, c).value = float(ws.cell(r, c).value) + 99.0
        wb.save(d / CR_REPORT)
        rc, lines = run_validator(cr=d, memo=write_memo("memo_ok3.md", fixed=True))
        st = statuses(lines)
        self.assertEqual(rc, 1)
        self.assertEqual(st["V18"][0], "FAIL", st["V18"])
        self.assertEqual(st["V09"][0], "FAIL", st["V09"])
        self.assertEqual(st["V01"][0], "PASS", st["V01"])

    def test_tampered_manifest_fails_v19(self):
        d = ROOT / "tamper_manifest"
        shutil.copytree(BUILD / "gauge_us", d)
        mp = d / C.manifest_name(MONTH)
        m = json.loads(mp.read_text(encoding="utf-8"))
        name = next(iter(m["outputs"]))
        m["outputs"][name]["sha256"] = "0" * 64
        mp.write_text(json.dumps(m, indent=1), encoding="utf-8")
        rc, lines = run_validator(gauge_us=d, memo=write_memo("memo_ok4.md", fixed=True))
        st = statuses(lines)
        self.assertEqual(st["V19"][0], "FAIL", st["V19"])
        self.assertIn(name, st["V19"][1])
        self.assertEqual(rc, 1)

    def test_v09_covers_share_fuel_count_and_app_matrix_tables(self):
        st, ev = statuses(self.lines_ok)["V09"]
        self.assertEqual(st, "PASS", ev)
        # CA: share table, fuel split, Innova!B3, app matrix, US Benchmark share, Same-ASIN!B3; US: share, fuel, Innova!B3
        self.assertIn("+ 9 share/fuel/count/app-matrix tables", ev)

    def test_new_rule_fixture_values(self):
        """The cells the new V09 rules check hold the values computed here straight from the fixture CSVs."""
        ca = X.read_normalized_csv(CA_FIXTURE)
        us = X.read_normalized_csv(US_FIXTURE)
        cr = ca[ca["source_set"].isin(["code_reader", "both"])]
        core = _core(ca)
        a = core[core["source_set"].isin(["code_reader", "both"])]
        g = core[core["source_set"] == "gauge"]
        self.assertEqual((len(cr), round(cr["revenue_month"].sum(), 2), cr["units_month"].sum()), (13, 59706.13, 382))
        self.assertEqual((len(a), round(a["revenue_month"].sum(), 2), len(core), len(cr) + len(g)), (4, 10562.71, 9, 18))
        cells = _special_cells()
        self.assertEqual(cells["share"]["(a) Gauge devices inside the code-reader export"]["# ASINs"], 4)
        self.assertAlmostEqual(cells["share"]["(a) Gauge devices inside the code-reader export"]["Share of revenue"],
                               a["revenue_month"].sum() / cr["revenue_month"].sum(), places=12)
        den_rev = cr["revenue_month"].sum() + g["revenue_month"].sum()
        self.assertAlmostEqual(cells["share"]["(b) All core gauge devices (CR ∪ gauge export)"]["Share of revenue"],
                               core["revenue_month"].sum() / den_rev, places=12)
        self.assertEqual(cells["share"]["(b) denominator: CR total + gauge-only rows"]["# ASINs"], 18)
        # CA fixture has no fuel_scope column: every core device is 'unspecified'
        self.assertEqual(cells["fuel_total"]["unspecified: # ASINs"], 9)
        self.assertEqual(cells["fuel_total"]["gas: # ASINs"], 0)
        us_core = _core(us)
        us_cr = us[us["source_set"].isin(["code_reader", "both"])]
        us_a = us_core[us_core["source_set"].isin(["code_reader", "both"])]
        self.assertAlmostEqual(cells["bench"]["(a) Gauge devices inside the code-reader export: share of units"]["US share"],
                               us_a["units_month"].sum() / us_cr["units_month"].sum(), places=12)
        self.assertEqual(cells["innova_b3"], 0)
        self.assertEqual(cells["same_asin_b3"], 1)
        self.assertEqual(cells["apps"], list(C.APP_FEATURE_MATRIX_APPS))
        yoy = cr["yoy_units_pct"]
        self.assertEqual(cells["trend_total"]["# Listings with Helium 10 YoY data"], int(yoy.notna().sum()))
        self.assertEqual(cells["trend_total"]["# Listings with YoY > 0"], int((yoy > 0).sum()))
        self.assertEqual(cells["trend_total"]["# Listings with YoY < 0"], int((yoy < 0).sum()))
        self.assertEqual(cells["trend_total"]["Last Year Sales (Helium 10 field; semantics unverified)"],
                         cr["last_year_units"].dropna().sum())

    def test_tampered_special_tables_fail_v09(self):
        d_g, d_cr = ROOT / "tamper_special_gauge", ROOT / "tamper_special_cr"
        shutil.copytree(BUILD / "gauge_ca", d_g)
        shutil.copytree(BUILD / "cr", d_cr)
        reg = _registry("CA")
        gname = C.gauge_report_name("CA", MONTH)
        wb = load_workbook(d_g / gname)
        edits = []
        for t in reg["tables"]:
            if t["workbook"] != gname:
                continue
            ws = wb[t["sheet"]]
            if t["title"] == V.SHARE_TABLE_TITLE:
                c = ws.cell(t["first_data_row"] + 1, t["first_col"] + t["columns"].index("Share of revenue"))
                c.value = c.value + 0.01
                edits.append("Summary/kpi")
            elif t["title"] == V.FUEL_TABLE_TITLE:
                c = ws.cell(t["total_row"], t["first_col"] + t["columns"].index("unspecified: # ASINs"))
                c.value = c.value + 1
                edits.append("Summary/subtype_mix")
            elif t["title"] == V.BENCH_SHARE_TABLE_TITLE:
                c = ws.cell(t["first_data_row"], t["first_col"] + t["columns"].index("US share"))
                c.value = c.value + 0.02
                edits.append("US Benchmark/kpi")
            elif t["role"] == "kpi" and t["sheet"] == "Innova":
                ws.cell(t["first_data_row"], t["first_col"]).value = 5
                edits.append("Innova/kpi")
            elif t["role"] == "kpi" and t["sheet"] == "US vs CA Same-ASIN":
                ws.cell(t["first_data_row"], t["first_col"] + 1).value = 7
                edits.append("US vs CA Same-ASIN/kpi")
            elif t["role"] == "modelb_app_matrix":
                src = t["first_col"] + t["columns"].index("Source")
                ws.cell(t["first_data_row"], src).value = 3.5
                edits.append("App-Gauge Proxy (Model B)/modelb_app_matrix")
        self.assertEqual(len(edits), 6, edits)
        wb.save(d_g / gname)
        aname = C.cr_analysis_name("CA", MONTH)
        trend = next(t for t in reg["tables"] if t["workbook"] == aname and t["role"] == "trend_proxy")
        wb = load_workbook(d_cr / aname)
        c = wb["Trend Proxy"].cell(trend["total_row"], trend["first_col"] + trend["columns"].index("# Listings with YoY > 0"))
        c.value = c.value + 1
        wb.save(d_cr / aname)
        rc, lines = run_validator(cr=d_cr, gauge_ca=d_g, memo=write_memo("memo_ok6.md", fixed=True))
        st, ev = statuses(lines)["V09"]
        self.assertEqual(rc, 1)
        self.assertEqual(st, "FAIL", ev)
        self.assertIn("7 problem(s)", ev)
        for label in edits:
            self.assertIn(f"{gname}!{label}", ev)
        self.assertIn(f"{aname}!Trend Proxy/trend_proxy Total '# Listings with YoY > 0'", ev)

    def test_v04_checks_every_ranked_table_row(self):
        st, ev = statuses(self.lines_ok)["V04"]
        self.assertEqual(st, "PASS", ev)
        reg = _registry("CA")
        n_ranked = sum(1 for t in reg["tables"] if t["role"] in V.RANK_BY_REVENUE_ROLES + V.RANK_BY_UNITS_ROLES
                       or (t["role"] == "innova" and t["workbook"].startswith("CA_Code_Reader")))
        n_ranked += sum(1 for t in _registry("US")["tables"] if t["role"] in V.RANK_BY_REVENUE_ROLES + V.RANK_BY_UNITS_ROLES)
        self.assertIn(f"{n_ranked} ranked tables", ev)

    def test_swapped_top50_rows_fail_v04(self):
        d = ROOT / "tamper_swap"
        shutil.copytree(BUILD / "cr", d)
        t = next(t for t in _registry("CA")["tables"] if t["workbook"] == CR_REPORT and t["role"] == "top_by_units")
        wb = load_workbook(d / CR_REPORT)
        ws = wb[t["sheet"]]
        r1, r2 = t["first_data_row"], t["first_data_row"] + 1
        for c in range(t["first_col"], t["first_col"] + len(t["columns"])):
            if t["columns"][c - t["first_col"]] == "Ranking":
                continue
            a, b = ws.cell(r1, c), ws.cell(r2, c)
            a.value, b.value = b.value, a.value
        wb.save(d / CR_REPORT)
        rc, lines = run_validator(cr=d, memo=write_memo("memo_ok7.md", fixed=True))
        st, ev = statuses(lines)["V04"]
        self.assertEqual(rc, 1)
        self.assertEqual(st, "FAIL", ev)
        self.assertIn(f"{CR_REPORT}!Top 50/top_by_units (by units): 2 rows out of order", ev)

    def test_v19_raw_inputs_follow_the_loader_discovery(self):
        raw = ROOT / "raw_discovery"
        (raw / "extra").mkdir(parents=True)
        src = (C.FIXTURES_DIR / "cr_page1.csv").read_bytes()
        top = raw / "CA_AMAZON_blackBoxProducts_1_2026-10-02.csv"
        sub = raw / "extra" / "CA_AMAZON_blackBoxProducts_bullydog_2026-10-02.csv"
        top.write_bytes(src)
        sub.write_bytes(src)
        (raw / "extra" / "._CA_AMAZON_blackBoxProducts_bullydog_2026-10-02.csv").write_bytes(b"appledouble")
        self.assertEqual(V.expected_raw_inputs([raw]), {str(top.resolve()), str(sub.resolve())})
        only_top = {str(top.resolve()): "x"}
        self.assertEqual(V.missing_raw_inputs(only_top, [raw]), [str(sub.resolve())])
        self.assertEqual(V.missing_raw_inputs({**only_top, str(sub.resolve()): "y"}, [raw]), [])

    def test_json_under_new_product_dir_is_refused(self):
        bad = C.NEW_PRODUCT_DIR / "CA-OBD-GAUGE" / "outputs" / "validation.json"
        err = io.StringIO()
        with self.assertRaises(SystemExit) as cm, contextlib.redirect_stderr(err):
            V.main(["--month", MONTH, "--json", str(bad)])
        self.assertEqual(cm.exception.code, 2)
        self.assertIn("--json must not write under", err.getvalue())
        self.assertFalse(bad.exists())

    def test_skip_needs_a_documented_reason(self):
        rc, lines = run_validator(memo=write_memo("memo_bad2.md", fixed=False), extra=["--skip", "V20"])
        st = statuses(lines)
        self.assertEqual(st["V20"][0], "SKIP")
        self.assertIn(V.SKIPPABLE["V20"], st["V20"][1])
        self.assertEqual(rc, 0)
        with self.assertRaises(SystemExit), contextlib.redirect_stderr(io.StringIO()):
            V.main(["--month", MONTH, "--skip", "V01"])

    def test_missing_dataset_source_fails_explicitly(self):
        argv = ["--month", MONTH, "--cr-out-dir", str(BUILD / "cr"), "--gauge-out-dir", str(BUILD / "gauge_ca"),
                "--us-gauge-out-dir", str(BUILD / "gauge_us"), "--runs-dir", str(RUNS),
                "--raw-dir", str(ABSENT / "ca_cr"), "--gauge-raw-dir", str(ABSENT / "ca_gauge"),
                "--us-gauge-raw-dir", str(ABSENT / "us_gauge"), "--us-cr-raw-dir", str(ABSENT / "us_cr"),
                "--memo", str(write_memo("memo_ok5.md", fixed=True)), "--sources", str(_Built.sources)]
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            rc = V.main(argv)
        st = statuses(buf.getvalue().splitlines())
        self.assertEqual(rc, 1)
        self.assertEqual(st["V01"][0], "FAIL")
        self.assertIn("no dataset source", st["V01"][1])
        self.assertEqual(st["V10"][0], "PASS")      # workbook-only checks still run


# --------------------------------------------------------------------------------------
# Combined CA + US gauge workbook (ca_common "Combined CA + US gauge workbook" block)
# --------------------------------------------------------------------------------------
COMBINED = C.combined_gauge_report_name(MONTH)
CMB_SHORT = COMBINED.split("_202")[0]
CMB_ROOT = ROOT / "combined"
CMB_BUILD = CMB_ROOT / "build"          # cr / gauge_ca / gauge_us built from CA_FIXTURE + CMB_US_CSV
CMB_RUNS = CMB_BUILD / "runs"
CMB_OUT = CMB_ROOT / "gauge_ca"          # CA gauge out dir copy + the combined workbook + the merged manifest
CMB_US_CSV = CMB_ROOT / "normalized_rows_us_plus.csv"   # US fixture + one US-only core brand (no CA listings -> '-')
US_ONLY_ASIN, US_ONLY_KEY, US_ONLY_DISPLAY = "B0TESTUSX1", "autool", "Autool"
FUELS = tuple(C.FEATURE_FUEL_SCOPE)
KEY_FIGURE_LABELS = C.COMBINED_KEY_FIGURE_LABELS
TIER_LABELS = [t for t, _, _ in C.GAUGE_TIERS] + ["All tiers"]


def write_us_plus(path: Path) -> None:
    """US_FIXTURE plus one US-only core gauge display (revenue >= GAUGE_BRAND_TAB_MIN_REVENUE -> a brand tab)."""
    df = pd.read_csv(US_FIXTURE, dtype=str, keep_default_na=False)
    row = df[df["asin"] == "B0TESTUS09"].iloc[0].copy()
    for k, v in {"asin": US_ONLY_ASIN, "title": "Autool X50 Plus OBD2 gauge display", "brand_raw": "AUTOOL", "brand_key": US_ONLY_KEY,
                 "brand_display": US_ONLY_DISPLAY, "units_month": "30", "revenue_month": "1799.70", "price": "59.99",
                 "url": f"https://amazon.com/dp/{US_ONLY_ASIN}", "image_url": ""}.items():
        row[k] = v
    pd.concat([df, row.to_frame().T], ignore_index=True).to_csv(path, index=False)


def _fixture_ctx(ds, bench=None):
    """GCtx exactly as build_gauge_workbook assembles it on the normalized (dev) path."""
    market = C.MARKETS[ds.market]
    u, notes, fa = G.gauge_union_with_flags(ds, preclassified=True, gauge_map_path=G.GAUGE_MAP_DEFAULT, runs_dir=CMB_RUNS,
                                            rederive=False)
    c = G.GCtx(ds=ds, market=market, u=u, notes=notes, month=MONTH, mon=X.month_label(MONTH),
               sub=X.subtitle_text(market, ds.export_dates, MONTH), ccy=market.currency,
               cr_totals=G.code_reader_totals(ds.code_reader, ds.market), fuel_absent=fa)
    if bench is not None:
        c.bench = bench.ds
        c.bu, c.bench_cr_totals = bench.u, bench.cr_totals
        c.modelb = G.modelb_universe(ds, G.APP_GAUGE_BRANDS_DEFAULT)
        c.app_matrix, c.app_matrix_path = G.read_app_feature_matrix(G.APP_FEATURE_MATRIX_DEFAULT), G.APP_FEATURE_MATRIX_DEFAULT
    return c


def build_combined_fixture(out_dir: Path, runs_dir: Path, us_csv: Path) -> Path:
    """A minimal combined workbook written with the engine (write_table / Book) to the frozen spec, plus its CAUS registry
    and the merged manifest entry. Stands in for build_combined_gauge_report.py (parallel track) in these tests."""
    CA, US = C.MARKETS["CA"], C.MARKETS["US"]
    Col = X.ColumnSpec
    ca_ds = X.dataset_from_normalized(X.read_normalized_csv(CA_FIXTURE), "CA", MONTH)
    us_ds = X.dataset_from_normalized(X.read_normalized_csv(us_csv), "US", MONTH)
    uc = _fixture_ctx(us_ds)
    cc = _fixture_ctx(ca_ds, bench=uc)
    ctx = {"CA": cc, "US": uc}
    u = {mk: ctx[mk].u for mk in ctx}
    core = {mk: u[mk][u[mk]["_core"]] for mk in u}
    dev = {mk: u[mk][u[mk]["gauge_device_scope"]] for mk in u}
    book = X.Book(COMBINED, CA)

    def put(ws, spec, top, market=CA):
        tr = X.write_table(ws, spec, top, market, freeze=False)
        book.tables.append(tr)
        return tr

    def spec(sheet, role, title, cols, rows, total=None, residual=None, allowed=("CA", "US"), flt=""):
        rows = rows if isinstance(rows, pd.DataFrame) else pd.DataFrame(rows)
        return X.TableSpec(sheet, role, title, cols, rows, total, residual, allowed, dataset_filter=flt)

    def mcol(mk):
        return None if mk == "CA" else "US"

    T = {k: v[0] for k, v in C.COMBINED_SUMMARY_TITLES.items()}
    book.text(book.sheet("Read Me"), "CA + US OBD Gauge Competitor Report — Read Me", ["Combined fixture workbook."], role="read_me")
    ws = book.sheet("Summary")
    # Key figures (Measure | CA | US | Unit), the 11 COMBINED_KEY_FIGURE_LABELS rows: money rows carry the market money
    # format per cell; the folded gauge-share rows (definition (b)) are bold with a "0.00%" format

    def kf(mk):
        uu, co, dv = u[mk], core[mk], dev[mk]
        acc, adj = uu[uu["gauge_class"].isin(C.GAUGE_ACCESSORY_CLASSES)], uu[uu["gauge_class"].isin(C.GAUGE_ADJACENT_CLASSES)]
        b = next(d for d in G.market_share_rows(uu, ctx[mk].cr_totals, mk) if d["label"] == G.SHARE_ROW_LABELS[2])
        cr = ctx[mk].cr_totals
        return [float(x) for x in (co["revenue_month"].sum(), co["units_month"].sum(), co["asin"].nunique(),
                                   (co["units_month"] > 0).sum(), dv["revenue_month"].sum(), acc["revenue_month"].sum(),
                                   adj["revenue_month"].sum(), cr["rev"], cr["units"], b["s_rev"], b["s_u"])]
    money_rows, share_rows = {0, 4, 5, 6, 7}, {9, 10}
    tr = put(ws, spec("Summary", "kpi", T["key_figures"], [Col("Measure", "label", "text", 34), Col("CA", "ca", "rating"),
                                                             Col("US", "us", "rating", market="US"), Col("Unit", "unit", "text")],
                      {"label": KEY_FIGURE_LABELS, "ca": kf("CA"), "us": kf("US"),
                       "unit": ["CAD / USD" if i in money_rows else "share" if i in share_rows else "count"
                                for i in range(len(KEY_FIGURE_LABELS))]}, flt="see metric labels"), 1)
    for i, r in enumerate(range(tr.first_data_row, tr.last_data_row + 1)):
        for col, mk in ((2, CA), (3, US)):
            c = ws.cell(r, col)
            c.number_format = mk.money_fmt if i in money_rows else "0.00%" if i in share_rows else X.FMT_INT
            if i in share_rows:
                c.font = Font(bold=True)
    # Brand summary — CA vs US (core devices; top brands by max revenue, residual, Total)
    disp = {**dict(zip(core["US"]["brand_key"], core["US"]["brand_display"])),
            **dict(zip(core["CA"]["brand_key"], core["CA"]["brand_display"]))}
    tot_rev = {mk: float(core[mk]["revenue_month"].sum()) for mk in core}

    def agg(mk, sub):
        rev = float(sub["revenue_month"].sum())
        return {f"{mk}_n": int(sub["asin"].nunique()), f"{mk}_rev": rev, f"{mk}_units": float(sub["units_month"].sum()),
                f"{mk}_share": X.safe_div(rev, tot_rev[mk]), f"{mk}_rating": X.weighted_rating(sub)}
    maxrev = {k: max(float(core[mk].loc[core[mk]["brand_key"] == k, "revenue_month"].sum()) for mk in core) for k in disp}
    # the builder's documented order: top-N by max(CA, US) revenue, displayed by CA revenue desc, then US revenue desc
    brev = {mk: core[mk].groupby("brand_key")["revenue_month"].sum() for mk in core}
    chosen = sorted(disp, key=lambda k: (-maxrev[k], disp[k], k))[:C.SUMMARY_TOP_BRANDS]
    rest = sorted(set(disp) - set(chosen))
    shown = sorted(chosen, key=lambda k: (-float(brev["CA"].get(k, 0.0)), -float(brev["US"].get(k, 0.0)), disp[k], k))
    rows, brand_dash = [], []           # a market with no listings of the brand shows DASH in its five cells
    for i, k in enumerate(shown):
        d = {"brand": disp[k], "brand_key": k}
        for mk in core:
            sub = core[mk][core[mk]["brand_key"] == k]
            d.update(agg(mk, sub) if len(sub) else {f"{mk}_{f}": "-" for f in ("n", "rev", "units", "share", "rating")})
            if not len(sub):
                brand_dash.append((i, mk))
        rows.append(d)
    total = {"brand": C.TOTAL_ROW_LABEL}
    for mk in core:
        total.update(agg(mk, core[mk]))
    residual = None
    if rest:
        residual = {"brand": C.RESIDUAL_ROW_LABEL.format(noun="brands", n=len(rest))}
        for mk in core:
            residual.update(agg(mk, core[mk][core[mk]["brand_key"].isin(rest)]))
    cols = [Col("Brand", "brand", "text", 24)]
    for mk in ("CA", "US"):
        ccy = C.MARKETS[mk].currency
        cols += [Col(f"{mk} # of Listings", f"{mk}_n", "int"), Col(f"{mk} Monthly Rev ({ccy})", f"{mk}_rev", "money", market=mcol(mk)),
                 Col(f"{mk} Monthly Units", f"{mk}_units", "int"), Col(f"{mk} Rev Share", f"{mk}_share", "pct"),
                 Col(f"{mk} Avg Rating", f"{mk}_rating", "rating")]
    tr_b = put(ws, spec("Summary", "summary_brands", T["brands"], cols, rows, X.fit(total, cols), X.fit(residual, cols),
                        flt="core devices, both markets"), X.table_end(tr) + 3)
    book.bar(ws, tr_b, "Brand", "CA Monthly Rev (CAD)", "CA core device revenue by brand (CAD)")
    # Sub-type mix — CA vs US (device scope incl. borderline; the GPS-only HUD row sits outside the Total, share blank)

    def sub_stats(mk, sub, den, with_share=True):
        rev = float(sub["revenue_month"].sum())
        return {f"{mk}_n": int(sub["asin"].nunique()), f"{mk}_rev": rev, f"{mk}_units": float(sub["units_month"].sum()),
                f"{mk}_share": X.safe_div(rev, float(den["revenue_month"].sum())) if with_share else float("nan")}
    rows = []
    for cls in C.GAUGE_DEVICE_CLASSES + C.GAUGE_ADJACENT_CLASSES:
        d = {"label": C.GAUGE_SUBTYPE_LABELS[cls]}
        for mk in u:
            src = dev[mk] if cls in C.GAUGE_DEVICE_CLASSES else u[mk]
            d.update(sub_stats(mk, src[src["gauge_class"] == cls], dev[mk], with_share=cls in C.GAUGE_DEVICE_CLASSES))
        rows.append(d)
    total = {"label": C.TOTAL_ROW_LABEL}
    for mk in u:
        total.update(sub_stats(mk, dev[mk], dev[mk]))
    cols = [Col("Sub-type", "label", "text", 28)]
    for mk in ("CA", "US"):
        ccy = C.MARKETS[mk].currency
        cols += [Col(f"{mk} # ASINs", f"{mk}_n", "int"), Col(f"{mk} Monthly Rev ({ccy})", f"{mk}_rev", "money", market=mcol(mk)),
                 Col(f"{mk} Monthly Units", f"{mk}_units", "int"), Col(f"{mk} Rev share", f"{mk}_share", "pct")]
    tr = tr_sub = put(ws, spec("Summary", "subtype_mix", T["subtypes"], cols, rows, total, flt="gauge_device_scope"), X.table_end(tr_b) + 3)
    # Price tier × sub-type per market (core devices)
    for key, mk in (("tier_ca", "CA"), ("tier_us", "US")):
        ccy, co = C.MARKETS[mk].currency, core[mk]
        cols = [Col("Sub-type", "label", "text", 28)]
        for tier in TIER_LABELS:
            cols += [Col(f"{tier} Rev ({ccy})", f"{tier}|rev", "money"), Col(f"{tier} Units", f"{tier}|units", "int")]

        def tier_row(sub, label):
            d = {"label": label}
            for tier in TIER_LABELS:
                part = sub if tier == "All tiers" else sub[sub["_gtier"] == tier]
                d[f"{tier}|rev"], d[f"{tier}|units"] = float(part["revenue_month"].sum()), float(part["units_month"].sum())
            return d
        rows = [tier_row(co[co["gauge_class"] == cls], C.GAUGE_SUBTYPE_LABELS[cls]) for cls in C.GAUGE_DEVICE_CLASSES]
        tr = put(ws, spec("Summary", "tier_matrix", T[key], cols, rows, tier_row(co, C.TOTAL_ROW_LABEL), allowed=(mk,),
                          flt="core devices; tiers = GAUGE_TIERS on price (half-open)"), X.table_end(tr) + 3, market=C.MARKETS[mk])
    # Fuel split — CA vs US (core devices): per fuel a subtotal row then its sub-type rows; Total
    fuel = {mk: (pd.Series("unspecified", index=core[mk].index) if ctx[mk].fuel_absent else core[mk]["fuel_scope"].astype(str))
            for mk in core}

    def fuel_vals(d, f, cls):
        for mk in core:
            co = core[mk]
            m = pd.Series(True, index=co.index) if f is None else (fuel[mk] == f)
            if cls is not None:
                m &= co["gauge_class"] == cls
            d.update({f"{mk}_n": int(co.loc[m, "asin"].nunique()), f"{mk}_rev": float(co.loc[m, "revenue_month"].sum()),
                      f"{mk}_units": float(co.loc[m, "units_month"].sum())})
        return d
    # one label column: "<fuel> — subtotal", then that fuel's sub-type rows with >= 1 ASIN in either market
    rows, sub_idx = [], []
    for f in FUELS:
        sub_idx.append(len(rows))
        rows.append(fuel_vals({"label": f"{f} — subtotal"}, f, None))
        for cls in C.GAUGE_DEVICE_CLASSES:
            d = fuel_vals({"label": C.GAUGE_SUBTYPE_LABELS[cls]}, f, cls)
            if d["CA_n"] + d["US_n"] >= 1:
                rows.append(d)
    cols = [Col("Fuel / Sub-type", "label", "text", 28)]
    for mk in ("CA", "US"):
        cols += [Col(f"{mk} # ASINs", f"{mk}_n", "int"),
                 Col(f"{mk} Monthly Rev ({C.MARKETS[mk].currency})", f"{mk}_rev", "money", market=mcol(mk)),
                 Col(f"{mk} Units", f"{mk}_units", "int")]
    tr_f = put(ws, spec("Summary", "subtype_mix", T["fuel"], cols, rows, fuel_vals({"label": C.TOTAL_ROW_LABEL}, None, None),
                        flt="core devices, both markets; fuel_scope in FEATURE_FUEL_SCOPE"), X.table_end(tr) + 3)
    extras = {id(tr_f): {"subtotal_rows": [tr_f.first_data_row + i for i in sub_idx]},
              id(tr_sub): {"excluded_from_total_rows": [tr_sub.last_data_row]},
              id(tr_b): {"column_markets": [None] + [cbm for mk in ("CA", "US") for cbm in [mk] * 5],
                         "dash_rows": [[tr_b.first_data_row + i, mk] for i, mk in brand_dash]}}
    # Top 50 CA / Top 50 US
    for mk in ("CA", "US"):
        ws = book.sheet(f"Top 50 {mk}")
        cols = G._top50_columns(C.MARKETS[mk].currency)
        r = 1
        for role, by in (("top_by_revenue", "revenue"), ("top_by_units", "units")):
            shown_, residual_, total_ = X.top_listings(core[mk], by=by, n=C.TOP_N)
            tr = put(ws, spec(ws.title, role, f"Top {C.TOP_N} core gauge devices — Rank by {by.title()}", cols, shown_,
                              X.fit(total_, cols), X.fit(residual_, cols), allowed=(mk,), flt="core devices"), r, market=C.MARKETS[mk])
            r = X.table_end(tr) + 4
    # Innova: B3 (CA) and B4 (US) device counts
    ws = book.sheet("Innova")
    X.write_sheet_header(ws, "Innova — CA + US OBD gauge view", None, 4)
    for row, mk in ((3, "CA"), (4, "US")):
        n = int(((u[mk]["brand_key"] == "innova") & u[mk]["gauge_device_scope"]).sum())
        X.set_text(ws.cell(row, 1), f"Innova gauge/HUD device listings — {mk}")
        book.tables.append(X.write_number(ws, row, 2, n, C.MARKETS[mk], label=f"Innova gauge/HUD device listings — {mk}",
                                          dataset_filter="brand_key == 'innova' & gauge_device_scope"))
    book.brand_sheet_map["innova"] = "Innova"
    # Brand tabs: the brand-tab rule in EITHER market; KPI rows 3-6 (Metric | CA | US) then the four ranking tables
    tab_keys = set()
    for mk in u:
        co = core[mk]
        for key, g in dev[mk].groupby("brand_key"):
            if key != "innova" and (float(co.loc[co["brand_key"] == key, "revenue_month"].sum()) >= C.GAUGE_BRAND_TAB_MIN_REVENUE
                                    or g["asin"].nunique() >= C.GAUGE_BRAND_TAB_MIN_ASINS):
                tab_keys.add(key)
    taken: set[str] = set()
    for key in sorted(tab_keys, key=lambda k: (-maxrev.get(k, 0.0), k)):
        name = C.sheet_name_for_brand(disp[key], taken)
        ws = book.sheet(name)
        X.write_sheet_header(ws, f"{disp[key]} — CA + US core gauge devices", None, 3)
        for j, mk in enumerate(("CA", "US")):
            X.set_text(ws.cell(2, 2 + j), mk)                          # KPI column labels B2 / C2
        flt = f"core devices & brand_key == {key!r}"
        for i, (label, kind) in enumerate((("Monthly Rev (CAD | USD)", "money"), ("Monthly Units", "int"), ("# of Listings", "int"),
                                           ("Rev share within market", "pct"))):
            X.set_text(ws.cell(3 + i, 1), label)
            for j, mk in enumerate(("CA", "US")):
                rows_ = core[mk][core[mk]["brand_key"] == key]
                rev = float(rows_["revenue_month"].sum())
                c = ws.cell(3 + i, 2 + j)
                if not len(rows_):
                    X.set_text(c, "-")
                    continue
                c.value = [rev, float(rows_["units_month"].sum()), int(rows_["asin"].nunique()), X.safe_div(rev, tot_rev[mk])][i]
                c.number_format = C.MARKETS[mk].money_fmt if kind == "money" else X.FMT_PCT if kind == "pct" else X.FMT_INT
        kpi_tr = X.TableRange(sheet=name, role="kpi", header_row=None, first_data_row=3, last_data_row=6, total_row=None,
                              residual_row=None, first_col=1, last_col=3, columns=["Metric", "CA", "US"], charts=[],
                              title="", dataset_filter=flt, allowed_markets=("CA", "US"))
        book.tables.append(kpi_tr)
        extras[id(kpi_tr)] = {"dash_rows": [[r_, mk] for mk in ("CA", "US")
                                            if not len(core[mk][core[mk]["brand_key"] == key]) for r_ in range(3, 7)]}
        r = C.BRAND_TAB_RESERVED_ROWS + 1
        titles = iter(C.COMBINED_BRAND_TABLE_TITLES)
        for mk in ("CA", "US"):
            cols = X.brand_tab_columns(C.MARKETS[mk].currency, type_header="Sub-type", type_field="_subtype")
            rows_ = core[mk][core[mk]["brand_key"] == key]
            total_ = X.fit(X.listing_totals(rows_, "title", C.TOTAL_ROW_LABEL), cols)
            if not len(rows_):
                total_.update({"revenue_month": "-", "units_month": "-",       # absent market: '-' on the Total row,
                               "review_count": None, "rating": None})          # blank in every other numeric column
            for role, by in (("brand_tab_revenue", "revenue"), ("brand_tab_units", "units")):
                data = X.rank_listings(rows_, by) if len(rows_) else pd.DataFrame(
                    [{**{c.field: None for c in cols}, "title": f"No {mk} core gauge devices for this brand"}])
                tr = put(ws, spec(name, role, next(titles), cols, data, total_, allowed=(mk,), flt=flt), r, market=C.MARKETS[mk])
                extras[id(tr)] = {"dash_rows": [] if len(rows_) else [[tr.total_row, mk]],
                                  "placeholder_rows": [] if len(rows_) else [tr.first_data_row]}
                r = X.table_end(tr) + 4
        book.brand_sheet_map[key] = name
    # CA analyses, unchanged (Model A/B), and the Same-ASIN sheet
    G._price_ladder(book, cc)
    G._feature_matrix(book, cc)
    G._modelb(book, cc)
    G._same_asin(book, cc)
    tr_sa = book.tables[-1]
    assert tr_sa.role == "same_asin", tr_sa.role
    extras[id(tr_sa)] = {"column_markets": [h[-2:] if h in ("Link CA", "Link US") else h[:2] if h[:3] in ("CA ", "US ") else None
                                            for h in tr_sa.columns]}
    # All Products / Dedupe & Classification Audit / Excluded: "<Sheet> — CA" then "<Sheet> — US"
    ws = book.sheet("All Products")
    r = 1
    for mk in ("CA", "US"):
        ccy = C.MARKETS[mk].currency
        cols = [Col("ASIN", "asin", "text", 13), Col("Product Name", "title", "text", 60), Col("Brand", "brand_display", "text", 16),
                Col("Gauge Class", "gauge_class", "text", 24), Col(f"Price ({ccy})", "price", "money2", 12),
                Col(f"Monthly Rev ({ccy})", "revenue_month", "money", 15), Col("Monthly Units", "units_month", "int", 12),
                Col("Gauge Tier", "_gtier", "text", 11), Col("URL", "_url", "text", 34), Col("Link", "asin", "link", 24)]
        tr = put(ws, spec("All Products", "all_rows", f"All Products — {mk}", cols, X.rank_listings(u[mk], "revenue"),
                          X.fit(X.listing_totals(u[mk], "asin", C.TOTAL_ROW_LABEL), cols), allowed=(mk,),
                          flt="code-reader ∪ gauge union (all classes)"), r, market=C.MARKETS[mk])
        r = X.table_end(tr) + 3
    ws = book.sheet("Dedupe & Classification Audit")
    r = 1
    for mk in ("CA", "US"):
        cols = [Col("ASIN", "asin", "text", 13), Col("Gauge Class", "gauge_class", "text", 24), Col("Rule ID", "gauge_rule_id", "text", 14)]
        tr = put(ws, spec(ws.title, "dedupe_audit", f"Dedupe & Classification Audit — {mk}", cols,
                          u[mk].sort_values(["gauge_class", "asin"], kind="mergesort"), allowed=(mk,),
                          flt="classification decisions (all union rows)"), r, market=C.MARKETS[mk])
        r = X.table_end(tr) + 3
    ws = book.sheet("Excluded")
    r = 1
    for mk in ("CA", "US"):
        ccy = C.MARKETS[mk].currency
        ex = u[mk][u[mk]["gauge_class"].isin(C.GAUGE_EXCLUDED_CLASSES + ("ambiguous",))]
        cols = [Col("ASIN", "asin", "text", 13), Col("Product Name", "title", "text", 60), Col("Gauge Class", "gauge_class", "text", 22),
                Col("Rule ID", "gauge_rule_id", "text", 14), Col(f"Monthly Rev ({ccy})", "revenue_month", "money", 15),
                Col("Monthly Units", "units_month", "int", 12), Col("URL", "_url", "text", 34), Col("Link", "asin", "link", 24)]
        tr = put(ws, spec("Excluded", "excluded", f"Excluded — {mk}", cols, X.rank_listings(ex, "revenue"),
                          X.fit(X.listing_totals(ex, "asin", C.TOTAL_ROW_LABEL), cols), allowed=(mk,),
                          flt="gauge_class in GAUGE_EXCLUDED_CLASSES + ambiguous"), r, market=C.MARKETS[mk])
        r = X.table_end(tr) + 3
    book.text(book.sheet("Source & Method"), "Source & Method", ["Combined fixture."], role="source_method")
    meta = {k: f"fixture {k}" for k in C.METADATA_REQUIRED_KEYS}
    meta.update({"Marketplace": "amazon.ca (CA) | amazon.com (US)", "Currency": "CA: CAD | US: USD (no FX conversion)"})
    book.metadata(book.sheet("Metadata"), meta)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / COMBINED
    book.save(path)
    reg = {"market": "CAUS", "month": MONTH,
           "workbooks": {COMBINED: {"sheets": book.sheetnames, "brand_sheet_map": dict(book.brand_sheet_map)}},
           "tables": [{**t.to_registry(COMBINED, book.brand_sheet_map), **extras.get(id(t), {})} for t in book.tables]}
    rp = C.run_file(runs_dir, MONTH, "table_registry", "CAUS", "json")
    rp.parent.mkdir(parents=True, exist_ok=True)
    rp.write_text(json.dumps(reg, indent=1, ensure_ascii=False), encoding="utf-8")
    with contextlib.redirect_stdout(io.StringIO()):
        X.write_manifest(out_dir, MONTH, [path], X.input_hashes([CA_FIXTURE, us_csv]))
    return path


class _Combined:
    done = False

    @classmethod
    def ensure(cls) -> None:
        _Built.ensure()
        if cls.done:
            return
        if CMB_ROOT.exists():
            shutil.rmtree(CMB_ROOT)
        CMB_ROOT.mkdir(parents=True)
        write_us_plus(CMB_US_CSV)
        build_set(CMB_BUILD, CMB_RUNS, CA_FIXTURE, CMB_US_CSV)
        shutil.copytree(CMB_BUILD / "gauge_ca", CMB_OUT)
        cls.path = build_combined_fixture(CMB_OUT, CMB_RUNS, CMB_US_CSV)
        cls.reg = json.loads(C.run_file(CMB_RUNS, MONTH, "table_registry", "CAUS", "json").read_text(encoding="utf-8"))
        cls.done = True

    @classmethod
    def table(cls, *, title: str | None = None, sheet: str | None = None, role: str | None = None) -> dict:
        ts = [t for t in cls.reg["tables"] if (title is None or t["title"] == title) and (sheet is None or t["sheet"] == sheet)
              and (role is None or t["role"] == role)]
        assert len(ts) == 1, (title, sheet, role, len(ts))
        return ts[0]


def tampered_copy(name: str) -> Path:
    d = CMB_ROOT / name
    if d.exists():
        shutil.rmtree(d)
    shutil.copytree(CMB_OUT, d)
    return d


def run_combined(**kw) -> tuple[int, list[str]]:
    """run_validator over the combined test set (CMB_BUILD workbooks, CMB_RUNS, the US-plus fixture)."""
    kw.setdefault("cr", CMB_BUILD / "cr")
    kw.setdefault("gauge_us", CMB_BUILD / "gauge_us")
    kw.setdefault("runs", CMB_RUNS)
    return run_validator(us_csv=CMB_US_CSV, **kw)


class CombinedWorkbookTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _Combined.ensure()
        cls.memo = write_memo("memo_combined.md", fixed=True)
        cls.rc, cls.lines = run_combined(gauge_ca=CMB_OUT, memo=cls.memo)
        cls.st = statuses(cls.lines)

    def _run(self, gauge_ca: Path, **kw) -> tuple[int, dict]:
        rc, lines = run_combined(gauge_ca=gauge_ca, memo=self.memo, **kw)
        self.assertEqual(len(lines), 24, "\n".join(lines))
        return rc, statuses(lines)

    def test_combined_workbook_is_picked_up_and_passes(self):
        failed = {k: v for k, v in self.st.items() if v[0] == "FAIL"}
        self.assertEqual(failed, {}, "\n".join(self.lines))
        self.assertEqual(self.rc, 0)
        self.assertEqual(len(self.lines), 24)
        self.assertRegex(self.lines[-1], r"^VALIDATION: PASS \(22/23\)$")
        for cid in ("V01", "V02", "V03", "V04", "V05", "V06", "V07", "V09", "V12", "V13", "V15", "V16", "V18", "V19", "V23"):
            self.assertIn(CMB_SHORT, self.st[cid][1], f"{cid}: {self.st[cid]}")
        self.assertIn("combined: present", self.st["V15"][1])
        self.assertIn("Top 50 US", self.st["V04"][1])

    def test_absent_combined_workbook_is_not_a_failure(self):
        rc, st = self._run(CMB_BUILD / "gauge_ca")
        self.assertEqual(rc, 0, st)
        self.assertEqual(st["V15"][0], "PASS")
        self.assertIn("combined: absent", st["V15"][1])
        self.assertNotIn(CMB_SHORT, st["V09"][1])

    def test_required_but_absent_combined_workbook_fails(self):
        rc, st = self._run(CMB_BUILD / "gauge_ca", extra=["--combined", "require"])
        self.assertEqual(rc, 1)
        self.assertEqual(st["V15"][0], "FAIL", st["V15"])
        self.assertIn(COMBINED, st["V15"][1])

    def test_tampered_share_cell_fails_v09(self):
        """The folded gauge-share rows of Key figures: a tampered value, a lost bold and a wrong format each FAIL V09
        (naming Key figures); V12 checks only the scope rows and stays PASS."""
        d = tampered_copy("tamper_share")
        title = C.COMBINED_SUMMARY_TITLES["key_figures"][0]
        t = _Combined.table(title=title, sheet="Summary")
        self.assertNotIn("Gauge share of the code-reader market", {e["title"] for e in _Combined.reg["tables"]})
        wb = load_workbook(d / COMBINED)
        ws = wb["Summary"]
        rows = {ws.cell(r, t["first_col"]).value: r for r in range(t["first_data_row"], t["last_data_row"] + 1)}
        self.assertEqual(list(rows), list(C.COMBINED_KEY_FIGURE_LABELS))
        rev_lbl, units_lbl = C.COMBINED_KEY_FIGURE_SHARE_LABELS
        us_col, ca_col = t["first_col"] + t["columns"].index("US"), t["first_col"] + t["columns"].index("CA")
        c = ws.cell(rows[rev_lbl], us_col)
        c.value = c.value + 0.01
        ws.cell(rows[units_lbl], ca_col).font = Font(bold=False)
        ws.cell(rows[units_lbl], us_col).number_format = "0.0%"
        wb.save(d / COMBINED)
        rc, st = self._run(d)
        self.assertEqual(rc, 1)
        st9, ev9 = st["V09"]
        self.assertEqual(st9, "FAIL", ev9)
        self.assertIn("3 problem(s)", ev9)
        self.assertIn(f"{COMBINED}!Summary/kpi[{title}] row {rows[rev_lbl]} 'US'", ev9)
        self.assertIn(f"share row {units_lbl!r} (CA) is not bold", ev9)
        self.assertIn(f"share row {units_lbl!r} (US) format '0.0%'", ev9)
        self.assertEqual(st["V12"][0], "PASS", st["V12"])

    def test_tampered_tier_units_cell_fails_v09(self):
        d = tampered_copy("tamper_tier")
        title = C.COMBINED_SUMMARY_TITLES["tier_us"][0]
        t = _Combined.table(title=title)
        h = f"{C.GAUGE_TIERS[2][0]} Units"
        wb = load_workbook(d / COMBINED)
        c = wb["Summary"].cell(t["first_data_row"], t["first_col"] + t["columns"].index(h))
        c.value = (c.value or 0) + 3
        wb.save(d / COMBINED)
        rc, st = self._run(d)
        self.assertEqual(rc, 1)
        self.assertEqual(st["V09"][0], "FAIL", st["V09"])
        self.assertIn(f"{COMBINED}!Summary/tier_matrix[{title}]", st["V09"][1])
        self.assertIn(repr(h), st["V09"][1])

    def test_swapped_top50_us_rows_fail_v04(self):
        d = tampered_copy("tamper_top50_us")
        t = _Combined.table(sheet="Top 50 US", role="top_by_revenue")
        wb = load_workbook(d / COMBINED)
        ws = wb["Top 50 US"]
        r1, r2 = t["first_data_row"], t["first_data_row"] + 1
        for j, h in enumerate(t["columns"]):
            if h == "Ranking":
                continue
            a, b = ws.cell(r1, t["first_col"] + j), ws.cell(r2, t["first_col"] + j)
            a.value, b.value = b.value, a.value
        wb.save(d / COMBINED)
        rc, st = self._run(d)
        self.assertEqual(rc, 1)
        self.assertEqual(st["V04"][0], "FAIL", st["V04"])
        self.assertIn(f"{COMBINED}!Top 50 US/top_by_revenue", st["V04"][1])
        self.assertIn("2 rows out of order", st["V04"][1])
        self.assertEqual(st["V01"][0], "PASS", st["V01"])

    def test_wrong_currency_header_in_us_block_fails_v07(self):
        d = tampered_copy("tamper_ccy")
        runs = CMB_ROOT / "tamper_ccy_runs"
        if runs.exists():
            shutil.rmtree(runs)
        shutil.copytree(CMB_RUNS, runs)
        title = C.COMBINED_SUMMARY_TITLES["brands"][0]
        t = _Combined.table(title=title)
        old, new = "US Monthly Rev (USD)", "US Monthly Rev (CAD)"
        wb = load_workbook(d / COMBINED)
        wb["Summary"].cell(t["header_row"], t["first_col"] + t["columns"].index(old)).value = new
        wb.save(d / COMBINED)
        rp = C.run_file(runs, MONTH, "table_registry", "CAUS", "json")
        reg = json.loads(rp.read_text(encoding="utf-8"))
        for e in reg["tables"]:
            if e["title"] == title:
                e["columns"] = [new if h == old else h for h in e["columns"]]
        rp.write_text(json.dumps(reg, indent=1, ensure_ascii=False), encoding="utf-8")
        rc, st = self._run(d, runs=runs)
        self.assertEqual(rc, 1)
        self.assertEqual(st["V07"][0], "FAIL", st["V07"])
        self.assertIn(f"{COMBINED}!Summary/summary_brands[{title}]", st["V07"][1])
        self.assertIn(repr(new), st["V07"][1])

    def test_tampered_brand_share_fails_v05(self):
        d = tampered_copy("tamper_brand_share")
        title = C.COMBINED_SUMMARY_TITLES["brands"][0]
        t = _Combined.table(title=title)
        wb = load_workbook(d / COMBINED)
        c = wb["Summary"].cell(t["first_data_row"], t["first_col"] + t["columns"].index("US Rev Share"))
        c.value = c.value + 0.05
        wb.save(d / COMBINED)
        rc, st = self._run(d)
        self.assertEqual(rc, 1)
        self.assertEqual(st["V05"][0], "FAIL", st["V05"])
        self.assertIn(f"{COMBINED}!Summary/summary_brands[{title}] 'US Rev Share'", st["V05"][1])
        self.assertNotIn("'CA Rev Share'", st["V05"][1])

    def test_tampered_fuel_kpi_innova_and_adjacent_cells_fail_v09_and_v12(self):
        d = tampered_copy("tamper_misc")
        T = C.COMBINED_SUMMARY_TITLES
        wb = load_workbook(d / COMBINED)
        ws = wb["Summary"]
        fuel = _Combined.table(title=T["fuel"][0])
        sub_r = next(r for r in range(fuel["first_data_row"], fuel["last_data_row"] + 1)
                     if ws.cell(r, fuel["first_col"]).value == "diesel-capable — subtotal")
        c = ws.cell(sub_r, fuel["first_col"] + fuel["columns"].index("US Units"))
        c.value = c.value + 1
        kf = _Combined.table(title=T["key_figures"][0])
        acc_r = next(r for r in range(kf["first_data_row"], kf["last_data_row"] + 1)
                     if ws.cell(r, kf["first_col"]).value == "Accessories revenue")
        c = ws.cell(acc_r, kf["first_col"] + kf["columns"].index("US"))
        c.value = c.value + 10
        st_t = _Combined.table(title=T["subtypes"][0])
        adj_r = next(r for r in range(st_t["first_data_row"], st_t["last_data_row"] + 1)
                     if ws.cell(r, st_t["first_col"]).value == C.GAUGE_SUBTYPE_LABELS["gps_hud"])
        ws.cell(adj_r, st_t["first_col"] + st_t["columns"].index("CA Rev share")).value = 0.1
        wb["Innova"]["B4"].value = 2
        brand_sheet = _Combined.reg["workbooks"][COMBINED]["brand_sheet_map"]["scangauge"]
        wb[brand_sheet]["C4"].value = wb[brand_sheet]["C4"].value + 5          # US Monthly Units KPI
        wb.save(d / COMBINED)
        rc, st = self._run(d)
        self.assertEqual(rc, 1)
        st9, ev9 = st["V09"]
        self.assertEqual(st9, "FAIL", ev9)
        self.assertIn("5 problem(s)", ev9)
        self.assertIn(f"{COMBINED}!Summary/subtype_mix[{T['fuel'][0]}] row {sub_r} 'US Units'", ev9)
        self.assertIn(f"{COMBINED}!Summary/kpi[{T['key_figures'][0]}] row {acc_r} 'US'", ev9)
        self.assertIn(f"{COMBINED}!Summary/subtype_mix[{T['subtypes'][0]}] row {adj_r} 'CA Rev share'", ev9)
        self.assertIn(f"{COMBINED}!Innova!B4", ev9)
        self.assertIn(f"{COMBINED}!{brand_sheet}/kpi[] row 4 'US'", ev9)
        st12, ev12 = st["V12"]
        self.assertEqual(st12, "FAIL", ev12)
        self.assertIn("2 problem(s)", ev12)       # Key figures + Innova B4
        self.assertEqual(st["V01"][0], "PASS", st["V01"])
        self.assertEqual(st["V05"][0], "PASS", st["V05"])   # the adjacent row is outside the share sum

    def test_fuel_leaf_rows_and_registry_extras_are_checked(self):
        d = tampered_copy("tamper_fuel_leaf")
        runs = CMB_ROOT / "tamper_fuel_leaf_runs"
        if runs.exists():
            shutil.rmtree(runs)
        shutil.copytree(CMB_RUNS, runs)
        T = C.COMBINED_SUMMARY_TITLES
        fuel = _Combined.table(title=T["fuel"][0])
        wb = load_workbook(d / COMBINED)
        ws = wb["Summary"]
        labels = {r: ws.cell(r, fuel["first_col"]).value for r in range(fuel["first_data_row"], fuel["last_data_row"] + 1)}
        sub_r = next(r for r, v in labels.items() if v == "diesel-capable — subtotal")
        leaf_r = sub_r + 1                                  # the first sub-type row of the diesel-capable group
        self.assertIn(labels[leaf_r], {C.GAUGE_SUBTYPE_LABELS[c] for c in C.GAUGE_DEVICE_CLASSES})
        c = ws.cell(leaf_r, fuel["first_col"] + fuel["columns"].index("US # ASINs"))
        c.value = c.value + 1
        wb.save(d / COMBINED)
        rp = C.run_file(runs, MONTH, "table_registry", "CAUS", "json")
        reg = json.loads(rp.read_text(encoding="utf-8"))
        for e in reg["tables"]:
            if e["title"] == T["fuel"][0]:
                e["subtotal_rows"] = e["subtotal_rows"][1:]
            if e["title"] == T["subtypes"][0]:
                e["excluded_from_total_rows"] = []
        rp.write_text(json.dumps(reg, indent=1, ensure_ascii=False), encoding="utf-8")
        rc, st = self._run(d, runs=runs)
        self.assertEqual(rc, 1)
        st9, ev9 = st["V09"]
        self.assertEqual(st9, "FAIL", ev9)
        self.assertIn("3 problem(s)", ev9)
        self.assertIn(f"{COMBINED}!Summary/subtype_mix[{T['fuel'][0]}] row {leaf_r} 'US # ASINs'", ev9)
        self.assertIn("registry subtotal_rows", ev9)
        self.assertIn("registry excluded_from_total_rows []", ev9)

    def test_dash_cells_only_where_the_market_has_no_listings(self):
        """'-' marks a brand with no listings in a market. A legitimate '-' replaced by 0, a '-' on a brand with
        listings (brand summary and a brand-tab Total) and a drifted dash_rows extra each FAIL V09; V05 / V07 skip '-'."""
        d = tampered_copy("tamper_dash")
        runs = CMB_ROOT / "tamper_dash_runs"
        if runs.exists():
            shutil.rmtree(runs)
        shutil.copytree(CMB_RUNS, runs)
        btitle = C.COMBINED_SUMMARY_TITLES["brands"][0]
        bt = _Combined.table(title=btitle)
        wb = load_workbook(d / COMBINED)
        ws = wb["Summary"]
        brow = {ws.cell(r, bt["first_col"]).value: r for r in range(bt["first_data_row"], bt["last_data_row"] + 1)}
        a_cell = ws.cell(brow[US_ONLY_DISPLAY], bt["first_col"] + bt["columns"].index("CA # of Listings"))
        self.assertEqual(a_cell.value, "-")                      # the fixture shows '-' for the US-only brand's CA block
        a_cell.value = 0
        s_cell = ws.cell(brow["ScanGauge"], bt["first_col"] + bt["columns"].index("US Monthly Units"))
        s_cell.value = "-"
        sheet = _Combined.reg["workbooks"][COMBINED]["brand_sheet_map"]["scangauge"]
        rt = _Combined.table(sheet=sheet, title="CA — Rank by Revenue")
        t_cell = wb[sheet].cell(rt["total_row"], rt["first_col"] + rt["columns"].index("Monthly Rev (CAD)"))
        t_cell.value = "-"
        wb.save(d / COMBINED)
        rp = C.run_file(runs, MONTH, "table_registry", "CAUS", "json")
        reg = json.loads(rp.read_text(encoding="utf-8"))
        a_sheet = reg["workbooks"][COMBINED]["brand_sheet_map"][US_ONLY_KEY]
        for e in reg["tables"]:
            if e["sheet"] == a_sheet and e["role"] == "kpi":
                e["dash_rows"] = e["dash_rows"][:-1]
        rp.write_text(json.dumps(reg, indent=1, ensure_ascii=False), encoding="utf-8")
        rc, st = self._run(d, runs=runs)
        self.assertEqual(rc, 1)
        st9, ev9 = st["V09"]
        self.assertEqual(st9, "FAIL", ev9)
        self.assertIn("4 problem(s)", ev9)
        self.assertIn(f"{COMBINED}!Summary/summary_brands[{btitle}] {a_cell.coordinate} 'CA # of Listings': 0 where '-' expected", ev9)
        self.assertIn(f"{COMBINED}!Summary/summary_brands[{btitle}] {s_cell.coordinate} 'US Monthly Units': '-' where listings exist",
                      ev9)
        self.assertIn(f"{COMBINED}!{sheet}/brand_tab_revenue[CA — Rank by Revenue] Total 'Monthly Rev (CAD)': '-' where listings exist",
                      ev9)
        self.assertIn(f"{COMBINED}!{a_sheet}/kpi[]: registry dash_rows", ev9)
        self.assertEqual(st["V05"][0], "PASS", st["V05"])
        self.assertEqual(st["V07"][0], "PASS", st["V07"])

    def _tamper_registry(self, name: str, edit) -> Path:
        runs = CMB_ROOT / f"{name}_runs"
        if runs.exists():
            shutil.rmtree(runs)
        shutil.copytree(CMB_RUNS, runs)
        rp = C.run_file(runs, MONTH, "table_registry", "CAUS", "json")
        reg = json.loads(rp.read_text(encoding="utf-8"))
        edit(reg)
        rp.write_text(json.dumps(reg, indent=1, ensure_ascii=False), encoding="utf-8")
        return runs

    def test_money_header_without_currency_token_fails_v07(self):
        """A 'US ' block prefix is not a currency label: every money header carries (CAD)/(USD) (only the Key figures /
        brand-tab KPI CA | US value columns are exempt, their rows name the currency)."""
        d = tampered_copy("tamper_no_token")
        title = C.COMBINED_SUMMARY_TITLES["brands"][0]
        t = _Combined.table(title=title)
        old, new = "US Monthly Rev (USD)", "US Monthly Rev"
        wb = load_workbook(d / COMBINED)
        wb["Summary"].cell(t["header_row"], t["first_col"] + t["columns"].index(old)).value = new

        def edit(reg):
            for e in reg["tables"]:
                if e["title"] == title:
                    e["columns"] = [new if h == old else h for h in e["columns"]]
        wb.save(d / COMBINED)
        rc, st = self._run(d, runs=self._tamper_registry("tamper_no_token", edit))
        self.assertEqual(rc, 1)
        st7, ev7 = st["V07"]
        self.assertEqual(st7, "FAIL", ev7)
        self.assertIn(f"{COMBINED}!Summary/summary_brands[{title}]: money column 'US Monthly Rev' lacks (USD)", ev7)
        self.assertIn("1 problem(s)", ev7)

    def test_swapped_brand_rows_fail_v09(self):
        """Brand summary rows are compared by position with the re-derived key order (CA revenue desc, then US)."""
        d = tampered_copy("tamper_brand_swap")
        title = C.COMBINED_SUMMARY_TITLES["brands"][0]
        t = _Combined.table(title=title)
        wb = load_workbook(d / COMBINED)
        ws = wb["Summary"]
        r1, r2 = t["first_data_row"], t["first_data_row"] + 1
        a, b = ws.cell(r1, t["first_col"]).value, ws.cell(r2, t["first_col"]).value
        for j in range(len(t["columns"])):
            c1, c2 = ws.cell(r1, t["first_col"] + j), ws.cell(r2, t["first_col"] + j)
            c1.value, c2.value = c2.value, c1.value
        wb.save(d / COMBINED)
        rc, st = self._run(d)
        self.assertEqual(rc, 1)
        st9, ev9 = st["V09"]
        self.assertEqual(st9, "FAIL", ev9)
        self.assertIn(f"{COMBINED}!Summary/summary_brands[{title}] row {r1}: brand {b!r} != expected {a!r} at this position", ev9)
        self.assertIn(f"row {r2}: brand {a!r} != expected {b!r}", ev9)

    def test_brand_display_collision_is_explicit(self):
        v = V.Validator(V.parse_args(["--month", MONTH]))
        frames = {"CA": pd.DataFrame({"brand_key": ["acme", "acme inc"], "brand_display": ["Acme", "Acme"],
                                      "revenue_month": [10.0, 5.0]}),
                  "US": pd.DataFrame({"brand_key": ["acme"], "brand_display": ["Acme"], "revenue_month": [7.0]})}
        shown, rest, disp, probs = v.expected_brand_rows(frames)
        self.assertEqual(shown, ["acme", "acme inc"])
        self.assertEqual(probs, ["brand keys ['acme', 'acme inc'] share the display label 'Acme': their Brand summary rows "
                                 "are ambiguous"])

    def test_brand_tab_kpi_block_structure_fails_v09(self):
        """Every brand tab: B2/C2 = CA/US, rows 3-6 labelled in order, one registered block over rows 3-6."""
        d = tampered_copy("tamper_kpi_block")
        sheet = _Combined.reg["workbooks"][COMBINED]["brand_sheet_map"]["keenso"]
        wb = load_workbook(d / COMBINED)
        ws = wb[sheet]
        for col in (1, 2, 3):
            ws.cell(6, col).value = None                       # the 'Rev share' row removed
        ws.cell(2, 3).value = None                             # the US column label removed
        wb.save(d / COMBINED)

        def edit(reg):
            for e in reg["tables"]:
                if e["sheet"] == sheet and e["role"] == "kpi":
                    e["last_data_row"] = 5                     # registry shortened to rows 3-5
        rc, st = self._run(d, runs=self._tamper_registry("tamper_kpi_block", edit))
        self.assertEqual(rc, 1)
        st9, ev9 = st["V09"]
        self.assertEqual(st9, "FAIL", ev9)
        self.assertIn(f"{COMBINED}!{sheet}: KPI column labels B2/C2 ['CA', None] != ['CA', 'US']", ev9)
        self.assertIn(f"{COMBINED}!{sheet}: KPI rows 3-6 labels", ev9)
        self.assertIn(f"{COMBINED}!{sheet}: expected one registered KPI block over rows 3-6", ev9)

    def test_absent_market_total_blank_ok_present_market_blank_fails(self):
        """Absent market (Autool CA): '-' or blank in every numeric Total column passes (clean run); a blank on a
        present market's Total ('# of Reviews') FAILs V09."""
        self.assertEqual(self.st["V09"][0], "PASS", self.st["V09"])
        a_sheet = _Combined.reg["workbooks"][COMBINED]["brand_sheet_map"][US_ONLY_KEY]
        at = _Combined.table(sheet=a_sheet, title="CA — Rank by Revenue")
        wb = load_workbook(_Combined.path)
        self.assertIsNone(wb[a_sheet].cell(at["total_row"], at["first_col"] + at["columns"].index("# of Reviews")).value)
        d = tampered_copy("tamper_total_blank")
        sheet = _Combined.reg["workbooks"][COMBINED]["brand_sheet_map"]["scangauge"]
        rt = _Combined.table(sheet=sheet, title="US — Rank by Units")
        wb = load_workbook(d / COMBINED)
        wb[sheet].cell(rt["total_row"], rt["first_col"] + rt["columns"].index("# of Reviews")).value = None
        wb.save(d / COMBINED)
        rc, st = self._run(d)
        self.assertEqual(rc, 1)
        st9, ev9 = st["V09"]
        self.assertEqual(st9, "FAIL", ev9)
        self.assertIn(f"{COMBINED}!{sheet}/brand_tab_units[US — Rank by Units] Total '# of Reviews' blank", ev9)
        self.assertIn("1 problem(s)", ev9)

    def test_combined_header_rules(self):
        probs = V.combined_header_problems
        self.assertEqual(probs("US Monthly Rev (USD)", "US"), [])
        self.assertEqual(probs("CA Rev", None), [])
        self.assertTrue(probs("US Monthly Rev (CAD)", None))                 # block prefix vs currency
        self.assertTrue(probs("Price (CAD)", "US"))                          # US table, CAD header
        self.assertTrue(probs("Rev ratio US/CA (CAD/USD)", None))            # cross-currency ratio
        self.assertTrue(probs("US/CA Rev ratio", None))                      # ratio over revenue without a currency
        self.assertEqual(probs("Units ratio US/CA", None), [])               # unit ratios are allowed


if __name__ == "__main__":
    unittest.main()
