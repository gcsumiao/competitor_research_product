"""Output validator for the CA code-reader + OBD gauge workbooks (checks V01-V23 of ca_common.VALIDATION_CHECKS).

  ca_market_reports/run.sh ca_market_reports/validate_outputs.py --month 202609
  dev/fixture: ... --from-normalized <ca csv> --us-from-normalized <us csv> --cr-out-dir ... --runs-dir ...

Prints one line per check, in V01..V23 order:
  PASS V01 <short name>: <evidence> | FAIL V01 <short name>: <evidence> | SKIP V01 <short name>: <reason>
then the final line (ca_common.VALIDATION_FINAL_RE): "VALIDATION: PASS (<n passed>/23)" or
"VALIDATION: FAIL (<n passed>/23; failed: V05, V07)". Exit 0 only when no check FAILs. A check that raises is a FAIL with the
exception text; nothing escapes main().

Independence: the dataset is RE-DERIVED from the raw Helium 10 CSVs (ca_load.load_month(..., freeze=False); CA typed, US with
assign_types=False), then the gauge union + classify_gauges(..., month=) + Model B join are re-run the way build_gauge_report does
(union_gauge_frames, the US candidate pre-filter, the frozen gauge decisions as prior, flag_type_conflicts, then the builder's
public enrich_union / modelb_universe helpers). Workbooks are read with openpyxl (never saved); the table registry
(runs/<m>/table_registry_<market>_<m>.json), the manifests, the frozen decision CSVs and the memo are read. Nothing is written
except the optional --json result file.

Dataset source: --from-normalized (dev/fixture) > raw CSVs (when the CA code-reader raw dir exists) > none (every dataset-backed
check FAILs with "no dataset source"). Real-data-only checks/sub-checks SKIP when the dataset is not raw-derived: V17 (the type
review queue is a loader artifact), the raw-input part of V19, the Bully Dog part of V12.

--skip accepts only the IDs in SKIPPABLE (each has a documented reason); any other ID is an argument error.

Memo (V20): tags inside inline code spans (`...`) are literal mentions of the tag syntax, not citations, and are ignored.
"""
from __future__ import annotations

import argparse
import ast
import json
import logging
import math
import posixpath
import re
import sys
import traceback
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

import pandas as pd
from openpyxl import load_workbook

from ca_market_reports import ca_common as C

SHORT_NAMES: dict[str, str] = {
    "V01": "summary revenue total", "V02": "summary units total", "V03": "summary listings", "V04": "top 50 rank 1",
    "V05": "shares sum to 1", "V06": "amazon links", "V07": "currency labels", "V08": "innova tabs", "V09": "total rows",
    "V10": "lock files", "V11": "error tokens", "V12": "gauge scope totals", "V13": "excluded tab", "V14": "duplicate asins",
    "V15": "metadata keys", "V16": "charts", "V17": "type review queue", "V18": "per-asin reconciliation",
    "V19": "manifest hashes", "V20": "memo tags", "V21": "dataset sanity", "V22": "classifier golden", "V23": "benchmark joins",
}
assert tuple(SHORT_NAMES) == tuple(C.VALIDATION_CHECKS)
SKIPPABLE: dict[str, str] = {
    "V20": "memo not filled yet: the memo is written from the validated workbooks (workbook-first build only)",
    "V23": "CA gauge workbook built without --benchmark-market US (no US Benchmark / Same-ASIN sheets)",
}
# Golden map-only rows carry no ASIN in tests/fixtures/gauge_golden_titles.csv; the seeded map ASINs, in file order, are
# replicated from tests/test_gauge.py MAP_ONLY_ASINS (FLAGGED: the fixture should carry the ASIN).
GOLDEN_MAP_ONLY_ASINS: tuple[str, ...] = ("B0957S3F3H", "B07FFF4457", "B09ZDNM987", "B088FZRTNZ", "B09XDXTV77")
ERROR_TOKENS: tuple[str, ...] = ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A")
CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
AXIS_TAGS = ("catAx", "valAx", "dateAx", "serAx")
CHART_KIND = {"bar": "barChart", "pie": "pieChart", "scatter": "scatterChart"}
MONEY_TOL = 0.01
EXACT_TOL = 1e-6
SHARE_TOL = 1e-6
MAX_LISTED = 8

LOG = logging.getLogger("ca_market_reports")


# --------------------------------------------------------------------------------------
# Results
# --------------------------------------------------------------------------------------
@dataclass
class Result:
    id: str
    status: str
    evidence: str

    @property
    def line(self) -> str:
        return f"{self.status} {self.id} {SHORT_NAMES[self.id]}: {self.evidence}"


class Skip(Exception):
    """Raised by a check to report SKIP with a reason."""


def _listed(items: list[str], n: int = MAX_LISTED) -> str:
    head = "; ".join(items[:n])
    return head + (f"; … (+{len(items) - n} more)" if len(items) > n else "")


def _outcome(problems: list[str], ok: str) -> tuple[str, str]:
    if problems:
        return "FAIL", f"{len(problems)} problem(s): {_listed(problems)}"
    return "PASS", ok


# --------------------------------------------------------------------------------------
# Lazy context values (an exception is cached and re-raised for every check that needs the value)
# --------------------------------------------------------------------------------------
class Lazy:
    def __init__(self, fn: Callable[[], Any]):
        self.fn = fn
        self.done = False
        self.value: Any = None
        self.exc: BaseException | None = None

    def get(self) -> Any:
        if not self.done:
            try:
                self.value = self.fn()
            except Exception as exc:  # cached: every dependent check FAILs with the same text
                self.exc = exc
            self.done = True
        if self.exc is not None:
            raise self.exc
        return self.value


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v)) or (isinstance(v, str) and v.strip() == "")


def _num(v: Any) -> float | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _close(a: float, b: float, tol: float) -> bool:
    if math.isnan(a) or math.isnan(b):
        return math.isnan(a) and math.isnan(b)
    return abs(a - b) <= tol + 1e-9 * max(abs(a), abs(b))


# --------------------------------------------------------------------------------------
# Workbooks + registry
# --------------------------------------------------------------------------------------
@dataclass
class WB:
    name: str
    path: Path
    market: str          # CA | US
    kind: str            # cr | gauge
    book: Any = None     # openpyxl workbook (formulas kept) or None when the file is missing
    missing: bool = False


class Table:
    """One registry entry bound to its worksheet."""

    def __init__(self, entry: dict, wb: WB):
        self.e = entry
        self.wb = wb
        self.sheet = entry["sheet"]
        if self.sheet not in wb.book.sheetnames:
            raise KeyError(f"{wb.name}: registered sheet {self.sheet!r} missing")
        self.ws = wb.book[self.sheet]
        self.role = entry["role"]
        self.columns: list[str] = list(entry["columns"])
        self.first_col = int(entry["first_col"])
        self.header_row = entry["header_row"]
        self.first = int(entry["first_data_row"])
        self.last = int(entry["last_data_row"])
        self.total_row = entry["total_row"]
        self.residual_row = entry["residual_row"]
        self.filter = entry.get("dataset_filter", "")
        if self.header_row is not None:
            got = [self.ws.cell(self.header_row, self.first_col + j).value for j in range(len(self.columns))]
            if got != self.columns:
                raise ValueError(f"{wb.name}!{self.sheet}/{self.role}: header row {self.header_row} {got} != registry {self.columns}")

    @property
    def label(self) -> str:
        return f"{self.wb.name}!{self.sheet}/{self.role}"

    @property
    def data_rows(self) -> range:
        return range(self.first, self.last + 1)

    @property
    def n_rows(self) -> int:
        return max(0, self.last - self.first + 1)

    def has(self, header: str) -> bool:
        return header in self.columns

    def col(self, header: str) -> int:
        if header not in self.columns:
            raise KeyError(f"{self.label}: no column {header!r}")
        return self.first_col + self.columns.index(header)

    def cell(self, row: int, header: str):
        return self.ws.cell(row, self.col(header))

    def values(self, header: str, rows: range | list[int] | None = None) -> list[Any]:
        c = self.col(header)
        return [self.ws.cell(r, c).value for r in (self.data_rows if rows is None else rows)]

    def header_like(self, *prefixes: str) -> str | None:
        for h in self.columns:
            if any(h.startswith(p) for p in prefixes):
                return h
        return None


# --------------------------------------------------------------------------------------
# Header -> metric resolution (V09 totals and additivity)
# --------------------------------------------------------------------------------------
_CCY_SUFFIX = re.compile(r" \((?:CAD|USD)\)$")
REV_BASES = {"Monthly Rev", "Est. Monthly Retail Rev", "Est. Monthly Rev", "Revenue/Mo"}
UNIT_BASES = {"Monthly Units", "Est. Monthly Units Sold", "Est. Monthly Units", "Quantity/Mo"}
LISTING_BASES = {"# of Listings", "# ASINs"}
REVIEW_BASES = {"Total Reviews", "# of Reviews", "# Reviews", "Reviews"}
RATING_BASES = {"Avg Rating", "Avg. Rating", "Tool Rating", "Rating", "Avg rating", "Units-weighted rating"}
PRICE_BASES = {"Price Per Unit", "Avg price", "Avg Price"}
REV_SHARE_BASES = {"Monthly Rev Market Share %", "Rev share", "Revenue share", "Revenue by %"}
SHARE_HEADERS = REV_SHARE_BASES | {"Qty by %", "% of Listings"}
# Trend Proxy (Helium 10 fields as reported): Last Year Sales = plain sum over listings carrying it; YoY only counted by sign
TREND_HEADERS: tuple[str, ...] = ("Last Year Sales (Helium 10 field; semantics unverified)", "# Listings with Helium 10 YoY data",
                                  "# Listings with YoY > 0", "# Listings with YoY < 0", "Revenue share of listings with YoY > 0",
                                  "Revenue share of listings with YoY data")
# Gauge share of the code-reader market (Summary kpi table, one per gauge workbook) and its CA vs US view (US Benchmark)
SHARE_TABLE_TITLE = "Gauge share of the code-reader market"
BENCH_SHARE_TABLE_TITLE = "Gauge share of the code-reader market — CA vs US"
SHARE_LABELS: tuple[str, ...] = ("Code-reader export total", "(a) Gauge devices inside the code-reader export",
                                 "(b) All core gauge devices (CR ∪ gauge export)", "(b) denominator: CR total + gauge-only rows")
BENCH_SHARE_LABELS: dict[str, tuple[str, str]] = {
    "(a) Gauge devices inside the code-reader export: share of revenue": ("a", "s_rev"),
    "(a) Gauge devices inside the code-reader export: share of units": ("a", "s_u"),
    "(b) All core gauge devices (CR ∪ gauge export): share of revenue": ("b", "s_rev"),
    "(b) All core gauge devices (CR ∪ gauge export): share of units": ("b", "s_u")}
FUEL_TABLE_TITLE = "Fuel split (core devices)"
KPI_BLOCK_TITLE = "Key figures"


def _base(h: str) -> str:
    return _CCY_SUFFIX.sub("", h)


def weighted_rating(df: pd.DataFrame) -> float:
    r = df["rating"].astype(float)
    u = df["units_month"].astype(float).fillna(0.0)
    m = r.notna() & (r > 0)
    w = float(u[m].sum())
    return float((r[m] * u[m]).sum() / w) if w > 0 else float("nan")


def _div(a: float, b: float) -> float:
    return float(a) / float(b) if b and not math.isnan(b) else float("nan")


def metric(header: str, sub: pd.DataFrame, den: pd.DataFrame) -> tuple[float, float] | None:
    """(expected value, tolerance) of a standard aggregate column over `sub` (shares over `den`), None when unknown."""
    b = _base(header)
    rev, units = float(sub["revenue_month"].sum()), float(sub["units_month"].sum())
    if b in REV_BASES:
        return rev, MONEY_TOL
    if b in UNIT_BASES:
        return units, EXACT_TOL
    if b in LISTING_BASES:
        return float(sub["asin"].nunique()), EXACT_TOL
    if b in REVIEW_BASES:
        return float(sub["review_count"].astype(float).sum(skipna=True)), EXACT_TOL
    if b in RATING_BASES:
        return weighted_rating(sub), SHARE_TOL
    if b in PRICE_BASES:
        return _div(rev, units), MONEY_TOL
    if b in REV_SHARE_BASES:
        return _div(rev, float(den["revenue_month"].sum())), SHARE_TOL
    if b == "Qty by %":
        return _div(units, float(den["units_month"].sum())), SHARE_TOL
    if b == "% of Listings":
        return _div(float(sub["asin"].nunique()), float(den["asin"].nunique())), SHARE_TOL
    if b == "# with sales > 0":
        return float((sub["units_month"] > 0).sum()), EXACT_TOL
    if b == "# app-gauge-capable ASINs":
        return float((sub["_app_capable"] == "Y").sum()), EXACT_TOL
    if b in TREND_HEADERS:
        ly = sub["last_year_units"].astype(float).dropna()
        yoy = sub["yoy_units_pct"].astype(float)
        return {TREND_HEADERS[0]: (float(ly.sum()) if len(ly) else float("nan"), EXACT_TOL),
                TREND_HEADERS[1]: (float(yoy.notna().sum()), EXACT_TOL),
                TREND_HEADERS[2]: (float((yoy > 0).sum()), EXACT_TOL),
                TREND_HEADERS[3]: (float((yoy < 0).sum()), EXACT_TOL),
                TREND_HEADERS[4]: (_div(float(sub.loc[yoy > 0, "revenue_month"].sum()), rev), SHARE_TOL),
                TREND_HEADERS[5]: (_div(float(sub.loc[yoy.notna(), "revenue_month"].sum()), rev), SHARE_TOL)}[b]
    return None


# --------------------------------------------------------------------------------------
# Dataset re-derivation
# --------------------------------------------------------------------------------------
@dataclass
class Data:
    source: str                         # "raw" | "normalized"
    cr: pd.DataFrame                    # CA code-reader rows (prepared like the CR builder: typed, tiered, _url)
    ca_u: pd.DataFrame                  # CA gauge union (classified + enriched)
    modelb: pd.DataFrame                # Model B universe (CR Type == Dongle + app map)
    us_u: pd.DataFrame | None           # US gauge union (classified + enriched)
    frames: dict[str, pd.DataFrame] = field(default_factory=dict)   # V21: every input frame
    notes: list[str] = field(default_factory=list)
    cr_full: dict[str, pd.DataFrame] = field(default_factory=dict)  # FULL code-reader export per market (US: before the candidate filter)
    fuel_absent: dict[str, bool] = field(default_factory=dict)      # dev frames without fuel_scope: core devices count as 'unspecified'


def _overlay_frozen_types(cr: pd.DataFrame, runs_dir: Path, month: str, type_map: Path, notes: list[str]) -> pd.DataFrame:
    """Replay the frozen type decisions of this month onto the fresh (freeze=False) derivation, as ca_load does on a replay run
    (minimal replication of ca_load._replay_type_decisions without its write): the override map wins over a frozen row."""
    from ca_market_reports import ca_tiers, ca_types
    path = C.run_file(runs_dir, month, "type_decisions", "CA_code_reader")
    if not path.exists():
        notes.append(f"no frozen type decisions ({path.name}): fresh types used")
        return cr
    fz = ca_types.read_type_decisions(path)
    overrides = ca_types.load_overrides(type_map, month)
    by = {r["asin"]: r for _, r in fz.iterrows()}
    out = cr.copy()
    replayed = 0
    for i, a in zip(out.index, out["asin"]):
        r = by.get(a)
        if r is None:
            continue
        ov = overrides.get(a)
        if ov is not None and (r["type"] != ov or r["type_source"] != "override"):
            continue
        for c in ("type", "type_source", "type_rule_id"):
            out.at[i, c] = r[c]
        out.at[i, "type_confidence"] = float(r["type_confidence"])
        out.at[i, "type_conflict"] = bool(r["type_conflict"])
        replayed += 1
    notes.append(f"frozen type decisions replayed for {replayed}/{len(out)} CR rows")
    return ca_tiers.assign_cr_tiers(out)


def _classified_union(ds: C.CaDataset, gauge_map_path: Path, runs_dir: Path) -> pd.DataFrame:
    """build_gauge_report._assemble_union_from_loader without its freeze write (minimal replication; FLAGGED)."""
    from ca_market_reports.build_gauge_report import read_gauge_map
    from ca_market_reports.ca_gauge_classification import candidate_mask, classify_gauges
    from ca_market_reports.ca_load import union_gauge_frames
    from ca_market_reports.ca_types import flag_type_conflicts
    if ds.gauge_set is None:
        raise ValueError(f"{ds.market}: no gauge export (gauge raw dir absent)")
    gm = read_gauge_map(gauge_map_path)
    cr = ds.code_reader
    if ds.market == "US":
        cr = cr[candidate_mask(cr, gm)].copy()
    union = union_gauge_frames(cr, ds.gauge_set, [])
    prior_path = C.run_file(runs_dir, ds.month, "gauge_decisions", f"{ds.market}_gauge")
    prior = pd.read_csv(prior_path, dtype=str, keep_default_na=False) if prior_path.exists() else None
    union = classify_gauges(union, gm, prior, month=ds.month)
    return flag_type_conflicts(union)


def derive_data(a: argparse.Namespace) -> Data:
    from ca_market_reports import build_ca_code_reader_report as CRB
    from ca_market_reports import build_gauge_report as GB
    from ca_market_reports import ca_xlsx_style as X
    notes: list[str] = []
    gauge_map = C.MAPS_DIR / "ca_gauge_map.csv"
    app_map = C.MAPS_DIR / "ca_app_gauge_brands.csv"
    if a.from_normalized:
        ca = X.dataset_from_normalized(X.read_normalized_csv(a.from_normalized), "CA", a.month)
        us = X.dataset_from_normalized(X.read_normalized_csv(a.us_from_normalized), "US", a.month) if a.us_from_normalized else None
        ca_u, _, fa_ca = GB.gauge_union(ca, preclassified=True, gauge_map_path=gauge_map, runs_dir=a.runs_dir, rederive=False)
        fuel_absent = {"CA": fa_ca}
        us_u = None
        if us is not None:
            us_u, _, fuel_absent["US"] = GB.gauge_union(us, preclassified=True, gauge_map_path=gauge_map, runs_dir=a.runs_dir,
                                                        rederive=False)
        cr = CRB.prepare_cr_frame(ca.code_reader)
        frames = {"CA normalized": X.read_normalized_csv(a.from_normalized)}
        if a.us_from_normalized:
            frames["US normalized"] = X.read_normalized_csv(a.us_from_normalized)
        source = "normalized"
        notes.append(f"dataset from --from-normalized {a.from_normalized}")
    else:
        if not Path(a.raw_dir).is_dir():
            raise FileNotFoundError(f"no dataset source: raw dir absent ({a.raw_dir}) and no --from-normalized")
        from ca_market_reports.ca_load import load_month
        prev = LOG.level
        LOG.setLevel(logging.ERROR)
        logging.getLogger("ca_market_reports.gauge").setLevel(logging.ERROR)
        try:
            ca = load_month("CA", a.month, cr_raw_dir=a.raw_dir, gauge_raw_dir=a.gauge_raw_dir, runs_dir=a.runs_dir,
                            assign_types=True, freeze=False)
            ca.code_reader = _overlay_frozen_types(ca.code_reader, Path(a.runs_dir), a.month, C.MAPS_DIR / "ca_type_overrides.csv",
                                                   notes)
            us = None
            if Path(a.us_gauge_raw_dir).is_dir() and Path(a.us_cr_raw_dir).is_dir():
                us = load_month("US", a.month, cr_raw_dir=a.us_cr_raw_dir, gauge_raw_dir=a.us_gauge_raw_dir, runs_dir=a.runs_dir,
                                assign_types=False, freeze=False)
            else:
                notes.append(f"US raw dirs absent ({a.us_cr_raw_dir}, {a.us_gauge_raw_dir}): no US dataset")
            ca_union = _classified_union(ca, gauge_map, Path(a.runs_dir))
            ca_u = GB.enrich_union(ca_union, C.MARKETS["CA"], [])
            us_u = GB.enrich_union(_classified_union(us, gauge_map, Path(a.runs_dir)), C.MARKETS["US"], []) if us else None
            fuel_absent = {"CA": False, "US": False}   # classify_gauges always sets fuel_scope
        finally:
            LOG.setLevel(prev)
            logging.getLogger("ca_market_reports.gauge").setLevel(logging.NOTSET)
        cr = CRB.prepare_cr_frame(ca.code_reader)
        frames = {"CA code_reader": ca.code_reader, "CA gauge": ca.gauge_set}
        if us is not None:
            frames.update({"US code_reader": us.code_reader, "US gauge": us.gauge_set})
        source = "raw"
    modelb = GB.modelb_universe(ca, app_map)
    if ca_u["asin"].duplicated().any() or (us_u is not None and us_u["asin"].duplicated().any()):
        raise ValueError("duplicate ASINs in a re-derived union")
    frames = {k: v for k, v in frames.items() if v is not None}
    cr_full = {"CA": ca.code_reader}
    if us is not None:
        cr_full["US"] = us.code_reader      # the full US code-reader export (the candidate filter only narrows a copy)
    return Data(source=source, cr=cr, ca_u=ca_u, modelb=modelb, us_u=us_u, frames=frames, notes=notes, cr_full=cr_full,
                fuel_absent=fuel_absent)


# --------------------------------------------------------------------------------------
# Validator
# --------------------------------------------------------------------------------------
class Validator:
    def __init__(self, a: argparse.Namespace):
        self.a = a
        self.m = a.month
        self.out_dirs = {"cr": Path(a.cr_out_dir), "gauge_CA": Path(a.gauge_out_dir), "gauge_US": Path(a.us_gauge_out_dir)}
        self.data = Lazy(lambda: derive_data(a))
        self.books = Lazy(self._load_books)
        self.registry = Lazy(self._load_registry)

    # ---------------------------------------------------------------- inputs
    def _load_books(self) -> dict[str, WB]:
        m = self.m
        specs = [(C.cr_report_name("CA", m), "cr", "CA", "cr"), (C.cr_analysis_name("CA", m), "cr", "CA", "cr"),
                 (C.gauge_report_name("CA", m), "gauge_CA", "CA", "gauge"), (C.gauge_report_name("US", m), "gauge_US", "US", "gauge")]
        out = {}
        for name, d, mk, kind in specs:
            p = self.out_dirs[d] / name
            wb = WB(name, p, mk, kind)
            if p.exists():
                wb.book = load_workbook(p, data_only=False)
            else:
                wb.missing = True
            out[name] = wb
        return out

    def _load_registry(self) -> dict[str, dict]:
        reg = {}
        for mk in ("CA", "US"):
            p = C.run_file(Path(self.a.runs_dir), self.m, "table_registry", mk, "json")
            if not p.exists():
                raise FileNotFoundError(f"table registry missing: {p}")
            r = json.loads(p.read_text(encoding="utf-8"))
            if (r.get("market"), r.get("month")) != (mk, self.m):
                raise ValueError(f"{p.name}: registry for {r.get('market')}/{r.get('month')}")
            reg[mk] = r
        return reg

    def wb(self, name: str) -> WB:
        w = self.books.get()[name]
        if w.missing:
            raise FileNotFoundError(f"workbook missing: {w.path}")
        return w

    def present_books(self) -> tuple[list[WB], list[str]]:
        ok, missing = [], []
        for w in self.books.get().values():
            (missing if w.missing else ok).append(w)
        return ok, [f"workbook missing: {w.path}" for w in missing]

    def tables(self, w: WB, role: str | None = None, sheet: str | None = None) -> list[Table]:
        reg = self.registry.get()[w.market]
        if w.name not in reg["workbooks"]:
            raise KeyError(f"{w.name} not in the {w.market} table registry")
        out = []
        for e in reg["tables"]:
            if e["workbook"] != w.name or (role and e["role"] != role) or (sheet and e["sheet"] != sheet):
                continue
            out.append(Table(e, w))
        return out

    def one(self, w: WB, role: str, sheet: str | None = None) -> Table:
        ts = self.tables(w, role, sheet)
        if len(ts) != 1:
            raise ValueError(f"{w.name}: expected one {role!r} table{' on ' + sheet if sheet else ''}, found {len(ts)}")
        return ts[0]

    def same_asin_set(self) -> set[str]:
        """CA ∩ US union ASINs that are device scope in EITHER market."""
        d = self.data.get()
        if d.us_u is None:
            raise ValueError("no US dataset for the Same-ASIN join")
        ca, us = d.ca_u.set_index("asin"), d.us_u.set_index("asin")
        both = set(ca.index) & set(us.index)
        return {a for a in both if bool(ca.at[a, "gauge_device_scope"]) or bool(us.at[a, "gauge_device_scope"])}

    def titled(self, w: WB, role: str, sheet: str, title: str) -> Table:
        ts = [t for t in self.tables(w, role, sheet) if t.e.get("title") == title]
        if len(ts) != 1:
            raise ValueError(f"{w.name}: expected one {role!r} table titled {title!r} on {sheet}, found {len(ts)}")
        return ts[0]

    @property
    def real(self) -> bool:
        """The dataset is raw-derived (same source rule as derive_data, decided without deriving)."""
        return not self.a.from_normalized and Path(self.a.raw_dir).is_dir()

    def real_reason(self) -> str:
        if not Path(self.a.raw_dir).is_dir():
            src = f"--from-normalized {self.a.from_normalized}" if self.a.from_normalized else "no dataset source"
            return f"raw dir absent ({self.a.raw_dir}); {src}"
        return f"dataset from --from-normalized {self.a.from_normalized}, not the raw exports"

    def union(self, w: WB) -> pd.DataFrame:
        d = self.data.get()
        u = d.ca_u if w.market == "CA" else d.us_u
        if u is None:
            raise ValueError(f"no {w.market} dataset (US raw dirs absent and no --us-from-normalized)")
        return u

    def frame_for(self, t: Table) -> tuple[pd.DataFrame, pd.DataFrame] | None:
        """(rows the table covers, share denominator) from the registry dataset_filter; None for tables without dataset rows."""
        f = t.filter
        d = self.data.get()
        if t.wb.kind == "cr":
            base = d.cr
            if f == "all CA code-reader rows":
                sub = base
            elif (mm := re.fullmatch(r"brand_key == ((['\"]).*\2)", f)):
                sub = base[base["brand_key"] == ast.literal_eval(mm.group(1))]
            elif (mm := re.fullmatch(r"price_tier in (\[.*\])", f)):
                sub = base[base["price_tier"].isin(ast.literal_eval(mm.group(1)))]
            elif f == "":
                return None
            else:
                raise ValueError(f"{t.label}: unknown dataset_filter {f!r}")
            den = base if t.role in ("tier_pivot", "category") else sub
            if t.role == "kpi":
                den = base
            return sub, den
        u = self.union(t.wb)
        core = u[u["_core"]]
        if f.startswith("core devices, both markets") or f.startswith("asin in CA union") or f in (
                "", "GAUGE_CLASSES", "apps", "see metric labels") or f.startswith(("dedupe_audit rows", "classification decisions")):
            return None
        if f.startswith("core devices"):
            sub, den_base = core, core
        elif f == "gauge_device_scope":
            sub = den_base = u[u["gauge_device_scope"]]
        elif f == "gauge_class in GAUGE_ADJACENT_CLASSES":
            sub = den_base = u[u["gauge_class"].isin(C.GAUGE_ADJACENT_CLASSES)]
        elif f.startswith("Model B universe"):
            sub = den_base = d.modelb
        elif f == "code-reader ∪ gauge union (all classes)":
            sub = den_base = u
        elif f == "gauge_class in GAUGE_EXCLUDED_CLASSES + ambiguous":
            sub = den_base = u[u["gauge_class"].isin(C.GAUGE_EXCLUDED_CLASSES + ("ambiguous",))]
        elif f == "brand_key == 'innova' & gauge_device_scope":
            return None
        else:
            raise ValueError(f"{t.label}: unknown dataset_filter {f!r}")
        mm = re.search(r"& brand_key == ((['\"]).*?\2)$", f)
        if mm:
            sub = sub[sub["brand_key"] == ast.literal_eval(mm.group(1))]
        den = den_base if t.role == "kpi" else sub
        return sub, den

    # ---------------------------------------------------------------- checks
    def _summary_total(self, w: WB, header_base: str) -> tuple[Table, float | None]:
        t = self.one(w, "summary_brands", "Summary")
        if t.total_row is None:
            raise ValueError(f"{t.label}: no Total row")
        h = next((c for c in t.columns if _base(c) == header_base), None)
        if h is None:
            raise KeyError(f"{t.label}: no {header_base!r} column")
        return t, _num(t.cell(t.total_row, h).value)

    def _v_summary(self, header_base: str, tol: float, what: str, fn: Callable[[pd.DataFrame], float]) -> tuple[str, str]:
        d = self.data.get()
        probs, ev = [], []
        targets = [(C.cr_report_name("CA", self.m), d.cr, "all CR rows")]
        for mk in ("CA", "US"):
            name = C.gauge_report_name(mk, self.m)
            w = self.books.get()[name]
            if not w.missing:
                u = self.union(w)
                targets.append((name, u[u["_core"]], "core devices"))
            else:
                probs.append(f"workbook missing: {w.path}")
        for name, frame, scope in targets:
            w = self.wb(name)
            t, got = self._summary_total(w, header_base)
            exp = fn(frame)
            if got is None or not _close(got, exp, tol):
                probs.append(f"{name} Summary Total {what} {got} != re-derived {exp:,.2f} ({scope})")
            else:
                ev.append(f"{name.split('_202')[0]} {got:,.2f}")
        return _outcome(probs, f"{what} == re-derived dataset ({'; '.join(ev)})")

    def v01(self):
        return self._v_summary("Monthly Rev", MONEY_TOL, "revenue", lambda f: float(f["revenue_month"].sum()))

    def v02(self):
        return self._v_summary("Monthly Units", EXACT_TOL, "units", lambda f: float(f["units_month"].sum()))

    def v03(self):
        return self._v_summary("# of Listings", EXACT_TOL, "# of listings", lambda f: float(f["asin"].nunique()))

    def v04(self):
        probs, ev = [], []
        names = [C.cr_report_name("CA", self.m), C.gauge_report_name("CA", self.m), C.gauge_report_name("US", self.m)]
        for name in names:
            w = self.wb(name)
            t = self.one(w, "top_by_revenue", "Top 50")
            sub, _ = self.frame_for(t)
            rev_h = next(h for h in t.columns if _base(h) in REV_BASES)
            if len(sub) == 0:
                if t.n_rows:
                    probs.append(f"{name}: Top 50 has rows but the dataset scope is empty")
                ev.append(f"{name.split('_202')[0]} empty scope")
                continue
            exp = float(sub["revenue_month"].max())
            got = _num(t.cell(t.first, rev_h).value)
            rank = t.cell(t.first, "Ranking").value
            if rank != 1 or got is None or not _close(got, exp, MONEY_TOL):
                probs.append(f"{name}: rank {rank} revenue {got} != dataset max {exp:,.2f}")
            else:
                ev.append(f"{name.split('_202')[0]} {got:,.2f}")
        return _outcome(probs, "rank 1 == dataset max revenue (" + "; ".join(ev) + ")")

    def v05(self):
        books, probs = self.present_books()
        n_tables = n_cols = 0
        for w in books:
            for t in self.tables(w):
                if t.header_row is None:
                    continue
                share_cols = [h for h in t.columns if h in SHARE_HEADERS]
                if not share_cols:
                    continue
                rows = list(t.data_rows) + ([t.residual_row] if t.residual_row else [])
                counted = False
                for h in share_cols:
                    vals = [_num(v) for v in t.values(h, rows)]
                    vals = [v for v in vals if v is not None]
                    if not vals:
                        continue
                    s = sum(vals)
                    counted = True
                    n_cols += 1
                    if abs(s - 1.0) > SHARE_TOL:
                        probs.append(f"{t.label} '{h}' sums to {s:.9f}")
                n_tables += counted
        return _outcome(probs, f"{n_cols} share columns in {n_tables} tables sum to 1 ± {SHARE_TOL:g} (displayed + residual)")

    def v06(self):
        books, probs = self.present_books()
        n_links = n_urls = 0
        rx = C.ALLOWED_FORMULA_RES
        for w in books:
            allowed_cells: set[tuple[str, str]] = set()
            for t in self.tables(w):
                link_cols = [h for h in t.columns if h == "Link" or h.startswith("Link ")]
                if not link_cols and not t.has("URL"):
                    continue
                if not t.has("ASIN"):
                    probs.append(f"{t.label}: link/URL column without an ASIN column")
                    continue
                for r in t.data_rows:
                    asin = t.cell(r, "ASIN").value
                    for h in link_cols:
                        tld = "com" if h == "Link US" else "ca" if h == "Link CA" else ("ca" if w.market == "CA" else "com")
                        c = t.cell(r, h)
                        allowed_cells.add((t.sheet, c.coordinate))
                        v = c.value
                        mm = next((x.match(v) for x in rx if isinstance(v, str) and x.match(v)), None)
                        if c.data_type != "f" or mm is None:
                            probs.append(f"{t.label} {c.coordinate}: not an allowed HYPERLINK: {str(v)[:60]!r}")
                        elif mm.group("asin") != asin or mm.group("tld") != tld:
                            probs.append(f"{t.label} {c.coordinate}: link {mm.group('tld')}/{mm.group('asin')} != row {tld}/{asin}")
                        n_links += 1
                    if t.has("URL"):
                        dom = C.MARKETS[w.market].domain
                        u = t.cell(r, "URL").value
                        if u != f"https://{dom}/dp/{asin}":
                            probs.append(f"{t.label} row {r}: URL {u!r} != https://{dom}/dp/{asin}")
                        n_urls += 1
                for extra in (t.total_row, t.residual_row):
                    for h in link_cols:
                        if extra and t.cell(extra, h).value is not None:
                            probs.append(f"{t.label} {t.cell(extra, h).coordinate}: link on a Total/residual row")
            for ws in w.book.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if c.data_type == "f" and (ws.title, c.coordinate) not in allowed_cells:
                            probs.append(f"{w.name}!{ws.title}!{c.coordinate}: formula outside a registered link column")
                        elif (w.market == "CA" and isinstance(c.value, str) and "amazon.com" in c.value
                              and ws.title not in C.CA_SHEETS_ALLOWING_US):
                            probs.append(f"{w.name}!{ws.title}!{c.coordinate}: amazon.com text on a CA-only sheet")
        return _outcome(probs, f"{n_links} HYPERLINKs and {n_urls} URL cells match ALLOWED_FORMULA_RES, row ASIN and market domain")

    @staticmethod
    def _col_market(h: str, own: str) -> str:
        if "(USD)" in h or h.startswith("US "):
            return "US"
        if "(CAD)" in h or h.startswith("CA "):
            return "CA"
        return own

    def v07(self):
        books, probs = self.present_books()
        n = 0
        for w in books:
            for t in self.tables(w):
                kpi_style = t.columns == ["Metric", "Value"]
                if t.header_row is None and not kpi_style:
                    continue
                rows = list(t.data_rows) + [r for r in (t.total_row, t.residual_row) if r]
                if kpi_style:
                    for r in t.data_rows:
                        label, vc = t.ws.cell(r, t.first_col).value, t.ws.cell(r, t.first_col + 1)
                        if _num(vc.value) is None or "$" not in (vc.number_format or ""):
                            continue
                        mk = self._col_market(str(label), w.market)
                        tok = "(CAD)" if mk == "CA" else "(USD)"
                        n += 1
                        if tok not in str(label):
                            probs.append(f"{t.label} {vc.coordinate}: money KPI label {label!r} lacks {tok}")
                        probs += self._fmt_problem(t, vc, mk)
                    continue
                for h in t.columns:
                    mk = self._col_market(h, w.market)
                    cells = [t.ws.cell(r, t.col(h)) for r in rows]
                    nums = [c for c in cells if _num(c.value) is not None]
                    money = [c for c in nums if "$" in (c.number_format or "")]
                    for c in money:
                        probs += self._fmt_problem(t, c, mk)
                    n += len(money)
                    tok = "(CAD)" if mk == "CA" else "(USD)"
                    rev_like = re.search(r"\b(Rev|Price|Revenue)\b", h) and nums and not any("%" in (c.number_format or "") for c in nums)
                    if (money or rev_like) and tok not in h:
                        probs.append(f"{t.label}: money column {h!r} lacks {tok}")
                    if w.market == "CA" and mk == "US" and (money or "(USD)" in h) and t.sheet not in C.CA_SHEETS_ALLOWING_US:
                        probs.append(f"{t.label}: USD column {h!r} on a sheet outside CA_SHEETS_ALLOWING_US")
                    if "(USD)" in h and w.market == "CA" and t.sheet not in C.CA_SHEETS_ALLOWING_US:
                        probs.append(f"{t.label}: header {h!r} outside CA_SHEETS_ALLOWING_US")
            for ws in w.book.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        fmt = c.number_format or ""
                        if "$" not in fmt or _num(c.value) is None:
                            continue
                        us_fmt = "CA$" not in fmt
                        if w.market == "US" and not us_fmt:
                            probs.append(f"{w.name}!{ws.title}!{c.coordinate}: CA$ format in the US workbook")
                        if w.market == "CA" and us_fmt and ws.title not in C.CA_SHEETS_ALLOWING_US:
                            probs.append(f"{w.name}!{ws.title}!{c.coordinate}: USD format outside CA_SHEETS_ALLOWING_US")
        return _outcome(probs, f"{n} money cells carry their table currency (CA$ / $), headers carry (CAD)/(USD)")

    @staticmethod
    def _fmt_problem(t: Table, c, mk: str) -> list[str]:
        fmt = c.number_format or ""
        if mk == "CA" and "CA$" not in fmt:
            return [f"{t.label} {c.coordinate}: CA money format {fmt!r} lacks CA$"]
        if mk == "US" and ("CA" in fmt or "$" not in fmt):
            return [f"{t.label} {c.coordinate}: US money format {fmt!r}"]
        return []

    def v08(self):
        d = self.data.get()
        probs, ev = [], []
        w = self.wb(C.cr_report_name("CA", self.m))
        t = self.one(w, "innova", "Innova")
        exp = set(d.cr.loc[d.cr["brand_key"] == "innova", "asin"])
        got = [v for v in t.values("ASIN")]
        if len(got) != len(exp) or set(got) != exp:
            probs.append(f"CR Innova tab rows {len(got)} != dataset innova rows {len(exp)} (diff {sorted(set(got) ^ exp)[:5]})")
        ev.append(f"CR Innova tab {len(got)} rows == dataset {len(exp)}")
        for mk in ("CA", "US"):
            gw = self.wb(C.gauge_report_name(mk, self.m))
            u = self.union(gw)
            line = self.one(gw, "innova", "Innova")
            text = str(line.ws.cell(line.first, line.first_col).value)
            mm = re.fullmatch(r"Innova gauge/HUD device listings in this dataset: (\d+)", text)
            n_dev = int(((u["brand_key"] == "innova") & u["gauge_device_scope"]).sum())
            if mm is None or int(mm.group(1)) != n_dev:
                probs.append(f"{mk} gauge Innova line {text!r} != re-derived {n_dev}")
            ev.append(f"{mk} gauge device line {n_dev}")
            if mk == "CA":
                hw = self.one(gw, "modelb_top", "Innova")
                exp_hw = set(d.cr.loc[(d.cr["brand_key"] == "innova") & (d.cr["type"] == "Dongle"), "asin"])
                got_hw = hw.values("ASIN")
                if len(got_hw) != len(exp_hw) or set(got_hw) != exp_hw:
                    probs.append(f"CA gauge Innova hardware rows {len(got_hw)} != Model B innova rows {len(exp_hw)}")
                ev.append(f"CA gauge hardware list {len(got_hw)} == Model B innova {len(exp_hw)}")
        return _outcome(probs, "; ".join(ev))

    def _bench_frames(self, t: Table) -> tuple[pd.DataFrame, pd.DataFrame]:
        d = self.data.get()
        if d.us_u is None:
            raise ValueError("benchmark table but no US dataset")
        return d.ca_u[d.ca_u["_core"]], d.us_u[d.us_u["_core"]]

    def _expected_cells(self, t: Table) -> list[tuple[str, float, float]] | None:
        """[(header, expected, tol)] for the Total row of a table (None: the table covers no dataset rows)."""
        if t.role in ("benchmark_brands", "benchmark_subtypes", "benchmark_tiers"):
            ca, us = self._bench_frames(t)
            out = []
            for h in t.columns:
                side = ca if h.startswith("CA ") else us if h.startswith("US ") else None
                if side is not None:
                    mtr = metric(h[3:], side, side)
                    if mtr is None:
                        raise KeyError(f"{t.label}: unknown column {h!r}")
                    out.append((h, *mtr))
            out.append(("Units ratio US/CA", _div(float(us["units_month"].sum()), float(ca["units_month"].sum())), SHARE_TOL))
            return out
        fr = self.frame_for(t)
        if fr is None:
            return None
        sub, den = fr
        if t.role == "tier_matrix":
            what = "revenue_month" if "Monthly Rev" in t.e.get("title", "") else "units_month"
            tol = MONEY_TOL if what == "revenue_month" else EXACT_TOL
            out = []
            for h in t.columns[1:]:
                b = _base(h)
                if b == "Total":
                    out.append((h, float(sub[what].sum()), tol))
                else:
                    cls = next((k for k, v in C.GAUGE_SUBTYPE_LABELS.items() if v == b), None)
                    if cls is None:
                        raise KeyError(f"{t.label}: unknown column {h!r}")
                    out.append((h, float(sub.loc[sub["gauge_class"] == cls, what].sum()), tol))
            return out
        out = []
        for h in t.columns:
            mtr = metric(h, sub, den)
            if mtr is not None:
                out.append((h, *mtr))
        return out

    def v09(self):
        books, probs = self.present_books()
        n_tot = n_cells = n_special = 0
        for w in books:
            for t in self.tables(w):
                special = self._special_rules(t)
                if special is not None:
                    sp, n = special
                    probs += sp
                    n_special += 1
                    n_cells += n
                    continue
                if t.role == "kpi":
                    if t.filter == "see metric labels":
                        continue          # the gauge Summary "Key figures" block: every row is re-derived by V12
                    probs += self._brand_kpi(t)
                    continue
                if t.total_row is None:
                    continue
                exp = self._expected_cells(t)
                if exp is None:
                    raise ValueError(f"{t.label}: Total row but dataset_filter {t.filter!r} covers no rows")
                n_tot += 1
                known = {h for h, _, _ in exp}
                for h in t.columns:
                    v = t.cell(t.total_row, h).value
                    if h not in known and _num(v) is not None:
                        probs.append(f"{t.label}: Total cell {h!r}={v} has no re-derivation rule")
                for h, e, tol in exp:
                    got = _num(t.cell(t.total_row, h).value)
                    n_cells += 1
                    if got is None:
                        if not math.isnan(e):
                            probs.append(f"{t.label} Total {h!r} blank != {e:,.4f}")
                    elif not _close(got, e, tol):
                        probs.append(f"{t.label} Total {h!r} {got:,.4f} != re-derived {e:,.4f}")
                # additivity: displayed rows + residual == Total for revenue/units
                rows = list(t.data_rows) + ([t.residual_row] if t.residual_row else [])
                for h in t.columns:
                    b = _base(h[3:] if t.role.startswith("benchmark") and h[:3] in ("CA ", "US ") else h)
                    if b not in REV_BASES | UNIT_BASES and not (t.role == "tier_matrix" and h != t.columns[0]):
                        continue
                    tv = _num(t.cell(t.total_row, h).value)
                    if tv is None:
                        continue
                    s = sum(x for x in (_num(v) for v in t.values(h, rows)) if x is not None)
                    if not _close(s, tv, MONEY_TOL * max(1, len(rows))):
                        probs.append(f"{t.label} '{h}': rows + residual {s:,.2f} != Total {tv:,.2f}")
                # residual present where truncated
                full = self._entities(t)
                if full is not None and full > t.n_rows and t.residual_row is None:
                    probs.append(f"{t.label}: {t.n_rows} rows shown of {full} but no residual row")
        return _outcome(probs, f"{n_tot} Total rows + {n_special} share/fuel/count/app-matrix tables ({n_cells} cells) == "
                               f"full filtered dataset; rows + residual add up")

    # ---------------------------------------------------------------- V09 special tables
    def _special_rules(self, t: Table) -> tuple[list[str], int] | None:
        title = t.e.get("title", "")
        if t.role == "kpi" and title == SHARE_TABLE_TITLE:
            return self._check_share_table(t)
        if t.role == "kpi" and title == BENCH_SHARE_TABLE_TITLE:
            return self._check_bench_share(t)
        if t.role == "subtype_mix" and title == FUEL_TABLE_TITLE:
            return self._check_fuel(t)
        if t.role == "kpi" and t.wb.kind == "gauge" and t.sheet == "Innova":
            u = self.union(t.wb)
            exp = int(((u["brand_key"] == "innova") & u["gauge_device_scope"]).sum())
            return self._check_number(t, t.ws.cell(t.first, t.first_col), exp), 1
        if t.role == "kpi" and t.sheet == "US vs CA Same-ASIN":
            return self._check_number(t, t.ws.cell(t.first, t.first_col + 1), len(self.same_asin_set())), 1
        if t.role == "modelb_app_matrix":
            return self._check_app_matrix(t)
        return None

    @staticmethod
    def _check_number(t: Table, cell, exp: float) -> list[str]:
        got = _num(cell.value)
        if got is None or not _close(got, float(exp), EXACT_TOL):
            return [f"{t.label} {cell.coordinate} {cell.value!r} != re-derived {exp}"]
        return []

    def share_rows(self, mk: str) -> tuple[dict[str, dict[str, float]], list[str]]:
        """Gauge share of the code-reader market for one market, re-derived from the FULL code-reader export."""
        d = self.data.get()
        if mk not in d.cr_full:
            raise ValueError(f"no full {mk} code-reader export (US raw dirs absent and no --us-from-normalized)")
        cr = d.cr_full[mk]
        u = d.ca_u if mk == "CA" else d.us_u
        core = u[u["_core"]]
        a = core[core["source_set"].isin(["code_reader", "both"])]
        g = core[core["source_set"] == "gauge"]
        cr_asins = set(cr["asin"])
        probs = []
        if set(a["asin"]) - cr_asins:
            probs.append(f"{mk}: code_reader/both core rows missing from the code-reader export {sorted(set(a['asin']) - cr_asins)[:5]}")
        if set(g["asin"]) & cr_asins:
            probs.append(f"{mk}: gauge-only core rows found in the code-reader export {sorted(set(g['asin']) & cr_asins)[:5]}")
        cr_n = float(cr["asin"].nunique())
        cr_rev = float(pd.to_numeric(cr["revenue_month"]).astype(float).sum())
        cr_u = float(pd.to_numeric(cr["units_month"]).astype(float).sum())
        den_n, den_rev = cr_n + g["asin"].nunique(), cr_rev + float(g["revenue_month"].sum())
        den_u = cr_u + float(g["units_month"].sum())
        a_rev, a_u = float(a["revenue_month"].sum()), float(a["units_month"].sum())
        b_rev, b_u = float(core["revenue_month"].sum()), float(core["units_month"].sum())

        def row(n, rev, units, s_rev, s_u):
            return {"# ASINs": float(n), "Monthly Rev": rev, "Monthly Units": units, "Share of revenue": s_rev, "Share of units": s_u}

        L = SHARE_LABELS
        return {L[0]: row(cr_n, cr_rev, cr_u, _div(cr_rev, cr_rev), _div(cr_u, cr_u)),
                L[1]: row(a["asin"].nunique(), a_rev, a_u, _div(a_rev, cr_rev), _div(a_u, cr_u)),
                L[2]: row(core["asin"].nunique(), b_rev, b_u, _div(b_rev, den_rev), _div(b_u, den_u)),
                L[3]: row(den_n, den_rev, den_u, _div(den_rev, den_rev), _div(den_u, den_u))}, probs

    _SHARE_TOL = {"# ASINs": EXACT_TOL, "Monthly Rev": MONEY_TOL, "Monthly Units": EXACT_TOL, "Share of revenue": SHARE_TOL,
                  "Share of units": SHARE_TOL}

    def _cmp(self, t: Table, r: int, h: str, exp: float, tol: float, probs: list[str]) -> None:
        got = _num(t.cell(r, h).value)
        if got is None:
            if not math.isnan(exp):
                probs.append(f"{t.label} row {r} {h!r} blank != {exp:,.6f}")
        elif not _close(got, exp, tol):
            probs.append(f"{t.label} row {r} {h!r} {got:,.6f} != re-derived {exp:,.6f}")

    def _check_share_table(self, t: Table) -> tuple[list[str], int]:
        exp, probs = self.share_rows(t.wb.market)
        labels = [str(v) for v in t.values("Measure")]
        if sorted(labels) != sorted(SHARE_LABELS) or len(labels) != len(set(labels)):
            probs.append(f"{t.label}: rows {labels} != {list(SHARE_LABELS)}")
        n = 0
        for r, label in zip(t.data_rows, labels):
            if label not in exp:
                continue
            for h in t.columns[1:]:
                b = _base(h)
                if b not in exp[label]:
                    probs.append(f"{t.label}: column {h!r} has no re-derivation rule")
                    continue
                self._cmp(t, r, h, exp[label][b], self._SHARE_TOL[b], probs)
                n += 1
        return probs, n

    def _check_bench_share(self, t: Table) -> tuple[list[str], int]:
        ca, p1 = self.share_rows("CA")
        us, p2 = self.share_rows("US")
        probs = p1 + p2
        labels = [str(v) for v in t.values("Measure")]
        if sorted(labels) != sorted(BENCH_SHARE_LABELS):
            probs.append(f"{t.label}: rows {labels} != {list(BENCH_SHARE_LABELS)}")
        n = 0
        for r, label in zip(t.data_rows, labels):
            if label not in BENCH_SHARE_LABELS:
                continue
            which, key = BENCH_SHARE_LABELS[label]
            src = SHARE_LABELS[1] if which == "a" else SHARE_LABELS[2]
            col = "Share of revenue" if key == "s_rev" else "Share of units"
            for h, rows in (("CA share", ca), ("US share", us)):
                self._cmp(t, r, h, rows[src][col], SHARE_TOL, probs)
                n += 1
        return probs, n

    def _check_fuel(self, t: Table) -> tuple[list[str], int]:
        d = self.data.get()
        mk = t.wb.market
        u = self.union(t.wb)
        core = u[u["_core"]]
        fuels = tuple(C.FEATURE_FUEL_SCOPE)
        fuel = pd.Series("unspecified", index=core.index) if d.fuel_absent.get(mk) else core["fuel_scope"].astype(str)
        probs = []
        bad = sorted(set(fuel) - set(fuels))
        if bad:
            probs.append(f"{mk}: core devices with fuel_scope outside FEATURE_FUEL_SCOPE: {bad}")
        heads = t.columns[1:]
        want = [h for f in fuels for h in (f"{f}: # ASINs", f"{f}: Monthly Rev ({C.MARKETS[mk].currency})", f"{f}: Monthly Units")]
        if heads != want:
            probs.append(f"{t.label}: fuel columns {heads} != {want}")
        by_label = {C.GAUGE_SUBTYPE_LABELS[c]: c for c in C.GAUGE_DEVICE_CLASSES}
        labels = [str(v) for v in t.values("Sub-type")]
        if labels != [C.GAUGE_SUBTYPE_LABELS[c] for c in C.GAUGE_DEVICE_CLASSES]:
            probs.append(f"{t.label}: rows {labels} != the device sub-types")
        if t.total_row is None:
            probs.append(f"{t.label}: no Total row")
        targets = [(r, core[core["gauge_class"] == by_label[lb]]) for r, lb in zip(t.data_rows, labels) if lb in by_label]
        if t.total_row is not None:
            targets.append((t.total_row, core))
        n = 0
        for r, sub in targets:
            f_sub = fuel.loc[sub.index]
            for h in heads:
                f, _, what = h.partition(": ")
                part = sub[f_sub == f]
                if what == "# ASINs":
                    e, tol = float(part["asin"].nunique()), EXACT_TOL
                elif what.startswith("Monthly Rev"):
                    e, tol = float(part["revenue_month"].sum()), MONEY_TOL
                elif what == "Monthly Units":
                    e, tol = float(part["units_month"].sum()), EXACT_TOL
                else:
                    probs.append(f"{t.label}: column {h!r} has no re-derivation rule")
                    continue
                self._cmp(t, r, h, e, tol, probs)
                n += 1
        return probs, n

    def _check_app_matrix(self, t: Table) -> tuple[list[str], int]:
        probs = []
        apps = [v for v in t.values("App")]
        if apps != list(C.APP_FEATURE_MATRIX_APPS):
            probs.append(f"{t.label}: {len(apps)} rows {apps} != APP_FEATURE_MATRIX_APPS {list(C.APP_FEATURE_MATRIX_APPS)}")
        src = t.values("Source")
        bad = [f"{a}: {v!r}" for a, v in zip(apps, src) if not isinstance(v, str)]
        if bad:
            probs.append(f"{t.label}: Source cells that are not strings: {bad[:5]}")
        return probs, len(apps) + len(src)

    def _entities(self, t: Table) -> int | None:
        if t.role.startswith("benchmark_brands"):
            ca, us = self._bench_frames(t)
            return len(set(ca["brand_key"]) | set(us["brand_key"]))
        if t.role.startswith("benchmark"):
            return None
        fr = self.frame_for(t)
        if fr is None:
            return None
        sub = fr[0]
        if t.has("ASIN"):
            return len(sub)
        if t.has("Brand"):
            return int(sub["brand_key"].nunique())
        return None

    def _brand_kpi(self, t: Table) -> list[str]:
        fr = self.frame_for(t)
        if fr is None:
            return [f"{t.label}: kpi table (filter {t.filter!r}) has no re-derivation rule"]
        sub, den = fr
        probs = []
        for r in t.data_rows:
            label = str(t.ws.cell(r, t.first_col).value)
            got = _num(t.ws.cell(r, t.first_col + 1).value)
            mtr = metric(label, sub, den)
            if mtr is None:
                probs.append(f"{t.label} KPI {label!r}: no re-derivation rule")
                continue
            e, tol = mtr
            if (got is None and not math.isnan(e)) or (got is not None and not _close(got, e, tol)):
                probs.append(f"{t.label} KPI {label!r} {got} != {e:,.4f}")
        return probs

    def v10(self):
        probs, n = [], 0
        for d in self.out_dirs.values():
            if not d.is_dir():
                probs.append(f"out dir missing: {d}")
                continue
            n += 1
            probs += [f"lock file {p}" for p in d.rglob("~$*")]
        return _outcome(probs, f"no ~$ lock files in {n} out dirs")

    def v11(self):
        books, probs = self.present_books()
        n = 0
        for w in books:
            for ws in w.book.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        n += 1
                        if isinstance(c.value, str) and c.value.startswith(ERROR_TOKENS):
                            probs.append(f"{w.name}!{ws.title}!{c.coordinate}: {c.value[:20]!r}")
        return _outcome(probs, f"{n} cells scanned in {len(books)} workbooks, no error tokens")

    def v12(self):
        probs, ev = [], []
        for mk in ("CA", "US"):
            w = self.wb(C.gauge_report_name(mk, self.m))
            u = self.union(w)
            core, dev = u[u["_core"]], u[u["gauge_device_scope"]]
            t, rev = self._summary_total(w, "Monthly Rev")
            units = _num(t.cell(t.total_row, "Monthly Units").value)
            n = _num(t.cell(t.total_row, "# of Listings").value)
            for what, got, e, tol in (("revenue", rev, core["revenue_month"].sum(), MONEY_TOL),
                                      ("units", units, core["units_month"].sum(), EXACT_TOL), ("listings", n, len(core), EXACT_TOL)):
                if got is None or not _close(got, float(e), tol):
                    probs.append(f"{mk} Summary core {what} {got} != {float(e):,.2f}")
            kpi = self.titled(w, "kpi", "Summary", KPI_BLOCK_TITLE)
            ccy = C.MARKETS[mk].currency
            adj = u[u["gauge_class"].isin(C.GAUGE_ADJACENT_CLASSES)]
            acc = u[u["gauge_class"].isin(C.GAUGE_ACCESSORY_CLASSES)]
            expect = {f"Core device revenue ({ccy})": (core["revenue_month"].sum(), MONEY_TOL),
                      "Core device units": (core["units_month"].sum(), EXACT_TOL),
                      "Core device # ASINs": (len(core), EXACT_TOL),
                      f"Device revenue incl. borderline ({ccy})": (dev["revenue_month"].sum(), MONEY_TOL),
                      f"Gauge accessories revenue ({ccy})": (acc["revenue_month"].sum(), MONEY_TOL),
                      f"Adjacent GPS-only HUD revenue ({ccy})": (adj["revenue_month"].sum(), MONEY_TOL),
                      "# core device ASINs with sales > 0": ((core["units_month"] > 0).sum(), EXACT_TOL)}
            seen = set()
            for r in kpi.data_rows:
                label = str(kpi.ws.cell(r, 1).value)
                got = _num(kpi.ws.cell(r, 2).value)
                if label not in expect:
                    probs.append(f"{mk} KPI {label!r} not re-derivable")
                    continue
                seen.add(label)
                e, tol = expect[label]
                if got is None or not _close(got, float(e), tol):
                    probs.append(f"{mk} KPI {label!r} {got} != {float(e):,.2f}")
            probs += [f"{mk} KPI {k!r} missing" for k in expect if k not in seen]
            ap = self.one(w, "all_rows", "All Products")
            if ap.n_rows != len(u):
                probs.append(f"{mk} All Products rows {ap.n_rows} != union rows {len(u)}")
            ev.append(f"{mk}: core {len(core)} rows / {float(core['revenue_month'].sum()):,.2f}, incl. borderline {len(dev)} rows "
                      f"/ {float(dev['revenue_month'].sum()):,.2f}, All Products {ap.n_rows} == union {len(u)}")
            if mk == "CA":
                if self.real:
                    bd = int(((u["brand_key"] == "bully dog") & (u["gauge_class"] == "tuner_with_gauge_display")).sum())
                    if bd < 3:
                        probs.append(f"CA Bully Dog tuner_with_gauge_display rows {bd} < 3")
                    ev.append(f"CA Bully Dog TD rows {bd}")
                else:
                    ev.append(f"Bully Dog check skipped ({self.real_reason()})")
        return _outcome(probs, "; ".join(ev))

    def v13(self):
        probs, ev = [], []
        for mk in ("CA", "US"):
            w = self.wb(C.gauge_report_name(mk, self.m))
            u = self.union(w)
            exp = set(u.loc[u["gauge_class"].isin(C.GAUGE_EXCLUDED_CLASSES + ("ambiguous",)), "asin"])
            t = self.one(w, "excluded", "Excluded")
            got = t.values("ASIN")
            if len(got) != len(exp) or set(got) != exp:
                probs.append(f"{mk} Excluded rows {len(got)} != excluded+ambiguous {len(exp)} (diff {sorted(set(got) ^ exp)[:5]})")
            no_rule = [a for a, r in zip(got, t.values("Rule ID")) if _is_blank(r)]
            if no_rule:
                probs.append(f"{mk} Excluded rows without a rule id: {no_rule[:5]}")
            ev.append(f"{mk} {len(got)} rows, all with a rule id")
        return _outcome(probs, "; ".join(ev))

    def v14(self):
        books, probs = self.present_books()
        n = 0
        for w in books:
            for t in self.tables(w):
                if t.role == "dedupe_audit" or not t.has("ASIN"):
                    continue
                n += 1
                vals = [v for v in t.values("ASIN") if v is not None]
                dup = sorted({v for v in vals if vals.count(v) > 1})
                if dup:
                    probs.append(f"{t.label}: duplicate ASINs {dup[:5]}")
        return _outcome(probs, f"{n} registered tables with an ASIN column, no duplicates (dedupe_audit tables exempt)")

    def v15(self):
        books, probs = self.present_books()
        for w in books:
            if "Metadata" not in w.book.sheetnames:
                probs.append(f"{w.name}: no Metadata sheet")
                continue
            ws = w.book["Metadata"]
            kv = {str(ws.cell(r, 1).value): ws.cell(r, 2).value for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value is not None}
            for k in C.METADATA_REQUIRED_KEYS:
                if k not in kv:
                    probs.append(f"{w.name}: Metadata key {k!r} missing")
                elif _is_blank(kv[k]):
                    probs.append(f"{w.name}: Metadata key {k!r} empty")
        return _outcome(probs, f"{len(books)} Metadata sheets carry all {len(C.METADATA_REQUIRED_KEYS)} required keys")

    @staticmethod
    def _sheet_charts(path: Path) -> dict[str, list[ET.Element]]:
        def rels(z: zipfile.ZipFile, part: str) -> dict[str, str]:
            rp = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
            if rp not in z.namelist():
                return {}
            out = {}
            for r in ET.fromstring(z.read(rp)).iter(f"{{{REL_NS}}}Relationship"):
                tgt = r.get("Target")
                out[r.get("Id")] = tgt.lstrip("/") if tgt.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(part), tgt))
            return out

        with zipfile.ZipFile(path) as z:
            wbrels = rels(z, "xl/workbook.xml")
            out: dict[str, list[ET.Element]] = {}
            for s in ET.fromstring(z.read("xl/workbook.xml")).iter(f"{{{MAIN_NS}}}sheet"):
                sheet_part = wbrels[s.get(f"{{{DOC_REL_NS}}}id")]
                charts = []
                for tgt in rels(z, sheet_part).values():
                    if "/drawings/" in tgt and tgt.endswith(".xml"):
                        for ct in rels(z, tgt).values():
                            if "/charts/" in ct:
                                charts.append(ET.fromstring(z.read(ct)))
                out[s.get("name")] = charts
            n_parts = sum(1 for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))
        out["__parts__"] = [None] * n_parts  # type: ignore[list-item]
        return out

    def v16(self):
        books, probs = self.present_books()
        n = 0
        for w in books:
            reg = [(e["sheet"], c) for e in self.registry.get()[w.market]["tables"] if e["workbook"] == w.name for c in e["charts"]]
            found = self._sheet_charts(w.path)
            if len(found.pop("__parts__")) != len(reg):
                probs.append(f"{w.name}: {len(reg)} registered charts but chart parts differ")
            for sheet in sorted(set(found) | {s for s, _ in reg}):
                want = sorted(CHART_KIND[c["type"]] for s, c in reg if s == sheet)
                trees = found.get(sheet, [])
                kinds = []
                for tr in trees:
                    pa = tr.find(f".//{{{CHART_NS}}}plotArea")
                    kinds += [el.tag.split("}")[1] for el in (pa if pa is not None else []) if el.tag.endswith("Chart")]
                if sorted(kinds) != want:
                    probs.append(f"{w.name}!{sheet}: charts {sorted(kinds)} != registered {want}")
                axes_wanted = sum(1 for s, c in reg if s == sheet and c["has_axes"])
                with_axes = [tr for tr in trees if any(tr.find(f".//{{{CHART_NS}}}{a}") is not None for a in AXIS_TAGS)]
                if len(with_axes) != axes_wanted:
                    probs.append(f"{w.name}!{sheet}: {len(with_axes)} charts with axes != registered {axes_wanted}")
                for tr in with_axes:
                    for a in AXIS_TAGS:
                        for ax in tr.iter(f"{{{CHART_NS}}}{a}"):
                            dl = ax.find(f"{{{CHART_NS}}}delete")
                            if dl is None or dl.get("val") not in ("0", "false"):
                                probs.append(f"{w.name}!{sheet}: {a} delete={None if dl is None else dl.get('val')}")
                n += len(trees)
        return _outcome(probs, f"{n} chart parts match the registry per sheet; every axis delete=0")

    def v17(self):
        if not self.real:
            raise Skip(f"real-data-only (loader review queue): {self.real_reason()}")
        d = self.data.get()
        p = C.run_file(Path(self.a.runs_dir), self.m, "type_review", "CA_code_reader")
        if not p.exists():
            return "FAIL", f"review file missing: {p}"
        review = set(pd.read_csv(p, dtype=str, keep_default_na=False, encoding="utf-8-sig")["asin"])
        cr = d.cr
        need = {"default_other": set(cr.loc[cr["type_source"] == "default_other", "asin"]),
                "low_confidence": set(cr.loc[(cr["type_source"] == "token_profile")
                                             & (cr["type_confidence"] < C.TYPE_REVIEW_CONFIDENCE), "asin"])}
        u = d.ca_u
        need["gauge_type_conflict"] = set(u.loc[u["type_conflict"].astype(bool), "asin"]) | set(cr.loc[cr["type_conflict"], "asin"])
        probs = [f"{k}: {len(v - review)} of {len(v)} not in {p.name} (e.g. {sorted(v - review)[:5]})" for k, v in need.items()
                 if v - review]
        return _outcome(probs, f"{p.name} ({len(review)} rows) lists every " +
                        ", ".join(f"{k} ({len(v)})" for k, v in need.items()))

    def v18(self):
        d = self.data.get()
        probs, ev = [], []
        w = self.wb(C.cr_report_name("CA", self.m))
        t = self.one(w, "all_rows", "All ASINs")
        probs += self._reconcile(t, d.cr, {"Brand": "brand_display", "Type": "type", "Price Tier": "price_tier"})
        leaf_bad = []
        for a, typ, price, tier in zip(d.cr["asin"], d.cr["type"], d.cr["price"], d.cr["price_tier"]):
            hits = [lbl for lbl, (types, lo, hi) in C.CR_TIERS.items() if typ in types and lo <= float(price) < hi]
            if len(hits) != 1 or hits[0] != tier:
                leaf_bad.append(f"{a} type {typ} price {price} leaves {hits} tier {tier!r}")
        probs += leaf_bad
        ev.append(f"CR All ASINs {t.n_rows} rows == dataset {len(d.cr)}, one leaf tier each")
        for mk in ("CA", "US"):
            gw = self.wb(C.gauge_report_name(mk, self.m))
            u = self.union(gw).copy()
            u["_tier_v"] = ["" if pd.isna(p) else next(lbl for lbl, lo, hi in C.GAUGE_TIERS if lo <= p < hi) for p in u["price"]]
            ap = self.one(gw, "all_rows", "All Products")
            probs += self._reconcile(ap, u, {"Brand": "brand_display", "Gauge Class": "gauge_class", "Gauge Tier": "_tier_v"})
            ev.append(f"{mk} All Products {ap.n_rows} rows == union {len(u)}")
        return _outcome(probs, "; ".join(ev) + " (units, revenue, brand, type/class, tier)")

    def _reconcile(self, t: Table, frame: pd.DataFrame, text_cols: dict[str, str]) -> list[str]:
        probs = []
        rev_h = next(h for h in t.columns if _base(h) == "Monthly Rev")
        rows = {}
        for r in t.data_rows:
            a = t.cell(r, "ASIN").value
            if a in rows:
                probs.append(f"{t.label}: ASIN {a} twice")
            rows[a] = r
        exp = frame.set_index("asin")
        missing, extra = sorted(set(exp.index) - set(rows)), sorted(set(rows) - set(exp.index))
        if missing:
            probs.append(f"{t.label}: {len(missing)} dataset ASINs missing (e.g. {missing[:5]})")
        if extra:
            probs.append(f"{t.label}: {len(extra)} ASINs not in the dataset (e.g. {extra[:5]})")
        for a, r in rows.items():
            if a not in exp.index:
                continue
            e = exp.loc[a]
            got_rev, got_u = _num(t.cell(r, rev_h).value), _num(t.cell(r, "Monthly Units").value)
            if got_rev is None or not _close(got_rev, float(e["revenue_month"]), 0.005):
                probs.append(f"{t.label} {a}: revenue {got_rev} != {float(e['revenue_month'])}")
            if got_u is None or not _close(got_u, float(e["units_month"]), EXACT_TOL):
                probs.append(f"{t.label} {a}: units {got_u} != {float(e['units_month'])}")
            for h, fld in text_cols.items():
                g = t.cell(r, h).value
                g = "" if g is None else str(g)
                if g != str(e[fld]):
                    probs.append(f"{t.label} {a}: {h} {g!r} != {e[fld]!r}")
        return probs

    def v19(self):
        probs, ev = [], []
        real = self.real
        from ca_market_reports.ca_xlsx_style import sha256_file
        raw_sets = {"cr": [self.a.raw_dir], "gauge_CA": [self.a.raw_dir, self.a.gauge_raw_dir],
                    "gauge_US": [self.a.us_cr_raw_dir, self.a.us_gauge_raw_dir]}
        if real and self.data.get().us_u is not None:
            gw = self.books.get()[C.gauge_report_name("CA", self.m)]
            if not gw.missing and "US Benchmark" in gw.book.sheetnames:
                raw_sets["gauge_CA"] += [self.a.us_cr_raw_dir, self.a.us_gauge_raw_dir]
        maps = {p.name: sha256_file(p) for p in C.MAPS_DIR.glob("*.csv")}
        for key, d in self.out_dirs.items():
            mp = d / C.manifest_name(self.m)
            if not mp.exists():
                probs.append(f"manifest missing: {mp}")
                continue
            man = json.loads(mp.read_text(encoding="utf-8"))
            if man.get("month") != self.m:
                probs.append(f"{mp}: month {man.get('month')!r}")
            outs, ins = man.get("outputs", {}), man.get("inputs", {})
            if mp.name in outs or str(mp.resolve()) in ins:
                probs.append(f"{mp}: manifest lists itself")
            for name, ent in outs.items():
                p = d / name
                if not p.exists():
                    probs.append(f"{mp.name} ({key}): output {name} not on disk")
                elif sha256_file(p) != ent.get("sha256"):
                    probs.append(f"{mp.name} ({key}): sha256 mismatch for output {name}")
            for x in d.glob("*.xlsx"):
                if x.name not in outs and not x.name.startswith("~$"):
                    probs.append(f"{mp.name} ({key}): {x.name} on disk but not in the manifest")
            for path, sha in ins.items():
                p = Path(path)
                if not p.exists():
                    probs.append(f"{mp.name} ({key}): input missing {path}")
                elif sha256_file(p) != sha:
                    probs.append(f"{mp.name} ({key}): sha256 mismatch for input {path}")
            hashed_maps = {Path(k).name: v for k, v in ins.items() if Path(k).parent.name == "maps"}
            miss_maps = sorted(n for n in maps if n not in hashed_maps)
            if miss_maps:
                probs.append(f"{mp.name} ({key}): maps not hashed {miss_maps}")
            stale = sorted(n for n in maps if n in hashed_maps and hashed_maps[n] != maps[n])
            if stale:
                probs.append(f"{mp.name} ({key}): maps changed since the build (rebuild): {stale}")
            if real:
                need = {str(p.resolve()) for rd in raw_sets[key] for p in Path(rd).glob("*.csv")}
                miss = sorted(need - set(ins))
                if miss:
                    probs.append(f"{mp.name} ({key}): {len(miss)} raw CSVs not hashed (e.g. {miss[:2]})")
                ev.append(f"{key}: {len(outs)} outputs, {len(ins)} inputs ({len(need)} raw CSVs)")
            else:
                ev.append(f"{key}: {len(outs)} outputs, {len(ins)} inputs")
        if not real:
            ev.append(f"raw-input hashes skipped ({self.real_reason()})")
        return _outcome(probs, "; ".join(ev) + "; manifests never list themselves")

    # ---------------------------------------------------------------- memo
    _NUM_RE = re.compile(r"(?<![\w.$])(?P<neg>[-−])?(?P<cur>CA\$|US\$|\$)?(?P<int>\d{1,3}(?:,\d{3})+|\d+)(?P<dec>\.\d+)?(?P<pct>%)?(?![\w])")
    _ANY_TAG_RE = re.compile(r"\[(?:WB|SRC):[^\]]*\]")
    _CODE_SPAN_RE = re.compile(r"`[^`]*`")

    @classmethod
    def _last_number(cls, text: str) -> dict | None:
        found = None
        for mm in cls._NUM_RE.finditer(text):
            found = mm
        if found is None:
            return None
        v = float(found.group("int").replace(",", "") + (found.group("dec") or ""))
        if found.group("neg"):
            v = -v
        return {"value": v, "cur": bool(found.group("cur")), "dec": bool(found.group("dec")), "pct": bool(found.group("pct")),
                "text": found.group(0)}

    def _find_workbook(self, name: str) -> Path | None:
        for d in self.out_dirs.values():
            if (d / name).exists():
                return d / name
        return None

    def v20(self):
        memo = Path(self.a.memo)
        if not memo.exists():
            return "FAIL", f"memo missing: {memo}"
        lines = memo.read_text(encoding="utf-8").splitlines()
        sources = Path(self.a.sources)
        src_urls: set[str] | None = None
        if sources.exists():
            src_urls = set(pd.read_csv(sources, dtype=str, keep_default_na=False)["url"].str.strip())
        probs: list[str] = []
        books: dict[Path, Any] = {}
        n_wb = n_src = n_tbd = 0
        for i, raw_line in enumerate(lines, 1):
            line = self._CODE_SPAN_RE.sub(lambda mm: " " * len(mm.group(0)), raw_line)   # `…` spans are literal mentions
            n_tbd += len(re.findall(r"\[WB: TBD\]", line))
            for mm in re.finditer(r"\[WB:[^\]]*\]", line):
                if mm.group(0) != "[WB: TBD]" and not C.MEMO_WB_REF_RE.fullmatch(mm.group(0)):
                    probs.append(f"line {i}: malformed tag {mm.group(0)!r}")
            for mm in C.MEMO_WB_REF_RE.finditer(line):
                n_wb += 1
                tag = f"{mm.group('file')}!{mm.group('sheet')}!{mm.group('cell')}"
                p = self._find_workbook(mm.group("file"))
                if p is None:
                    probs.append(f"line {i} [WB: {tag}]: workbook not in the out dirs")
                    continue
                if p not in books:
                    books[p] = load_workbook(p, data_only=False)
                wb = books[p]
                if mm.group("sheet") not in wb.sheetnames:
                    probs.append(f"line {i} [WB: {tag}]: sheet missing")
                    continue
                cv = wb[mm.group("sheet")][mm.group("cell")].value
                if cv is None:
                    probs.append(f"line {i} [WB: {tag}]: cell empty")
                    continue
                before = self._ANY_TAG_RE.sub(" ", line[:mm.start()])
                tok = self._last_number(before)
                if tok is None:
                    probs.append(f"line {i} [WB: {tag}]: no number before the tag")
                    continue
                if isinstance(cv, str):
                    ctok = self._last_number(cv)
                    if ctok is None or cv.startswith("="):
                        probs.append(f"line {i} [WB: {tag}]: cell {cv[:40]!r} holds no number")
                        continue
                    cell = ctok["value"]
                else:
                    cell = float(cv)
                memo_v = tok["value"]
                if tok["pct"]:
                    memo_v, tol = memo_v / 100.0, 0.001
                elif tok["cur"]:
                    tol = 0.01 if tok["dec"] else 1.0
                elif tok["dec"]:
                    tol = 0.001
                else:
                    tol = 0.0 if float(cell).is_integer() else 1.0
                if abs(memo_v - cell) > tol + 1e-9:
                    probs.append(f"line {i} [WB: {tag}]: memo {tok['text']} != cell {cell:,.4f}")
            for mm in C.MEMO_SRC_REF_RE.finditer(line):
                n_src += 1
                if src_urls is None:
                    probs.append(f"line {i}: [SRC: {mm.group('url')}] but sources CSV missing ({sources})")
                elif mm.group("url").strip() not in src_urls:
                    probs.append(f"line {i}: [SRC: {mm.group('url')}] not in {sources.name}")
        if n_tbd:
            probs.insert(0, f"{n_tbd} [WB: TBD] placeholder(s) left")
        return _outcome(probs, f"{n_wb} [WB:] tags equal their cells; {n_src} [SRC:] urls in {sources.name}; 0 [WB: TBD]")

    def v21(self):
        d = self.data.get()
        probs, ev = [], []
        for name, f in d.frames.items():
            bad = [a for a in f["asin"] if not C.ASIN_RE.match(str(a))]
            if bad:
                probs.append(f"{name}: invalid ASINs {bad[:5]}")
            for c in ("units_month", "revenue_month", "price"):
                x = pd.to_numeric(f[c], errors="coerce").astype(float)
                badn = f.loc[~x.map(math.isfinite) | (x < 0), "asin"].tolist()
                if badn:
                    probs.append(f"{name}: {c} not finite/non-negative for {badn[:5]}")
            types = set(f["type"].fillna("").astype(str)) - set(C.TYPES) - {""}
            if types:
                probs.append(f"{name}: unknown type values {sorted(types)}")
            gc = set(f["gauge_class"].fillna("").astype(str)) - set(C.GAUGE_CLASS_NAMES) - {""}
            if gc:
                probs.append(f"{name}: unknown gauge_class values {sorted(gc)}")
            rev = pd.to_numeric(f["revenue_month"], errors="coerce")
            units = pd.to_numeric(f["units_month"], errors="coerce")
            ev.append(f"{name} {len(f)} rows ({int(((rev > 0) & (units == 0)).sum())} with revenue > 0 and units == 0)")
        for name, u in (("CA union", d.ca_u), ("US union", d.us_u)):
            if u is None:
                continue
            gc = set(u["gauge_class"]) - set(C.GAUGE_CLASS_NAMES)
            if gc:
                probs.append(f"{name}: unknown gauge_class {sorted(gc)}")
        return _outcome(probs, "; ".join(ev))

    def v22(self):
        from ca_market_reports.ca_gauge_classification import classify_gauges, classify_title
        p = C.FIXTURES_DIR / "gauge_golden_titles.csv"
        g = pd.read_csv(p, dtype=str, keep_default_na=False)
        probs = []
        pure = g[g["note"] != "map-only"]
        for _, r in pure.iterrows():
            cls, rid = classify_title(r["title"], r["brand"])
            if cls != r["expected_class"]:
                probs.append(f"{r['note'][:40]}: expected {r['expected_class']} got {cls} ({rid})")
        mo = g[g["note"] == "map-only"].reset_index(drop=True)
        if len(mo) != len(GOLDEN_MAP_ONLY_ASINS):
            probs.append(f"{len(mo)} map-only golden rows but {len(GOLDEN_MAP_ONLY_ASINS)} known seed ASINs")
        else:
            df = pd.DataFrame({"asin": list(GOLDEN_MAP_ONLY_ASINS), "title": mo["title"], "brand_raw": mo["brand"]})
            gm = pd.read_csv(C.MAPS_DIR / "ca_gauge_map.csv", dtype=str, keep_default_na=False)
            out = classify_gauges(df, gm, month=self.m)
            for a, cls, rid, exp in zip(out["asin"], out["gauge_class"], out["gauge_rule_id"], mo["expected_class"]):
                if cls != exp or rid != "MAP":
                    probs.append(f"map-only {a}: expected {exp} via MAP got {cls} ({rid})")
        return _outcome(probs, f"{len(pure)} pure-rule golden titles and {len(mo)} map-only rows classify as expected")

    def v23(self):
        d = self.data.get()
        probs, ev = [], []
        w = self.wb(C.gauge_report_name("CA", self.m))
        if d.us_u is None:
            raise ValueError("no US dataset for the Same-ASIN join")
        t = self.one(w, "same_asin", "US vs CA Same-ASIN")
        both = set(d.ca_u["asin"]) & set(d.us_u["asin"])
        exp = self.same_asin_set()
        got = t.values("ASIN")
        if len(got) != len(exp) or set(got) != exp:
            probs.append(f"Same-ASIN rows {len(got)} != re-derived {len(exp)} (diff {sorted(set(got) ^ exp)[:5]})")
        ev.append(f"Same-ASIN {len(got)} rows == CA∩US device-scope-in-either {len(exp)} (of {len(both)} shared ASINs)")
        dongle = d.cr[d.cr["type"] == "Dongle"]
        n_dongle = int(dongle["asin"].nunique())
        if set(d.modelb["asin"]) != set(dongle["asin"]):
            probs.append(f"Model B universe {len(d.modelb)} != CR Type==Dongle {n_dongle}")
        tiers = self.one(w, "modelb_tiers", "App-Gauge Proxy (Model B)")
        brands = self.one(w, "modelb_brands", "App-Gauge Proxy (Model B)")
        for tt, h in ((tiers, "# ASINs"), (brands, "# of Listings")):
            got_n = _num(tt.cell(tt.total_row, h).value) if tt.total_row else None
            if got_n is None or int(got_n) != n_dongle:
                probs.append(f"{tt.label} Total {h} {got_n} != CR Type==Dongle {n_dongle}")
        ev.append(f"Model B universe {n_dongle} == CR Type==Dongle rows")
        return _outcome(probs, "; ".join(ev))

    # ---------------------------------------------------------------- driver
    def run(self, skip: set[str]) -> list[Result]:
        results = []
        for cid in C.VALIDATION_CHECKS:
            if cid in skip:
                results.append(Result(cid, "SKIP", f"--skip: {SKIPPABLE[cid]}"))
                continue
            try:
                st, ev = getattr(self, cid.lower())()
            except Skip as s:
                st, ev = "SKIP", str(s)
            except Exception as exc:  # a crashed check is a FAIL with the exception text
                tb = traceback.extract_tb(exc.__traceback__)[-1]
                st, ev = "FAIL", f"{type(exc).__name__}: {exc} (at {Path(tb.filename).name}:{tb.lineno})"
            results.append(Result(cid, st, " ".join(str(ev).split())))
        return results


def final_line(results: list[Result]) -> str:
    failed = [r.id for r in results if r.status == "FAIL"]
    passed = sum(r.status == "PASS" for r in results)
    if failed:
        return f"VALIDATION: FAIL ({passed}/{len(results)}; failed: {', '.join(failed)})"
    return f"VALIDATION: PASS ({passed}/{len(results)})"


def parse_args(argv: list[str] | None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Validate the CA code-reader and OBD gauge workbooks (V01-V23)")
    p.add_argument("--month", required=True, help="report month YYYYMM")
    p.add_argument("--cr-out-dir", type=Path, default=C.MARKETS["CA"].cr_out_dir())
    p.add_argument("--gauge-out-dir", type=Path, default=C.MARKETS["CA"].gauge_out_dir())
    p.add_argument("--us-gauge-out-dir", type=Path, default=C.MARKETS["US"].gauge_out_dir())
    p.add_argument("--runs-dir", type=Path, default=C.RUNS_DIR)
    p.add_argument("--raw-dir", type=Path, help="CA code-reader raw folder (default NewProductCategory/CA-CODE-READER/raw_data/<m>)")
    p.add_argument("--gauge-raw-dir", type=Path, help="CA gauge raw folder (default NewProductCategory/CA-OBD-GAUGE/raw_data/<m>)")
    p.add_argument("--us-gauge-raw-dir", type=Path, help="US gauge raw folder (default NewProductCategory/US-OBD-GAUGE/raw_data/<m>)")
    p.add_argument("--us-cr-raw-dir", type=Path, help="US code-reader raw folder (default Amazon_Raw_Data/raw_data/<m>)")
    p.add_argument("--memo", type=Path, help="memo markdown (default ca_market_reports/memo/CA_OBD_Gauge_Market_Memo_<m>.md)")
    p.add_argument("--sources", type=Path, help="memo sources CSV (default ca_market_reports/memo/sources_<m>.csv)")
    p.add_argument("--skip", default="", help=f"comma-separated check IDs to skip; only {sorted(SKIPPABLE)} (documented reasons)")
    p.add_argument("--json", type=Path, help="write machine-readable results to this path")
    p.add_argument("--from-normalized", type=Path, help="DEV: CA dataset from a normalized fixture CSV instead of the raw exports")
    p.add_argument("--us-from-normalized", type=Path, help="DEV: US dataset from a normalized fixture CSV")
    a = p.parse_args(argv)
    if not C.MONTH_RE.match(a.month):
        p.error(f"--month must be a calendar month YYYYMM, got {a.month!r}")
    a.raw_dir = a.raw_dir or C.MARKETS["CA"].cr_raw_dir(a.month)
    a.gauge_raw_dir = a.gauge_raw_dir or C.MARKETS["CA"].gauge_raw_dir(a.month)
    a.us_gauge_raw_dir = a.us_gauge_raw_dir or C.MARKETS["US"].gauge_raw_dir(a.month)
    a.us_cr_raw_dir = a.us_cr_raw_dir or C.MARKETS["US"].cr_raw_dir(a.month)
    a.memo = a.memo or C.MEMO_DIR / C.memo_name(a.month)
    a.sources = a.sources or C.MEMO_DIR / f"sources_{a.month}.csv"
    if a.us_from_normalized and not a.from_normalized:
        p.error("--us-from-normalized needs --from-normalized (dev and raw inputs are never mixed)")
    skip = {s.strip().upper() for s in a.skip.split(",") if s.strip()}
    bad = sorted(skip - set(SKIPPABLE))
    if bad:
        p.error(f"--skip {bad}: only {sorted(SKIPPABLE)} may be skipped ({'; '.join(f'{k}: {v}' for k, v in SKIPPABLE.items())})")
    a.skip_set = skip
    return a


def main(argv: list[str] | None = None) -> int:
    a = parse_args(argv)
    try:
        results = Validator(a).run(a.skip_set)
    except Exception as exc:  # never escapes: every check FAILs with the setup error
        results = [Result(cid, "FAIL", f"validator setup: {type(exc).__name__}: {exc}") for cid in C.VALIDATION_CHECKS]
    final = final_line(results)
    for r in results:
        print(r.line)
    print(final)
    rc = 1 if any(r.status == "FAIL" for r in results) else 0
    if a.json:
        try:
            Path(a.json).write_text(json.dumps({"month": a.month, "checks": [r.__dict__ for r in results], "final": final,
                                                "exit_code": rc}, indent=1, ensure_ascii=False), encoding="utf-8")
        except OSError as exc:
            print(f"ERROR: cannot write --json {a.json}: {exc}", file=sys.stderr)
            return 1
    assert C.VALIDATION_FINAL_RE.match(final), final
    return rc


if __name__ == "__main__":
    sys.exit(main())
