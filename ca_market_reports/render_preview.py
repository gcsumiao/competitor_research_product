"""Static HTML preview + markdown excerpt of a built workbook (read-only; the workbook is never written).

  ca_market_reports/run.sh ca_market_reports/render_preview.py --workbook <xlsx> --out <html> [--excerpt <md>]

HTML: Summary (visible rows), Top 50 (banner + header + first 20 data rows of the first table) and Metadata, with the
workbook's fills, font colours and number formats. Markdown: <= 20 rows of the Summary's first table + top 10 of Top 50.
"""
from __future__ import annotations

import argparse
import html
import re
import sys
from pathlib import Path

import openpyxl

from ca_market_reports import ca_common as C

TOP50_ROWS = 20
EXCERPT_SUMMARY_ROWS = 20
EXCERPT_TOP_ROWS = 10
_LINK_RE = re.compile(r'^=HYPERLINK\("(?P<url>[^"]+)","(?P<text>[^"]+)"\)$')


def _rgb(color) -> str | None:
    if color is None or color.type != "rgb" or not isinstance(color.rgb, str):
        return None
    return "#" + color.rgb[-6:]


def fmt_value(cell) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, str):
        m = _LINK_RE.match(v)
        return m.group("text") if m else v
    if isinstance(v, bool):
        return "Y" if v else "N"
    fmt = cell.number_format or "General"
    if isinstance(v, (int, float)):
        if "%" in fmt:
            return f"{v * 100:.2f}%" if "0.00%" in fmt else f"{v * 100:.1f}%"
        sym = "CA$" if '"CA$"' in fmt else ("$" if '"$"' in fmt else "")
        if sym:
            return f"{sym}{v:,.2f}" if ".00" in fmt else f"{sym}{v:,.0f}"
        if fmt == "0.0":
            return f"{v:.1f}"
        if fmt == "#,##0":
            return f"{v:,.0f}"
        return f"{v:g}" if isinstance(v, float) else str(v)
    return str(v)


def _cell_html(cell, colspan: int) -> str:
    style = []
    if cell.fill is not None and cell.fill.patternType == "solid":
        bg = _rgb(cell.fill.fgColor)
        if bg:
            style.append(f"background:{bg}")
    f = cell.font
    if f is not None:
        fc = _rgb(f.color)
        if fc and fc.lower() != "#000000":
            style.append(f"color:{fc}")
        if f.b:
            style.append("font-weight:600")
        if f.i:
            style.append("font-style:italic")
        if f.sz and f.sz >= 13:
            style.append(f"font-size:{f.sz + 2:.0f}px")
    text = fmt_value(cell)
    v = cell.value
    m = _LINK_RE.match(v) if isinstance(v, str) else None
    body = f'<a href="{html.escape(m.group("url"))}">{html.escape(m.group("text"))}</a>' if m else html.escape(text)
    num = isinstance(v, (int, float)) and not isinstance(v, bool)
    if num:
        style.append("text-align:right")
    span = f' colspan="{colspan}"' if colspan > 1 else ""
    return f'<td{span} style="{";".join(style)}">{body}</td>'


def _merges(ws) -> tuple[dict[tuple[int, int], int], set[tuple[int, int]]]:
    spans, covered = {}, set()
    for rng in ws.merged_cells.ranges:
        spans[(rng.min_row, rng.min_col)] = rng.max_col - rng.min_col + 1
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    covered.add((r, c))
    return spans, covered


def first_table(ws) -> tuple[int, int, int]:
    """(header_row, first_data_row, last_row) of the first table: header = first row whose A cell has the header fill."""
    for r in range(1, ws.max_row + 1):
        c = ws.cell(r, 1)
        if c.fill is not None and c.fill.patternType == "solid" and _rgb(c.fill.fgColor) == "#" + C.FILL_HEADER and c.value:
            last = r
            while last + 1 <= ws.max_row and any(ws.cell(last + 1, k).value is not None for k in range(1, ws.max_column + 1)):
                last += 1
            return r, r + 1, last
    raise ValueError(f"{ws.title}: no header row found")


def _table_width(ws, header_row: int) -> int:
    n = 0
    for c in range(1, ws.max_column + 1):
        if ws.cell(header_row, c).value is not None:
            n = c
    return n


def sheet_html(ws, rows: list[int], ncols: int) -> str:
    spans, covered = _merges(ws)
    out = [f'<section><h2>{html.escape(ws.title)}</h2><table>']
    hidden = 0
    for r in rows:
        if ws.row_dimensions[r].hidden:
            hidden += 1
            continue
        tds = []
        for c in range(1, ncols + 1):
            if (r, c) in covered:
                continue
            span = min(spans.get((r, c), 1), ncols - c + 1)
            tds.append(_cell_html(ws.cell(r, c), span))
        out.append("<tr>" + "".join(tds) + "</tr>")
    out.append("</table>")
    notes = []
    if hidden:
        notes.append(f"{hidden} zero-revenue row(s) hidden in the workbook")
    if ws._charts:
        notes.append(f"{len(ws._charts)} chart(s) on this sheet in the workbook (not rendered here)")
    if notes:
        out.append(f'<p class="note">{html.escape("; ".join(notes))}</p>')
    out.append("</section>")
    return "\n".join(out)


def render_html(path: Path) -> str:
    wb = openpyxl.load_workbook(path)
    parts = []
    if "Summary" in wb.sheetnames:
        ws = wb["Summary"]
        width = max(_table_width(ws, first_table(ws)[0]), 2)
        parts.append(sheet_html(ws, list(range(1, ws.max_row + 1)), width))
    if "Top 50" in wb.sheetnames:
        ws = wb["Top 50"]
        h, f, last = first_table(ws)
        parts.append(sheet_html(ws, list(range(1, h + 1)) + list(range(f, min(last, f + TOP50_ROWS - 1) + 1)), _table_width(ws, h)))
    if "Metadata" in wb.sheetnames:
        ws = wb["Metadata"]
        parts.append(sheet_html(ws, list(range(1, ws.max_row + 1)), 2))
    css = ("body{font-family:Calibri,'Segoe UI',Arial,sans-serif;font-size:13px;margin:16px;background:#fff;color:#111}"
           "table{border-collapse:collapse;margin-bottom:8px}td{border:1px solid #d0d0d0;padding:3px 6px;white-space:nowrap;"
           "max-width:520px;overflow:hidden;text-overflow:ellipsis;vertical-align:top}a{color:#" + C.COLOR_LINK + "}"
           "h1{font-size:18px}h2{font-size:15px;margin:22px 0 6px}.note{color:#555;font-style:italic;margin:0 0 12px}")
    return (f"<!doctype html><html lang=\"en\"><head><meta charset=\"utf-8\"><meta name=\"viewport\" content=\"width=device-width, "
            f"initial-scale=1\"><title>{html.escape(path.stem)}</title><style>{css}</style></head><body>"
            f"<h1>{html.escape(path.name)}</h1><p class=\"note\">Static preview: Summary, Top 50 (first {TOP50_ROWS} rows), "
            f"Metadata.</p>{''.join(parts)}</body></html>")


def _md_row(vals: list[str]) -> str:
    return "| " + " | ".join(v.replace("|", "\\|").replace("\n", " ") for v in vals) + " |"


def render_excerpt(path: Path) -> str:
    wb = openpyxl.load_workbook(path)
    out = [f"# {path.name} — excerpt", ""]
    ws = wb["Summary"]
    h, f, last = first_table(ws)
    n = _table_width(ws, h)
    visible = [r for r in range(f, last + 1) if not ws.row_dimensions[r].hidden]
    if len(visible) > EXCERPT_SUMMARY_ROWS:
        visible = visible[:EXCERPT_SUMMARY_ROWS - 1] + visible[-1:]
    out += [f"## Summary (first table, {len(visible)} rows shown)", "",
            _md_row([str(ws.cell(h, c).value) for c in range(1, n + 1)]), _md_row(["---"] * n)]
    out += [_md_row([fmt_value(ws.cell(r, c)) for c in range(1, n + 1)]) for r in visible]
    ws = wb["Top 50"]
    h, f, last = first_table(ws)
    n = _table_width(ws, h)
    keep = [c for c in range(1, n + 1) if ws.cell(h, c).value not in ("URL", "Link")][:10]
    rows = list(range(f, min(last, f + EXCERPT_TOP_ROWS - 1) + 1))
    out += ["", f"## Top 50 (top {len(rows)})", "", _md_row([str(ws.cell(h, c).value) for c in keep]), _md_row(["---"] * len(keep))]
    out += [_md_row([fmt_value(ws.cell(r, c)) for c in keep]) for r in rows]
    return "\n".join(out) + "\n"


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Render a static HTML preview (and optional markdown excerpt) of a workbook")
    p.add_argument("--workbook", type=Path, required=True)
    p.add_argument("--out", type=Path, required=True, help="HTML output path")
    p.add_argument("--excerpt", type=Path, help="markdown excerpt output path")
    a = p.parse_args(argv)
    if not a.workbook.is_file():
        raise FileNotFoundError(a.workbook)
    a.out.parent.mkdir(parents=True, exist_ok=True)
    a.out.write_text(render_html(a.workbook), encoding="utf-8")
    print(f"preview: {a.out}")
    if a.excerpt:
        a.excerpt.parent.mkdir(parents=True, exist_ok=True)
        a.excerpt.write_text(render_excerpt(a.workbook), encoding="utf-8")
        print(f"excerpt: {a.excerpt}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
